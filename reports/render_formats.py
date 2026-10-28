"""Рендереры HTML, Markdown, XLSX и CSV из одного промежуточного документа."""
from __future__ import annotations

import base64
import datetime as dt
import html
import io
import re
import struct
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

from utils.export import CsvWriter, XLSX_MAX_CELL_CHARS, export_text
from utils.files import sanitize_filename
from .document import (Cell, DocBullets, DocChart, DocHeading, DocKeyValues, DocPageBreak, DocParagraph, DocTable,
                       DocTitlePage, DocToc, ReportDoc)


@dataclass
class RenderResult:
    files: list[Path] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _lang(doc: ReportDoc) -> str:
    return "ru" if (doc.formatter is None or doc.formatter.ru) else "en"


def _png_size(png: bytes) -> tuple[int, int]:
    return struct.unpack(">II", png[16:24])


# =================================================================== HTML
def render_html(doc: ReportDoc, stage: Path, stem: str, **_opts) -> RenderResult:
    st, t = doc.style, doc.template
    accent = t.accent or st.accent
    font = t.font_name or st.font_name
    esc = html.escape
    heads = [e for e in doc.elements if isinstance(e, DocHeading)]
    ids = {id(h): f"h{i + 1}" for i, h in enumerate(heads)}
    ru = _lang(doc) == "ru"
    out: list[str] = []
    zebra = f"tbody tr:nth-child(even) td {{ background: {st.zebra}; }}" if st.zebra else ""
    border = {"grid": f"1px solid {st.border}", "horizontal": "none", "rules": "none", "minimal": "none"}[st.table_style]
    row_line = {"grid": "", "horizontal": f"border-bottom: 1px solid {st.border};", "rules": "",
                "minimal": f"border-bottom: 1px solid {st.border};"}[st.table_style]
    head_style = f"background: {st.header_bg}; color: {st.header_fg};" if st.header_bg.upper() != "#FFFFFF" else \
        f"border-bottom: 1.5px solid #000; font-weight: 700;"
    out.append(f"""<!DOCTYPE html>
<html lang="{_lang(doc)}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{esc(doc.title)}</title>
<style>
:root {{ --accent: {accent}; --muted: {st.muted}; --border: {st.border}; }}
body {{ font-family: '{font}', 'Segoe UI', Arial, sans-serif; font-size: {max(12, int(t.font_size or st.font_size) + 3)}px;
  line-height: {st.line_spacing + 0.3:.2f}; color: #1a1a1a; max-width: 1100px; margin: 0 auto; padding: 32px 24px; }}
h1, h2, h3 {{ color: {accent if st.id not in ('academic', 'minimal') else st.heading_color};
  font-family: '{st.heading_font_name}', 'Segoe UI', Arial, sans-serif; margin-top: 1.6em; }}
.title-page {{ min-height: 40vh; display: flex; flex-direction: column; justify-content: center;
  text-align: {st.title_align}; border-bottom: 3px solid var(--accent); margin-bottom: 2em; padding-bottom: 2em; }}
.title-page h1 {{ font-size: 2.4em; margin: .2em 0; }} .muted {{ color: var(--muted); }}
nav.toc {{ background: #f7f7f7; padding: 8px 20px; border-radius: 6px; }} nav.toc ul {{ list-style: none; padding-left: 1em; }}
.table-wrap {{ overflow-x: auto; margin: 10px 0 4px; }}
table {{ border-collapse: collapse; width: 100%; font-family: '{st.table_font_name}', monospace, sans-serif;
  font-size: {st.table_font_size + 2}px; }}
th, td {{ padding: 4px 8px; border: {border}; {row_line} vertical-align: top; }}
thead th {{ {head_style} text-align: left; position: sticky; top: 0; }}
td.num, th.num {{ text-align: right; white-space: nowrap; }} td.empty {{ color: var(--muted); font-style: italic; }}
tfoot td {{ font-weight: 700; border-top: {'3px double #000' if st.totals_double_rule else '2px solid ' + accent}; }}
{zebra}
.note {{ color: var(--muted); font-style: italic; font-size: .9em; margin: 4px 0 14px; }}
pre {{ background: #f4f4f4; padding: 10px; border-left: 3px solid var(--border); overflow-x: auto; }}
img {{ max-width: 100%; height: auto; display: block; margin: 12px auto; }}
dl.kv {{ display: grid; grid-template-columns: max-content 1fr; gap: 4px 18px; }} dl.kv dt {{ color: var(--muted); font-weight: 600; }}
dl.kv dd {{ margin: 0; }}
footer {{ margin-top: 3em; padding-top: 1em; border-top: 1px solid var(--border); color: var(--muted); font-size: .85em; }}
@media print {{ body {{ padding: 0; max-width: none; }} .pb {{ page-break-after: always; }} thead th {{ position: static; }}
  tr {{ page-break-inside: avoid; }} thead {{ display: table-header-group; }} }}
</style></head><body>""")
    for e in doc.elements:
        if isinstance(e, DocTitlePage):
            out.append('<header class="title-page">')
            if e.organization:
                out.append(f'<div class="muted"><b>{esc(e.organization)}</b></div>')
            out.append(f"<h1>{esc(e.title)}</h1>")
            if e.subtitle:
                out.append(f'<div class="muted" style="font-size:1.3em">{esc(e.subtitle)}</div>')
            if e.part_label:
                out.append(f"<div><b>{esc(e.part_label)}</b></div>")
            meta = []
            if e.author:
                meta.append(("Автор: " if ru else "Author: ") + e.author)
            meta.append(("Дата формирования: " if ru else "Generated: ") + e.generated)
            out.append('<p class="muted">' + "<br>".join(esc(m) for m in meta) + "</p></header>")
        elif isinstance(e, DocToc):
            out.append(f'<nav class="toc"><h2 style="margin-top:.4em">{"Содержание" if ru else "Contents"}</h2><ul>')
            for h in heads:
                if h.level <= 3:
                    out.append(f'<li style="margin-left:{(h.level - 1) * 1.2}em"><a href="#{ids[id(h)]}">{esc(h.text)}</a></li>')
            out.append("</ul></nav>")
        elif isinstance(e, DocHeading):
            lvl = min(3, max(1, e.level)) + (0 if e.level > 1 else 0)
            out.append(f'<h{lvl} id="{ids[id(e)]}">{esc(e.text)}</h{lvl}>')
        elif isinstance(e, DocParagraph):
            if e.style == "code":
                out.append(f"<pre>{esc(e.text)}</pre>")
            elif e.style in ("note", "caption"):
                out.append(f'<p class="note">{esc(e.text)}</p>')
            else:
                out.append(f"<p>{esc(e.text)}</p>")
        elif isinstance(e, DocBullets):
            out.append("<ul>" + "".join(f"<li>{esc(i)}</li>" for i in e.items) + "</ul>")
        elif isinstance(e, DocChart):
            b64 = base64.b64encode(e.png).decode("ascii")
            out.append(f'<img alt="{esc(e.alt or e.title)}" src="data:image/png;base64,{b64}">')
        elif isinstance(e, DocKeyValues):
            if e.title:
                out.append(f"<h2>{esc(e.title)}</h2>")
            out.append('<dl class="kv">' + "".join(f"<dt>{esc(k)}</dt><dd>{esc(v)}</dd>" for k, v in e.items) + "</dl>")
        elif isinstance(e, DocPageBreak):
            out.append('<div class="pb"></div>')
        elif isinstance(e, DocTable):
            if e.title and e.show_title:
                out.append(f"<p><b>{esc(e.title)}</b></p>")
            out.append('<div class="table-wrap"><table><thead><tr>' + "".join(
                f'<th class="{"num" if c.align == "right" else ""}">{esc(c.title)}</th>' for c in e.columns)
                + "</tr></thead><tbody>")
            for row in e.rows:
                out.append("<tr>" + "".join(
                    (f'<td class="empty">—</td>' if c.empty else
                     f'<td class="{"num" if c.align == "right" else ""}">{esc(c.text)}</td>') for c in row) + "</tr>")
            out.append("</tbody>")
            if e.totals:
                out.append("<tfoot><tr>" + "".join(
                    f'<td class="{"num" if c.align == "right" else ""}">{esc(c.text)}</td>' for c in e.totals)
                    + "</tr></tfoot>")
            out.append("</table></div>")
            if e.note:
                out.append(f'<p class="note">{esc(e.note)}</p>')
    foot = " · ".join(x for x in (doc.footer, doc.header) if x)
    out.append(f"<footer>{esc(foot)}</footer></body></html>" if foot else "</body></html>")
    path = stage / f"{stem}.html"
    path.write_text("\n".join(out), encoding="utf-8")
    return RenderResult([path], [])


# ============================================================== Markdown
def _md(text: str) -> str:
    return (text.replace("\\", "\\\\").replace("|", "\\|").replace("<", "&lt;").replace(">", "&gt;")
            .replace("\r\n", "<br>").replace("\n", "<br>"))


def render_md(doc: ReportDoc, stage: Path, stem: str, **_opts) -> RenderResult:
    res = RenderResult()
    ru = _lang(doc) == "ru"
    lines: list[str] = []
    images = 0
    assets = stage / f"{stem}_files"
    for e in doc.elements:
        if isinstance(e, DocTitlePage):
            lines += [f"# {_md(e.title)}", ""]
            if e.subtitle:
                lines += [f"*{_md(e.subtitle)}*", ""]
            meta = [f"{'Организация' if ru else 'Organization'}: {e.organization}" if e.organization else "",
                    f"{'Автор' if ru else 'Author'}: {e.author}" if e.author else "",
                    f"{'Дата формирования' if ru else 'Generated'}: {e.generated}", e.part_label]
            lines += [" \n".join(_md(m) for m in meta if m), ""]
        elif isinstance(e, DocToc):
            lines += [f"## {'Содержание' if ru else 'Contents'}", ""]
            for h in doc.headings(3):
                lines.append("  " * (h.level - 1) + f"- {_md(h.text)}")
            lines.append("")
        elif isinstance(e, DocHeading):
            lines += [("#" * min(6, e.level + 1)) + " " + _md(e.text), ""]
        elif isinstance(e, DocParagraph):
            if e.style == "code":
                lines += ["```sql", e.text.replace("```", "'''"), "```", ""]
            elif e.style in ("note", "caption"):
                lines += [f"*{_md(e.text)}*", ""]
            else:
                lines += [_md(e.text).replace("<br>", "  \n"), ""]
        elif isinstance(e, DocBullets):
            lines += [f"- {_md(i)}" for i in e.items] + [""]
        elif isinstance(e, DocChart):
            images += 1
            assets.mkdir(parents=True, exist_ok=True)
            name = f"chart_{images:02d}.png"
            (assets / name).write_bytes(e.png)
            res.files.append(assets / name)
            lines += [f"![{_md(e.alt or e.title)}]({stem}_files/{name})", ""]
        elif isinstance(e, DocKeyValues):
            lines += [f"| {'Параметр' if ru else 'Item'} | {'Значение' if ru else 'Value'} |", "| --- | --- |"]
            lines += [f"| {_md(k)} | {_md(v)} |" for k, v in e.items] + [""]
        elif isinstance(e, DocPageBreak):
            lines += ["---", ""]
        elif isinstance(e, DocTable):
            if e.title and e.show_title:
                lines += [f"**{_md(e.title)}**", ""]
            lines.append("| " + " | ".join(_md(c.title) for c in e.columns) + " |")
            lines.append("| " + " | ".join("---:" if c.align == "right" else "---" for c in e.columns) + " |")
            for row in e.rows:
                lines.append("| " + " | ".join("" if c.empty else _md(c.text) for c in row) + " |")
            if e.totals:
                lines.append("| " + " | ".join(f"**{_md(c.text)}**" if c.text else "" for c in e.totals) + " |")
            lines.append("")
            if e.note:
                lines += [f"*{_md(e.note)}*", ""]
    path = stage / f"{stem}.md"
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    res.files.insert(0, path)
    return res


# ================================================================== XLSX
def _sheet_name(base: str, used: set[str]) -> str:
    name = re.sub(r"[\[\]:*?/\\]", "_", base).strip().strip("'") or "Лист"
    name = name[:31]
    n, cand = 1, name
    while cand.lower() in used:
        n += 1
        suffix = f" ({n})"
        cand = name[:31 - len(suffix)] + suffix
    used.add(cand.lower())
    return cand


def _num_format(col, ru: bool, symbol: str) -> str | None:
    if col.fmt == "percent":
        return "0.0%"
    if col.fmt == "currency":
        return '#,##0.00 "' + symbol + '"' if symbol else "#,##0.00"
    if col.kind == "integer":
        return "0" if "№" in col.title or col.title.lower() in ("id", "год", "year") else "#,##0"
    if col.kind == "numeric" or col.fmt in ("decimal",):
        return "#,##0.00"
    if col.kind == "datetime":
        return "DD.MM.YYYY HH:MM" if ru else "YYYY-MM-DD HH:MM"
    if col.kind == "date":
        return "DD.MM.YYYY" if ru else "YYYY-MM-DD"
    return None


def render_xlsx(doc: ReportDoc, stage: Path, stem: str, **opts) -> RenderResult:
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    res = RenderResult()
    st, ru = doc.style, _lang(doc) == "ru"
    symbol = doc.template.currency_symbol
    wb = Workbook()
    used: set[str] = set()
    main = wb.active
    main.title = _sheet_name("Отчёт" if ru else "Report", used)
    head_fill = PatternFill("solid", fgColor=st.header_bg.lstrip("#"))
    head_font = Font(bold=True, color=st.header_fg.lstrip("#") if st.header_bg.upper() != "#FFFFFF" else "000000")
    zebra_fill = PatternFill("solid", fgColor=st.zebra.lstrip("#")) if st.zebra else None
    thin = Side(style="thin", color=st.border.lstrip("#"))
    trunc = [0]

    def put(ws, row, col, value, number_format=None, bold=False, italic=False, fill=None, align=None):
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub("", value)
            if len(value) > XLSX_MAX_CELL_CHARS:
                value, trunc[0] = value[:XLSX_MAX_CELL_CHARS], trunc[0] + 1
            c = ws.cell(row=row, column=col, value=value)
            c.data_type = "s"                              # текст никогда не становится формулой
        else:
            if isinstance(value, Decimal):
                value = float(value) if abs(value) < 10 ** 15 else str(value)
            elif isinstance(value, dt.datetime) and value.tzinfo:
                value = value.replace(tzinfo=None)
            elif isinstance(value, int) and not isinstance(value, bool) and abs(value) >= 10 ** 15:
                value = str(value)
            c = ws.cell(row=row, column=col, value=value)
        if number_format and not isinstance(value, str):
            c.number_format = number_format
        if bold or italic:
            c.font = Font(bold=bold, italic=italic)
        if fill:
            c.fill = fill
        if align:
            c.alignment = Alignment(horizontal=align, vertical="top", wrap_text=False)
        return c

    def write_table(ws, t: DocTable, top: int, with_title: bool = True) -> int:
        r = top
        if t.title and with_title:
            put(ws, r, 1, t.title, bold=True)
            r += 1
        for ci, col in enumerate(t.columns, 1):
            c = put(ws, r, ci, col.title, bold=True, fill=head_fill if st.header_bg.upper() != "#FFFFFF" else None)
            c.font = head_font
            c.alignment = Alignment(horizontal="center" if col.align == "right" else "left", vertical="center", wrap_text=True)
        header_row = r
        formats = [_num_format(c, ru, symbol) for c in t.columns]
        for ri, row in enumerate(t.rows):
            r += 1
            for ci, cell in enumerate(row, 1):
                raw = cell.raw
                if cell.empty or raw is None:
                    put(ws, r, ci, None)
                elif isinstance(raw, (int, float, Decimal)) and not isinstance(raw, bool):
                    put(ws, r, ci, raw, formats[ci - 1])
                elif isinstance(raw, (dt.date, dt.datetime)):
                    put(ws, r, ci, raw, formats[ci - 1] or "DD.MM.YYYY")
                elif isinstance(raw, bool):
                    put(ws, r, ci, cell.text)
                else:
                    put(ws, r, ci, cell.text if raw is not None else "")
                if zebra_fill is not None and ri % 2 == 1 and len(t.rows) <= 20000:
                    ws.cell(row=r, column=ci).fill = zebra_fill
        if t.totals:
            r += 1
            for ci, c in enumerate(t.totals, 1):
                cell = put(ws, r, ci, c.raw if isinstance(c.raw, (int, float, Decimal)) else c.text, formats[ci - 1], bold=True)
                cell.border = Border(top=Side(style="double" if st.totals_double_rule else "medium", color="000000"))
        for ci, col in enumerate(t.columns, 1):
            letter = get_column_letter(ci)
            ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, min(60, max(9, col.width + 2)))
        if t.note:
            r += 1
            put(ws, r, 1, t.note, italic=True)
        return header_row

    # главный лист: титул, повествование (заголовки, текст, выводы)
    r = 1
    narrative: list = []
    stats_sheet = None
    stats_row = 1
    charts_sheet = None
    chart_row = 1
    pending_heading = ""
    table_n = 0
    for e in doc.elements:
        if isinstance(e, DocTitlePage):
            put(main, r, 1, e.title, bold=True)
            main.cell(row=r, column=1).font = Font(bold=True, size=16)
            r += 1
            for text in (e.subtitle, e.organization, ("Автор: " if ru else "Author: ") + e.author if e.author else "",
                         ("Дата формирования: " if ru else "Generated: ") + e.generated, e.part_label):
                if text:
                    put(main, r, 1, text)
                    r += 1
            r += 1
        elif isinstance(e, DocHeading):
            pending_heading = e.text
            put(main, r, 1, e.text, bold=True)
            main.cell(row=r, column=1).font = Font(bold=True, size=14 - min(e.level, 3))
            r += 1
        elif isinstance(e, DocParagraph):
            put(main, r, 1, e.text, italic=e.style in ("note", "caption"))
            r += 1
        elif isinstance(e, DocBullets):
            for it in e.items:
                put(main, r, 1, "• " + it)
                r += 1
        elif isinstance(e, DocKeyValues):
            for k, v in e.items:
                put(main, r, 1, k, bold=True)
                put(main, r, 2, v)
                r += 1
            r += 1
        elif isinstance(e, DocChart):
            if charts_sheet is None:
                charts_sheet = wb.create_sheet(_sheet_name("Диаграммы" if ru else "Charts", used))
            put(charts_sheet, chart_row, 1, e.title, bold=True)
            img = XlImage(io.BytesIO(e.png))
            w, h = _png_size(e.png)
            scale = min(1.0, 720 / w)
            img.width, img.height = int(w * scale), int(h * scale)
            charts_sheet.add_image(img, f"A{chart_row + 1}")
            chart_row += int(img.height / 20) + 3
        elif isinstance(e, DocTable):
            if e.compact:
                if stats_sheet is None:
                    stats_sheet = wb.create_sheet(_sheet_name("Статистика" if ru else "Statistics", used))
                write_table(stats_sheet, e, stats_row)
                stats_row += len(e.rows) + 4 + (1 if e.totals else 0) + (1 if e.note else 0)
            else:
                table_n += 1
                base = e.title if (e.title and e.show_title) else (pending_heading or f"Таблица {table_n}")
                ws = wb.create_sheet(_sheet_name(base, used))
                hdr = write_table(ws, e, 1, with_title=False)      # на листе данных шапка — в первой строке
                ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
                if e.rows:
                    ws.auto_filter.ref = f"A{hdr}:{get_column_letter(len(e.columns))}{hdr + len(e.rows)}"
                put(main, r, 1, f"→ лист «{ws.title}»" if ru else f"→ sheet '{ws.title}'", italic=True)
                r += 1
        elif isinstance(e, DocToc) or isinstance(e, DocPageBreak):
            continue
    main.column_dimensions["A"].width = 60
    main.column_dimensions["B"].width = 45
    for ws in wb.worksheets:
        if ws.max_row > 1_048_576:
            res.warnings.append(f"Лист «{ws.title}» превышает лимит строк Excel.")
    if trunc[0]:
        res.warnings.append(f"{trunc[0]} знач. усечены до {XLSX_MAX_CELL_CHARS} символов (лимит Excel).")
    path = stage / f"{stem}.xlsx"
    wb.save(path)
    res.files.append(path)
    return res


# =================================================================== CSV
def render_csv(doc: ReportDoc, stage: Path, stem: str, *, delimiter: str = ";", formula_guard: bool = True,
               **_opts) -> RenderResult:
    """CSV: по одному файлу на каждую таблицу данных (без таблиц статистики и итоговых строк)."""
    res = RenderResult()
    tables = [t for t in doc.tables() if not t.compact]
    if not tables:
        res.warnings.append("CSV: в отчёте нет таблиц данных — файлы не созданы.")
        return res
    heading = ""
    names: dict[int, str] = {}
    for e in doc.elements:
        if isinstance(e, DocHeading):
            heading = e.text
        elif isinstance(e, DocTable) and not e.compact:
            names[id(e)] = e.title if (e.title and e.show_title) else heading
    for i, t in enumerate(tables, 1):
        if len(tables) == 1:
            name = f"{stem}.csv"
        else:
            title = sanitize_filename(names.get(id(t)) or f"table{i}", f"table{i}")[:60]
            name = f"{stem}_{i:02d}_{title}.csv"
        w = CsvWriter(stage / name, [c.title for c in t.columns], delimiter=delimiter, formula_guard=formula_guard)
        try:
            w.write_rows([[c.raw if not c.empty else None for c in row] for row in t.rows])
            r = w.close()
        except BaseException:
            w.abort()
            raise
        res.files.append(stage / name)
        res.warnings.extend(f"{name}: {x}" for x in r.warnings)
    if len(tables) > 1:
        res.warnings.append(f"CSV: создано файлов — {len(tables)} (по одному на таблицу).")
    if any(t.totals for t in tables):
        res.warnings.append("CSV: строки итогов не включаются (файл содержит только данные).")
    return res
