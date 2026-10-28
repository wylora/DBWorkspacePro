"""Рендерер DOCX (python-docx)."""
from __future__ import annotations

import io
import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

from .document import (DocBullets, DocChart, DocHeading, DocKeyValues, DocPageBreak, DocParagraph, DocTable, DocToc,
                       DocTitlePage, ReportDoc)
from .formatting import clean_xml_text
from .styles import StylePreset
from .variables import render_text

ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT, "center": WD_ALIGN_PARAGRAPH.CENTER}
DOTX_TYPE = b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
DOCX_TYPE = b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
MIN_COL_CM = 1.3


_W = nsdecls("w")


def _esc(text: str) -> str:
    text = clean_xml_text(text).replace("\t", " ")
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _runs_xml(text: str, rpr: str) -> str:
    parts = _esc(text).replace("\r\n", "\n").replace("\r", "\n").split("\n")
    body = '<w:br/>'.join(f'<w:t xml:space="preserve">{p}</w:t>' for p in parts)
    return f"<w:r>{rpr}{body}</w:r>" if text else ""


def _fast_row(cells, widths_tw: list[int], aligns: list[str], style_id: str, *, zebra_fill: str = "", bold: bool = False,
              half_pt: int = 0, muted: str = "", top_border: tuple[str, int, str] | None = None) -> object:
    """Строка таблицы одним XML-фрагментом (в десятки раз быстрее, чем через python-docx)."""
    out = [f"<w:tr {_W}><w:trPr><w:cantSplit/></w:trPr>"]
    for ci, c in enumerate(cells):
        rpr_bits = ("<w:b/>" if bold else "") + ("<w:i/>" if c.empty else "") + \
                   (f'<w:color w:val="{muted.lstrip("#")}"/>' if c.empty else "") + \
                   (f'<w:sz w:val="{half_pt}"/>' if half_pt else "")
        rpr = f"<w:rPr>{rpr_bits}</w:rPr>" if rpr_bits else ""
        tcpr = f'<w:tcW w:w="{widths_tw[ci]}" w:type="dxa"/>'
        if top_border:
            val, sz, color = top_border
            tcpr += f'<w:tcBorders><w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color.lstrip("#")}"/></w:tcBorders>'
        if zebra_fill:
            tcpr += f'<w:shd w:val="clear" w:color="auto" w:fill="{zebra_fill.lstrip("#")}"/>'
        jc = {"right": "right", "center": "center"}.get(aligns[ci], "")
        ppr = f'<w:pPr><w:pStyle w:val="{style_id}"/>' + (f'<w:jc w:val="{jc}"/>' if jc else "") + "</w:pPr>"
        text = "—" if c.empty else c.text
        out.append(f"<w:tc><w:tcPr>{tcpr}</w:tcPr><w:p>{ppr}{_runs_xml(text, rpr)}</w:p></w:tc>")
    out.append("</w:tr>")
    return parse_xml("".join(out))


def _rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _set_font(font, name: str | None = None, size: float | None = None, bold=None, italic=None, color: str | None = None):
    if name:
        font.name = name
        rpr = font.element.get_or_add_rPr() if hasattr(font.element, "get_or_add_rPr") else font.element
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), name)
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    if color:
        font.color.rgb = _rgb(color)


def _shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill.lstrip("#"))
    tcpr.append(shd)


def _cell_border(cell, **edges) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge, (val, sz, color) in edges.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color.lstrip("#"))
        borders.append(el)


def _table_borders(table, style: StylePreset) -> None:
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    kind, c = style.table_style, style.border

    def edge(name, val, sz, color):
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color.lstrip("#"))
        borders.append(el)

    none = ("nil", 0, "FFFFFF")
    spec = {
        "grid": dict(top=("single", 4, c), left=("single", 4, c), bottom=("single", 4, c), right=("single", 4, c),
                     insideH=("single", 4, c), insideV=("single", 4, c)),
        "horizontal": dict(top=("single", 8, style.accent), left=none, bottom=("single", 8, style.accent), right=none,
                           insideH=("single", 4, c), insideV=none),
        "rules": dict(top=("single", 12, "#000000"), left=none, bottom=("single", 12, "#000000"), right=none,
                      insideH=none, insideV=none),
        "minimal": dict(top=("single", 4, c), left=none, bottom=("single", 4, c), right=none,
                        insideH=("single", 2, c), insideV=none),
    }[kind]
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge(name, *spec[name])
    tblpr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for side, w in (("top", 40), ("left", 80), ("bottom", 40), ("right", 80)):
        m = OxmlElement(f"w:{side}")
        m.set(qn("w:w"), str(w))
        m.set(qn("w:type"), "dxa")
        mar.append(m)
    tblpr.append(mar)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)


def _field(paragraph, instr: str, placeholder: str = "1", dirty: bool = False):
    def run_with(child):
        r = OxmlElement("w:r")
        r.append(child)
        paragraph._p.append(r)
        return r

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    if dirty:
        begin.set(qn("w:dirty"), "true")
    run_with(begin)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    run_with(it)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run_with(sep)
    t = OxmlElement("w:t")
    t.text = placeholder
    run_with(t)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run_with(end)


def _para_border(paragraph, edge: str, sz: int, color: str, space: int = 4) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pb = ppr.find(qn("w:pBdr"))
    if pb is None:
        pb = OxmlElement("w:pBdr")
        ppr.append(pb)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color.lstrip("#"))
    pb.append(el)


def _orient(section, landscape: bool) -> None:
    w, h = section.page_width, section.page_height
    if (landscape and w < h) or (not landscape and w > h):
        section.page_width, section.page_height = h, w
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT


def _open_base(path: str, warnings: list[str]):
    p = Path(path)
    if not p.is_file():
        warnings.append(f"Пользовательский шаблон Word не найден: {path}. Использован стандартный.")
        return None
    try:
        with zipfile.ZipFile(p) as zin:
            data = zin.read("[Content_Types].xml")
            if DOTX_TYPE in data:
                tmp = Path(tempfile.mkdtemp()) / "base.docx"
                with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
                    for item in zin.infolist():
                        content = zin.read(item.filename)
                        if item.filename == "[Content_Types].xml":
                            content = content.replace(DOTX_TYPE, DOCX_TYPE)
                        zout.writestr(item, content)
                return Document(str(tmp))
        return Document(str(p))
    except Exception as exc:  # noqa: BLE001
        warnings.append(f"Не удалось открыть пользовательский шаблон Word ({exc.__class__.__name__}). "
                        "Использован стандартный.")
        return None


def _replace_vars_in_paragraph(par, variables: dict, missing: list[str]) -> None:
    if "{{" not in par.text:
        return
    new, miss = render_text(par.text, variables)
    missing.extend(m for m in miss if m not in missing)
    if new != par.text and par.runs:
        par.runs[0].text = new
        for r in par.runs[1:]:
            r.text = ""


def _replace_vars(d, variables: dict, warnings: list[str]) -> None:
    missing: list[str] = []
    containers = [d]
    for sec in d.sections:
        containers += [sec.header, sec.footer, sec.first_page_header, sec.first_page_footer]

    def walk(c):
        for par in c.paragraphs:
            _replace_vars_in_paragraph(par, variables, missing)
        for tbl in getattr(c, "tables", []):
            for row in tbl.rows:
                for cell in row.cells:
                    walk(cell)

    for c in containers:
        try:
            walk(c)
        except Exception:  # noqa: BLE001
            continue
    for m in missing:
        warnings.append(f"В шаблоне Word неизвестная переменная {{{{{m}}}}} (оставлена как есть).")


class DocxRenderer:
    def __init__(self, doc: ReportDoc):
        self.doc = doc
        self.t = doc.template
        self.st = doc.style
        self.warnings: list[str] = []
        self.base = None
        if self.t.base_docx:
            self.base = _open_base(self.t.base_docx, self.warnings)
        self.d = self.base or Document()
        self.custom_base = self.base is not None
        self.font = self.t.font_name or self.st.font_name
        self.size = self.t.font_size or self.st.font_size
        self.accent = self.t.accent or self.st.accent
        self.landscape_now = self.t.orientation == "landscape"
        self.tbl_font = self.st.table_font_name if not self.t.font_name else self.t.font_name

    # ---------------------------------------------------------- подготовка
    def _setup_page(self) -> None:
        sec = self.d.sections[0]
        if not self.custom_base:
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
            m = Cm(self.t.margin_cm or self.st.margin_cm)
            sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = m
        _orient(sec, self.t.orientation == "landscape")

    def _setup_styles(self) -> None:
        styles = self.d.styles
        if not self.custom_base:
            normal = styles["Normal"]
            _set_font(normal.font, self.font, self.size, color="#000000" if self.st.id != "minimal" else "#222222")
            pf = normal.paragraph_format
            pf.space_after = Pt(6)
            pf.line_spacing = self.st.line_spacing
            for i, key in enumerate(("Heading 1", "Heading 2", "Heading 3")):
                h = styles[key]
                _set_font(h.font, self.st.heading_font_name if not self.t.font_name else self.t.font_name,
                          self.st.heading_sizes[i], bold=True, italic=False,
                          color=self.accent if self.st.id not in ("academic", "minimal") else self.st.heading_color)
                h.paragraph_format.space_before = Pt(16 if i == 0 else 12)
                h.paragraph_format.space_after = Pt(6)
                h.paragraph_format.keep_with_next = True
        for name, size in (("ReportTable", None), ("ReportNote", None), ("ReportCode", None)):
            if name not in [s.name for s in styles]:
                s = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
                s.base_style = styles["Normal"]
        ts = styles["ReportTable"]
        _set_font(ts.font, self.tbl_font, self.st.table_font_size if not self.t.font_size else self.t.font_size - 1)
        ts.paragraph_format.space_after = Pt(0)
        ts.paragraph_format.space_before = Pt(0)
        ts.paragraph_format.line_spacing = 1.0
        ns = styles["ReportNote"]
        _set_font(ns.font, None, max(8, self.size - 1.5), italic=True, color=self.st.muted)
        ns.paragraph_format.space_after = Pt(8)
        cs = styles["ReportCode"]
        _set_font(cs.font, "Consolas", 8.5)
        cs.paragraph_format.space_after = Pt(6)
        cs.paragraph_format.line_spacing = 1.0
        for lvl in (1, 2, 3):
            nm = f"TOC {lvl}"
            if nm not in [s.name for s in styles]:
                s = styles.add_style(nm, WD_STYLE_TYPE.PARAGRAPH)
                s.base_style = styles["Normal"]
                s.paragraph_format.left_indent = Cm(0.6 * (lvl - 1))
                s.paragraph_format.space_after = Pt(3)

    def _header_footer(self) -> None:
        sec = self.d.sections[0]
        title_page = any(isinstance(e, DocTitlePage) for e in self.doc.elements)
        sec.different_first_page_header_footer = title_page
        has_own = self.custom_base and any(p.text.strip() for p in sec.header.paragraphs + sec.footer.paragraphs)
        if has_own:
            return
        ru = self.doc.formatter.ru if self.doc.formatter else True
        head = self.doc.header
        if self.doc.part and self.doc.part[1] > 1:
            head = (head + "  ·  " if head else "") + (f"часть {self.doc.part[0] + 1} из {self.doc.part[1]}" if ru
                                                       else f"part {self.doc.part[0] + 1} of {self.doc.part[1]}")
        if head:
            hp = sec.header.paragraphs[0]
            hp.text = head
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _set_font(hp.runs[0].font, None, 8.5, color=self.st.muted)
            _para_border(hp, "bottom", 4, self.st.border)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if self.doc.footer:
            r = fp.add_run(self.doc.footer + ("    " if self.t.page_numbers else ""))
            _set_font(r.font, None, 8.5, color=self.st.muted)
        if self.t.page_numbers:
            for kind, text in (("txt", "Страница " if ru else "Page "), ("PAGE", "1"), ("txt", " из " if ru else " of "),
                               ("NUMPAGES", "1")):
                if kind == "txt":
                    r = fp.add_run(text)
                    _set_font(r.font, None, 8.5, color=self.st.muted)
                else:
                    _field(fp, kind, text)
            for r in fp._p.iterchildren(qn("w:r")):
                rpr = r.find(qn("w:rPr"))
                if rpr is None:
                    rpr = OxmlElement("w:rPr")
                    r.insert(0, rpr)
                    sz = OxmlElement("w:sz")
                    sz.set(qn("w:val"), "17")
                    rpr.append(sz)
                    col = OxmlElement("w:color")
                    col.set(qn("w:val"), self.st.muted.lstrip("#"))
                    rpr.append(col)

    # --------------------------------------------------------------- размеры
    def _avail_cm(self) -> float:
        sec = self.d.sections[-1]
        return (sec.page_width - sec.left_margin - sec.right_margin) / 360000.0

    # -------------------------------------------------------------- элементы
    def render(self, path) -> list[str]:
        d = self.d
        if self.custom_base:
            _replace_vars(d, self.doc.variables, self.warnings)
        self._setup_page()
        self._setup_styles()
        self._header_footer()
        elements = self.doc.elements
        self._landscape_open = False
        for i, e in enumerate(elements):
            if isinstance(e, DocHeading) and i + 1 < len(elements) and isinstance(elements[i + 1], DocTable) \
                    and self._is_wide(elements[i + 1]) and not self.landscape_now:
                self._enter_landscape()            # заголовок должен оказаться на одной странице со своей таблицей
            if isinstance(e, DocTitlePage):
                self._title_page(e)
            elif isinstance(e, DocToc):
                self._toc()
            elif isinstance(e, DocHeading):
                nxt = elements[i + 1] if i + 1 < len(elements) else None
                if self._landscape_open and not (isinstance(nxt, DocTable) and self._is_wide(nxt)):
                    self._leave_landscape()
                d.add_heading(e.text, level=min(3, max(1, e.level)))
            elif isinstance(e, DocParagraph):
                if self._landscape_open and e.style != "note":
                    self._leave_landscape()
                self._paragraph(e)
            elif isinstance(e, DocBullets):
                for it in e.items:
                    self._bullet(it)
            elif isinstance(e, DocChart):
                self._leave_landscape()
                self._chart(e)
            elif isinstance(e, DocKeyValues):
                self._keyvalues(e)
            elif isinstance(e, DocTable):
                self._table(e)
            elif isinstance(e, DocPageBreak):
                self._page_break()
        cp = d.core_properties
        cp.title = self.doc.title
        cp.author = self.doc.author or "DB Workspace Pro"
        cp.subject = self.t.name
        cp.comments = "Создано в DB Workspace Pro"
        cp.created = cp.modified = self.doc.generated_at
        d.save(str(path))
        return self.warnings

    def _page_break(self) -> None:
        last = self.d.paragraphs[-1] if self.d.paragraphs else None
        p = self.d.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_break(WD_BREAK.PAGE)

    def _title_page(self, tp: DocTitlePage) -> None:
        d, st = self.d, self.st
        al = ALIGN[st.title_align]
        if tp.organization:
            p = d.add_paragraph()
            p.alignment = al
            r = p.add_run(tp.organization)
            _set_font(r.font, None, self.size + 1, bold=True, color=self.st.muted)
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(150 if tp.organization else 170)
        p.alignment = al
        r = p.add_run(tp.title)
        _set_font(r.font, st.heading_font_name if not self.t.font_name else self.t.font_name, 28, bold=True,
                  color=self.accent if st.id not in ("academic", "minimal") else st.heading_color)
        _para_border(p, "bottom", 12, self.accent, 10)
        if tp.subtitle:
            p = d.add_paragraph()
            p.alignment = al
            p.paragraph_format.space_before = Pt(10)
            _set_font(p.add_run(tp.subtitle).font, None, self.size + 4, color=st.muted)
        if tp.part_label:
            p = d.add_paragraph()
            p.alignment = al
            _set_font(p.add_run(tp.part_label).font, None, self.size + 1, bold=True, color=self.accent)
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(120)
        p.alignment = al
        lines = []
        ru = self.doc.formatter.ru if self.doc.formatter else True
        if tp.author:
            lines.append(("Автор: " if ru else "Author: ") + tp.author)
        lines.append(("Дата формирования: " if ru else "Generated: ") + tp.generated)
        for i, line in enumerate(lines):
            if i:
                p.add_run().add_break()
            _set_font(p.add_run(line).font, None, self.size, color=st.muted)

    def _toc(self) -> None:
        d = self.d
        ru = self.doc.formatter.ru if self.doc.formatter else True
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        _set_font(p.add_run("Содержание" if ru else "Contents").font,
                  self.st.heading_font_name if not self.t.font_name else self.t.font_name,
                  self.st.heading_sizes[0], bold=True, color=self.accent if self.st.id not in ("academic", "minimal")
                  else self.st.heading_color)
        heads = self.doc.headings(3)
        if not heads:
            return
        paras = []
        for h in heads:
            hp = d.add_paragraph(style=f"TOC {min(3, h.level)}")
            hp.add_run(h.text)
            paras.append(hp)
        first, last = paras[0], paras[-1]
        # поле TOC охватывает все абзацы-заглушки: Word обновит его (с номерами страниц) при открытии/F9
        begin = OxmlElement("w:r")
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), "begin")
        fc.set(qn("w:dirty"), "true")
        begin.append(fc)
        instr = OxmlElement("w:r")
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = ' TOC \\o "1-3" \\h \\z \\u '
        instr.append(it)
        sep = OxmlElement("w:r")
        sc = OxmlElement("w:fldChar")
        sc.set(qn("w:fldCharType"), "separate")
        sep.append(sc)
        ppr = first._p.find(qn("w:pPr"))
        pos = list(first._p).index(ppr) + 1 if ppr is not None else 0
        for k, el in enumerate((begin, instr, sep)):
            first._p.insert(pos + k, el)
        end = OxmlElement("w:r")
        ec = OxmlElement("w:fldChar")
        ec.set(qn("w:fldCharType"), "end")
        end.append(ec)
        last._p.append(end)
        note = d.add_paragraph(style="ReportNote")
        note.add_run("Номера страниц появятся после обновления оглавления в Word (правая кнопка → «Обновить поле»)."
                     if ru else "Page numbers appear after updating the table of contents in Word.")
        self._page_break()

    def _paragraph(self, e: DocParagraph) -> None:
        if e.style == "code":
            p = self.d.add_paragraph(style="ReportCode")
            for i, line in enumerate(e.text.split("\n")):
                if i:
                    p.add_run().add_break()
                p.add_run(line)
            _para_border(p, "left", 12, self.st.border, 8)
        elif e.style in ("note", "caption"):
            self.d.add_paragraph(e.text, style="ReportNote")
        else:
            self.d.add_paragraph(e.text)

    def _bullet(self, text: str) -> None:
        try:
            self.d.add_paragraph(text, style="List Bullet")
        except KeyError:
            p = self.d.add_paragraph("•  " + text)
            p.paragraph_format.left_indent = Cm(0.6)

    def _chart(self, e: DocChart) -> None:
        width = min(self._avail_cm(), 16.5)
        self.d.add_picture(io.BytesIO(e.png), width=Cm(width))
        pic_par = self.d.paragraphs[-1]
        pic_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic_par.paragraph_format.keep_together = True
        try:
            for dp in pic_par._p.iter(qn("wp:docPr")):
                dp.set("descr", e.alt or e.title)
                dp.set("title", e.title)
        except Exception:  # noqa: BLE001
            pass

    def _keyvalues(self, e: DocKeyValues) -> None:
        d = self.d
        if e.title:
            d.add_heading(e.title, 2)
        table = d.add_table(rows=0, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        _table_borders(table, self.st if self.st.table_style != "rules" else self.st)
        avail = self._avail_cm()
        for k, v in e.items:
            row = table.add_row()
            self._cell(row.cells[0], k, "left", bold=True, color=self.st.muted)
            self._cell(row.cells[1], v, "left")
            row.cells[0].width, row.cells[1].width = Cm(avail * 0.32), Cm(avail * 0.68)
            self._row_props(row, cant_split=True)
        d.add_paragraph().paragraph_format.space_after = Pt(4)

    def _is_wide(self, t: DocTable) -> bool:
        return len(t.columns) > 8 or sum(c.width for c in t.columns) > (105 if not self.landscape_now else 190)

    def _enter_landscape(self) -> None:
        if self._landscape_open or self.landscape_now:
            return
        sec = self.d.add_section(WD_SECTION.NEW_PAGE)
        sec.different_first_page_header_footer = False
        _orient(sec, True)
        self._landscape_open = True

    def _leave_landscape(self) -> None:
        if not self._landscape_open:
            return
        sec = self.d.add_section(WD_SECTION.NEW_PAGE)
        sec.different_first_page_header_footer = False
        _orient(sec, False)
        self._landscape_open = False

    # ---------------------------------------------------------------- таблицы
    def _cell(self, cell, text: str, align: str, *, bold=False, italic=False, color=None, size=None) -> None:
        p = cell.paragraphs[0]
        p.style = self.d.styles["ReportTable"]
        p.alignment = ALIGN.get(align, WD_ALIGN_PARAGRAPH.LEFT)
        if text:
            r = p.add_run(text)
            if bold or italic or color or size:
                _set_font(r.font, None, size, bold=bold or None, italic=italic or None, color=color)

    @staticmethod
    def _row_props(row, header: bool = False, cant_split: bool = True) -> None:
        trpr = row._tr.get_or_add_trPr()
        if cant_split:
            trpr.append(OxmlElement("w:cantSplit"))
        if header:
            trpr.append(OxmlElement("w:tblHeader"))

    def _table(self, t: DocTable) -> None:
        d, st = self.d, self.st
        n = len(t.columns)
        if n == 0:
            return
        wide = self._is_wide(t)
        if wide and not self.landscape_now:
            self._enter_landscape()
        elif not wide and self._landscape_open:
            self._leave_landscape()
        if t.title and t.show_title:
            cap = d.add_paragraph()
            cap.paragraph_format.keep_with_next = True
            cap.paragraph_format.space_before = Pt(8)
            cap.paragraph_format.space_after = Pt(3)
            _set_font(cap.add_run(t.title).font, None, self.size, bold=True, color=self.st.muted
                      if self.st.id in ("minimal",) else None)
        avail = self._avail_cm()
        weights = [max(c.width, 4) for c in t.columns]
        total_w = sum(weights)
        widths = [max(MIN_COL_CM, avail * w / total_w) for w in weights]
        scale = avail / sum(widths)
        widths = [w * scale for w in widths]
        fsize = (st.table_font_size if not self.t.font_size else self.t.font_size - 1) - (1 if n > 8 else 0) - (1 if n > 12 else 0)
        fsize = max(7.0, fsize)

        table = d.add_table(rows=1, cols=n)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _table_borders(table, st)
        for i, w in enumerate(widths):
            table.columns[i].width = Cm(w)
        hdr = table.rows[0]
        self._row_props(hdr, header=True)
        for i, c in enumerate(t.columns):
            cell = hdr.cells[i]
            cell.width = Cm(widths[i])
            self._cell(cell, c.title, "center" if c.align == "right" else "left", bold=True, color=st.header_fg,
                       size=fsize)
            if st.header_bg.upper() != "#FFFFFF":
                _shade(cell, st.header_bg)
            if st.table_style == "rules":
                _cell_border(cell, bottom=("single", 6, "#000000"))
        tw = [int(w / 2.54 * 1440) for w in widths]
        aligns = [c.align for c in t.columns]
        style_id = d.styles["ReportTable"].style_id
        half_pt = int(round(fsize * 2))
        base_pt = int(round((st.table_font_size if not self.t.font_size else self.t.font_size - 1) * 2))
        tbl = table._tbl
        for ri, row in enumerate(t.rows):
            fill = st.zebra if (st.zebra and ri % 2 == 1) else ""
            tbl.append(_fast_row(row, tw, [c.align for c in row], style_id, zebra_fill=fill,
                                 half_pt=half_pt if half_pt != base_pt or self.t.font_size else 0, muted=st.muted))
        if t.totals:
            rule = ("double" if st.totals_double_rule else "single", 6 if st.totals_double_rule else 8,
                    "#000000" if st.totals_double_rule or st.table_style == "rules" else st.accent)
            fill = ("#E7EFE9" if st.id == "financial" else st.zebra) if st.zebra else ""
            tbl.append(_fast_row(t.totals, tw, [c.align for c in t.totals], style_id, zebra_fill=fill, bold=True,
                                 half_pt=half_pt, top_border=rule, muted=st.muted))
        # подпись/примечание под таблицей
        if t.note:
            d.add_paragraph(t.note, style="ReportNote").paragraph_format.space_before = Pt(4)
        else:
            sp = d.add_paragraph()
            sp.paragraph_format.space_after = Pt(4)


def render_docx(doc: ReportDoc, path) -> list[str]:
    """Сохраняет документ в path; возвращает предупреждения рендерера."""
    return DocxRenderer(doc).render(path)
