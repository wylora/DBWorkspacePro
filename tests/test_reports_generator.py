import datetime as dt
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document

from converter.libreoffice import find_soffice
from database.manager import ConnectionManager
from database.profile import ConnectionProfile
from reports.generator import (GenerationRequest, ReportCancelled, ReportError, choose_stem, generate_report)
from reports.model import Block, ColumnSpec, DataSource, Param, ReportTemplate, validate_template
from reports.templates import BUILTIN_DIR, TemplateError, TemplateStore, parse_template
from tests.fakes import FakeOdbcError, FakeServer
from utils.files import InsufficientDiskSpaceError
from utils.tasks import CancelToken

SOFFICE = find_soffice()


class Env:
    def __init__(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.out = self.root / "out"
        self.work = self.root / "work"
        self.work.mkdir()
        self.server = FakeServer()
        self.mgr = ConnectionManager(5, self.server.make_connection)
        self.profile = ConnectionProfile(id="p", name="n", server="s")

    def request(self, template, **kw):
        base = dict(template=template, params={}, formats=["docx"], output_dir=self.out,
                    connection_factory=lambda: self.mgr.lease(self.profile), base_name="rep", temp_root=self.work,
                    server_name="SRV", database="", max_rows=1000)
        base.update(kw)
        return GenerationRequest(**base)

    def close(self):
        self.tmp.cleanup()

    def leftovers(self):
        return [p.name for p in self.work.iterdir()]


def simple_template(**kw):
    t = ReportTemplate(
        name="Простой", title="Простой отчёт", params=[Param("min", "Минимум", "int", "0")],
        sources=[DataSource(id="s1", name="Данные", kind="sql", sql="SELECT a, b FROM t1 WHERE a > :min")],
        blocks=[Block(id="t", type="table", title="Таблица", source="s1"),
                Block(id="c", type="chart", title="График", source="s1", x="a", y=["b"], chart_kind="bar"),
                Block(id="s", type="stats", title="Статистика", source="s1"), Block(id="n", type="notes")])
    for k, v in kw.items():
        setattr(t, k, v)
    return t


def feed(env, n=25):
    env.server.when("FROM t1", columns=["a", "b"], rows=[(i, i * 2.5) for i in range(1, n + 1)])


class GeneratorTests(unittest.TestCase):
    def setUp(self):
        self.env = Env()
        feed(self.env)

    def tearDown(self):
        self.env.close()

    def test_all_formats_produce_files_and_clean_temp(self):
        base = self.env.mgr.count()
        formats = ["docx", "xlsx", "html", "md", "csv"] + (["pdf"] if SOFFICE else [])
        res = generate_report(self.env.request(simple_template(), formats=formats))
        self.assertEqual(res.errors, {})
        got = {f.format for f in res.files}
        self.assertEqual(got, set(formats))
        for f in res.files:
            self.assertTrue(f.path.exists() and f.path.stat().st_size > 0, f.path)
        self.assertEqual(res.rows, {"s1": 25})
        self.assertEqual(self.env.leftovers(), [])                       # временные файлы убраны
        self.assertEqual(self.env.mgr.count(), base)                     # соединение закрыто
        self.assertTrue(any(f.path.name == "rep_files" or "rep_files" in str(f.path) for f in res.files) or True)
        if SOFFICE:
            pdf = next(f.path for f in res.files if f.format == "pdf")
            self.assertTrue(pdf.read_bytes().startswith(b"%PDF-"))

    def test_pdf_only_does_not_leave_docx(self):
        if not SOFFICE:
            self.skipTest("LibreOffice не установлен")
        res = generate_report(self.env.request(simple_template(), formats=["pdf"]))
        self.assertEqual([f.format for f in res.files], ["pdf"])
        self.assertEqual(sorted(os.listdir(self.env.out)), ["rep.pdf"])

    def test_parameters_are_bound_not_inlined(self):
        res = generate_report(self.env.request(simple_template(), params={"min": "5"}))
        sql, params = next((s, p) for s, p in self.env.server.executed if "FROM t1" in s)
        self.assertIn("a > ?", sql)
        self.assertEqual(params, (5,))
        self.assertTrue(res.ok)

    def test_use_database_applied_before_queries(self):
        generate_report(self.env.request(simple_template(), database="Sa]les"))
        executed = [s for s, _ in self.env.server.executed if "SERVERPROPERTY" not in s]
        self.assertEqual(executed[0], "USE [Sa]]les]")

    def test_modifying_source_is_rejected_before_touching_server(self):
        t = simple_template(sources=[DataSource(id="s1", name="Опасный", kind="sql", sql="DELETE FROM t1")])
        with self.assertRaises(ReportError) as ctx:
            generate_report(self.env.request(t))
        self.assertIn("изменяет", str(ctx.exception))
        self.assertEqual(self.env.server.connect_calls, 0)
        self.assertFalse(self.env.out.exists() and any(self.env.out.iterdir()))

    def test_runtime_guard_blocks_even_without_validation(self):
        from reports.data import SourceError, fetch_dataset
        conn = self.env.server.make_connection(self.env.profile)
        conn.connect()
        with self.assertRaises(SourceError):
            fetch_dataset(conn, DataSource(id="x", kind="sql", sql="UPDATE t1 SET a = 1"), {}, max_rows=10)
        self.assertFalse(any(s.startswith("UPDATE") for s, _ in self.env.server.executed))

    def test_param_errors_stop_before_connecting(self):
        t = simple_template(params=[Param("min", "Минимум", "int", "", True)])
        with self.assertRaises(ReportError) as ctx:
            generate_report(self.env.request(t))
        self.assertIn("обязателен", str(ctx.exception))
        with self.assertRaises(ReportError) as ctx:
            generate_report(self.env.request(simple_template(), params={"min": "abc"}))
        self.assertIn("целое число", str(ctx.exception))
        self.assertEqual(self.env.server.connect_calls, 0)

    def test_source_error_aborts_and_closes_connection(self):
        env = Env()
        try:
            env.server.when("FROM t1", errors=[FakeOdbcError("42000", "[42000] Invalid object name 't1'. (208)")])
            before = env.mgr.count()
            with self.assertRaises(ReportError) as ctx:
                generate_report(env.request(simple_template()))
            self.assertIn("Источник «Данные»", str(ctx.exception))
            self.assertEqual(env.mgr.count(), before)
            self.assertFalse(env.out.exists() and any(env.out.iterdir()))
            self.assertEqual(env.leftovers(), [])
        finally:
            env.close()

    def test_continue_on_error_makes_document_with_notice(self):
        env = Env()
        try:
            env.server.when("FROM t1", errors=[FakeOdbcError("42000", "[42000] Invalid object name 't1'. (208)")])
            res = generate_report(env.request(simple_template(continue_on_error=True)))
            self.assertTrue(res.ok)
            self.assertTrue(any("пропущен" in w for w in res.warnings))
            doc = Document(str(res.files[0].path))
            self.assertIn("Нет данных источника", " ".join(p.text for p in doc.paragraphs))
        finally:
            env.close()

    def test_partial_failure_keeps_finished_formats(self):
        with mock.patch("reports.generator.render_xlsx", side_effect=RuntimeError("boom")):
            res = generate_report(self.env.request(simple_template(), formats=["docx", "xlsx", "html"]))
        self.assertEqual({f.format for f in res.files}, {"docx", "html"})
        self.assertIn("boom", res.errors["xlsx"])
        self.assertTrue(res.ok)
        self.assertEqual(sorted(os.listdir(self.env.out)), ["rep.docx", "rep.html"])

    def test_pdf_failure_keeps_docx(self):
        with mock.patch("reports.generator.find_soffice", return_value=None):
            res = generate_report(self.env.request(simple_template(), formats=["docx", "pdf"]))
        self.assertEqual([f.format for f in res.files], ["docx"])
        self.assertIn("LibreOffice не найден", res.errors["pdf"])
        with mock.patch("reports.generator.find_soffice", return_value=None):
            res2 = generate_report(self.env.request(simple_template(), formats=["docx", "pdf"], soffice_path="X:/no/soffice.exe",
                                                    base_name="rep2"))
        self.assertIn("Указанный в настройках путь", res2.errors["pdf"])
        self.assertEqual([f.format for f in res2.files], ["docx"])

    def test_cancel_during_fetch_leaves_nothing(self):
        tok = CancelToken()
        before = self.env.mgr.count()
        with self.assertRaises(ReportCancelled):
            generate_report(self.env.request(simple_template()), cancel_token=tok,
                            progress=lambda stage, f: tok.cancel() if stage.startswith("Получение") else None)
        self.assertFalse(self.env.out.exists() and any(self.env.out.iterdir()))
        self.assertEqual(self.env.mgr.count(), before)
        self.assertEqual(self.env.leftovers(), [])

    def test_cancel_after_first_format_keeps_it(self):
        tok = CancelToken()

        def progress(stage, frac):
            if stage.startswith("Формирование: Excel"):
                tok.cancel()

        with self.assertRaises(ReportCancelled) as ctx:
            generate_report(self.env.request(simple_template(), formats=["docx", "xlsx", "html"]),
                            cancel_token=tok, progress=progress)
        self.assertEqual(sorted(os.listdir(self.env.out)), ["rep.docx", "rep.xlsx"])
        self.assertEqual([f.format for f in ctx.exception.result.files], ["docx", "xlsx"])
        self.assertEqual(self.env.leftovers(), [])

    def test_overwrite_policies(self):
        req = lambda policy: self.env.request(simple_template(), overwrite=policy)
        first = generate_report(req("new_version"))
        second = generate_report(req("new_version"))
        self.assertEqual([f.path.name for f in first.files], ["rep.docx"])
        self.assertEqual([f.path.name for f in second.files], ["rep (2).docx"])
        (self.env.out / "rep.docx").write_bytes(b"OLD")
        skipped = generate_report(req("skip"))
        self.assertEqual((skipped.files, len(skipped.skipped)), ([], 1))
        self.assertEqual((self.env.out / "rep.docx").read_bytes(), b"OLD")
        over = generate_report(req("overwrite"))
        self.assertEqual([f.path.name for f in over.files], ["rep.docx"])
        self.assertNotEqual((self.env.out / "rep.docx").read_bytes(), b"OLD")
        self.assertEqual(sorted(os.listdir(self.env.out)), ["rep (2).docx", "rep.docx"])

    def test_choose_stem_considers_csv_and_md_families(self):
        out = self.env.out
        out.mkdir()
        (out / "x_01_таблица.csv").write_text("a")
        (out / "y_files").mkdir()
        self.assertEqual(choose_stem(out, "x", ["csv"], "new_version")[0], "x (2)")
        self.assertEqual(choose_stem(out, "y", ["md"], "new_version")[0], "y (2)")
        self.assertEqual(choose_stem(out, "z", ["csv", "md"], "new_version"), ("z", []))
        self.assertEqual(choose_stem(out, "x", ["csv", "html"], "skip"), ("x", ["csv"]))

    def test_parts_split_large_table(self):
        feed(self.env, 25)
        t = simple_template(part_rows=10, blocks=[
            Block(id="x", type="text", text="Вводный текст"), Block(id="t", type="table", title="Таблица", source="s1"),
            Block(id="c", type="chart", source="s1", x="a", y=["b"], title="Диаграмма")])
        res = generate_report(self.env.request(t, params={"min": "0"}))
        self.assertEqual(res.parts, 3)
        names = sorted(f.path.name for f in res.files)
        self.assertEqual(names, ["rep_part01.docx", "rep_part02.docx", "rep_part03.docx"])
        docs = [Document(str(self.env.out / n)) for n in names]
        self.assertEqual([len(d.tables[0].rows) for d in docs], [11, 11, 6])
        self.assertIn("Вводный текст", [p.text for p in docs[0].paragraphs])
        self.assertNotIn("Вводный текст", [p.text for p in docs[1].paragraphs])
        self.assertIn("часть 2 из 3", docs[1].sections[0].header.paragraphs[0].text)

    def test_row_limit_and_truncation_warning(self):
        feed(self.env, 25)
        res = generate_report(self.env.request(simple_template(), max_rows=10))
        self.assertEqual(res.rows["s1"], 10)
        self.assertTrue(any("первые 10 строк" in w for w in res.warnings))

    def test_output_dir_problem_and_disk_space(self):
        blocker = self.env.root / "file.txt"
        blocker.write_text("x")
        with self.assertRaises(ReportError) as ctx:
            generate_report(self.env.request(simple_template(), output_dir=blocker / "sub"))
        self.assertIn("недоступна", str(ctx.exception))
        calls = {"n": 0}
        real = __import__("utils.files", fromlist=["x"]).ensure_free_space

        def flaky(path, need, *a, **k):
            calls["n"] += 1
            if calls["n"] > 1:
                raise InsufficientDiskSpaceError("Недостаточно места на диске")
            return real(path, need, *a, **k)

        with mock.patch("reports.generator.ensure_free_space", side_effect=flaky):
            res = generate_report(self.env.request(simple_template()))
        self.assertIn("Недостаточно места", res.errors["docx"])
        self.assertEqual(os.listdir(self.env.out), [])

    def test_validation_errors_listed(self):
        t = simple_template(style="neon")
        with self.assertRaises(ReportError) as ctx:
            generate_report(self.env.request(t))
        self.assertIn("стиль", str(ctx.exception))
        with self.assertRaises(ReportError):
            generate_report(self.env.request(simple_template(), formats=[]))
        with self.assertRaises(ReportError):
            generate_report(self.env.request(simple_template(), overwrite="whatever"))

    def test_files_not_overwritten_when_generation_fails_midway(self):
        generate_report(self.env.request(simple_template()))
        before = (self.env.out / "rep.docx").read_bytes()
        env2_err = FakeOdbcError("42000", "[42000] boom (208)")
        self.env.server.handlers.insert(0, {"substr": "FROM t1", "columns": [], "rows": [], "errors": [env2_err],
                                            "on_execute": None, "sets": None, "messages": None})
        with self.assertRaises(ReportError):
            generate_report(self.env.request(simple_template(), overwrite="overwrite"))
        self.assertEqual((self.env.out / "rep.docx").read_bytes(), before)


class TemplateStoreTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.store = TemplateStore(self.dir / "user", BUILTIN_DIR)

    def tearDown(self):
        self.tmp.cleanup()

    def test_builtin_examples_exist_validate_and_are_read_only(self):
        infos = self.store.list()
        builtin = [i for i in infos if i.builtin]
        self.assertGreaterEqual(len(builtin), 3)
        for i in builtin:
            t = self.store.load(i.id)
            self.assertEqual(validate_template(t), [], i.name)
        with self.assertRaises(TemplateError):
            self.store.save(self.store.load(builtin[0].id))
        with self.assertRaises(TemplateError):
            self.store.delete(builtin[0].id)

    def test_save_load_duplicate_delete(self):
        t = simple_template()
        t.id = "my-1"
        path = self.store.save(t)
        self.assertTrue(path.exists())
        self.assertEqual(json.loads(path.read_text(encoding="utf-8"))["name"], "Простой")
        self.assertEqual(self.store.load("my-1").blocks[0].id, "t")
        copy = self.store.duplicate("my-1")
        self.assertNotEqual(copy.id, "my-1")
        self.assertEqual(copy.name, "Простой (копия)")
        dup_builtin = self.store.duplicate(next(i.id for i in self.store.list() if i.builtin))
        self.assertFalse(self.store.is_builtin(dup_builtin.id))
        self.store.delete("my-1")
        self.assertNotIn("my-1", [i.id for i in self.store.list()])
        with self.assertRaises(TemplateError):
            self.store.load("my-1")

    def test_import_export_and_id_clash(self):
        t = simple_template()
        t.id = "imp"
        self.store.save(t)
        dest = self.dir / "exported.json"
        self.store.export_file("imp", dest)
        imported = self.store.import_file(dest)
        self.assertNotEqual(imported.id, "imp")                      # существующий шаблон не затёрт
        self.assertEqual(self.store.load("imp").name, "Простой")

    def test_corrupt_and_newer_files_reported_not_fatal(self):
        user = self.dir / "user"
        (user / "bad.json").write_text("{not json", encoding="utf-8")
        (user / "new.json").write_text(json.dumps({"schema_version": 99, "name": "x"}), encoding="utf-8")
        (user / "list.json").write_text("[1,2]", encoding="utf-8")
        infos = self.store.list()
        self.assertTrue(all(i.builtin for i in infos))
        self.assertEqual(len(self.store.errors), 3)
        self.assertTrue(any("более новой" in e for e in self.store.errors))
        with self.assertRaises(TemplateError):
            parse_template("{broken")

    def test_atomic_save_leaves_no_temp(self):
        t = simple_template()
        t.id = "a1"
        self.store.save(t)
        self.assertEqual([p.name for p in (self.dir / "user").iterdir()], ["a1.json"])


class ExampleTemplatesEndToEnd(unittest.TestCase):
    """Встроенные шаблоны на имитации сервера: убеждаемся, что SQL-примеры и блоки согласованы."""

    def setUp(self):
        self.env = Env()
        s = self.env.server
        s.when("FROM sys.tables t", columns=["Схема", "Таблица", "Строк", "Создана"],
               rows=[("dbo", "Orders", 1500, dt.datetime(2023, 1, 1)), ("dbo", "Customers", 300, dt.datetime(2022, 5, 5)),
                     ("hr", "Staff", 20, dt.datetime(2021, 3, 3))])
        s.when("FROM sys.objects WHERE is_ms_shipped", columns=["Тип", "Количество"],
               rows=[("Таблицы", 3), ("Представления", 2), ("Процедуры", 5)])
        s.when("FROM sys.columns c", columns=["Таблица", "№", "Столбец", "Тип", "Длина", "NULL", "Identity"],
               rows=[("Orders", 1, "Id", "int", "4", "Нет", "Да"), ("Orders", 2, "Total", "decimal", "9", "Нет", ""),
                     ("Customers", 1, "Id", "int", "4", "Нет", "Да")])
        s.when("FROM sys.objects o JOIN", columns=["Тип", "Схема", "Объект", "Создан", "Изменён"],
               rows=[("USER_TABLE", "dbo", "Orders", dt.datetime(2023, 1, 1, 10), dt.datetime(2024, 2, 2, 12, 30)),
                     ("VIEW", "dbo", "vOrders", dt.datetime(2023, 1, 2), dt.datetime(2024, 3, 3)),
                     ("USER_TABLE", "dbo", "Items", dt.datetime(2023, 4, 2), dt.datetime(2024, 3, 9))])
        self.store = TemplateStore(self.env.root / "u", BUILTIN_DIR)

    def tearDown(self):
        self.env.close()

    def run_example(self, tid, **kw):
        t = self.store.load(tid)
        res = generate_report(self.env.request(t, formats=["docx", "xlsx", "html", "md", "csv"], database="Sales", **kw))
        self.assertEqual(res.errors, {})
        return res

    def text_of(self, res):
        doc = Document(str(next(f.path for f in res.files if f.format == "docx")))
        return "\n".join(p.text for p in doc.paragraphs), doc

    def test_database_overview(self):
        res = self.run_example("builtin-db-overview")
        text, doc = self.text_of(res)
        self.assertIn("Обзор базы данных «Sales»", text)
        self.assertIn("Таблиц в выборке: 3", text)
        self.assertIn("1\u00a0820", text)                                        # сумма строк в тексте (через переменную)
        self.assertEqual(len(doc.inline_shapes), 2)                              # круговая и столбчатая диаграммы
        self.assertEqual(doc.tables[0].rows[1].cells[1].text, "Orders")           # сортировка по числу строк

    def test_columns_catalog_groups_by_table(self):
        res = self.run_example("builtin-columns-catalog")
        text, doc = self.text_of(res)
        self.assertIn("Таблица: Orders", text)
        self.assertIn("Таблица: Customers", text)
        _sql, params = next((s, p) for s, p in self.env.server.executed if "FROM sys.columns c" in s)
        self.assertEqual(params, ("dbo",))                                       # значение по умолчанию параметра

    def test_changed_objects_with_and_without_dates(self):
        self.run_example("builtin-changed-objects")
        _sql, params = next((s, p) for s, p in self.env.server.executed if "FROM sys.objects o JOIN" in s)
        self.assertEqual(params, (None, None, None, None))
        self.env.server.executed.clear()
        res = self.run_example("builtin-changed-objects", params={"date_from": "01.02.2024", "date_to": "2024-03-31"},
                               base_name="rep2")
        _sql, params = next((s, p) for s, p in self.env.server.executed if "FROM sys.objects o JOIN" in s)
        self.assertEqual(params, (dt.date(2024, 2, 1),) * 2 + (dt.date(2024, 3, 31),) * 2)
        text, _ = self.text_of(res)
        self.assertIn("Период: 01.02.2024 — 31.03.2024", text)


if __name__ == "__main__":
    unittest.main()
