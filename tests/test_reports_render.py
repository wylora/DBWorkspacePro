import datetime as dt
import glob
import os
import shutil
import subprocess
import tempfile
import threading
import time
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from docx import Document
from docx.oxml.ns import qn

from converter.libreoffice import (ConversionCancelled, ConversionError, convert, find_soffice, verify_output)
from reports.data import Dataset
from reports.document import (BuildContext, DocChart, DocHeading, DocTable, DocToc, build_document)
from reports.model import (Aggregate, Block, ColumnSpec, ComputedColumn, DataSource, Metric, Param, ReportTemplate,
                           SortSpec)
from reports.render_docx import render_docx
from reports.render_formats import render_csv, render_html, render_md, render_xlsx
from utils.tasks import CancelToken

SOFFICE = find_soffice()
HAVE_PDFINFO = shutil.which("pdfinfo") is not None and shutil.which("pdftotext") is not None


def sales_ds(n=30):
    rows = [(i, f"Клиент {i}", "Север" if i % 3 else "Юг", i * 10.5, dt.date(2024, 1, 1) + dt.timedelta(days=i * 5),
             None if i % 7 == 0 else i) for i in range(1, n + 1)]
    return Dataset("s1", "Продажи", ["Id", "Клиент", "Регион", "Сумма", "Дата", "Кол-во"], rows)


def make_template(**kw):
    t = ReportTemplate(
        name="Тест", title="Отчёт за {{param.year}}", subtitle="Подзаголовок", organization="ООО «Ромашка»",
        author="Иванов", params=[Param("year", "Год", "int", "2024")],
        sources=[DataSource(id="s1", name="Продажи", kind="sql", sql="SELECT 1")],
        blocks=[Block(id="h", type="heading", title="Введение", level=1),
                Block(id="x", type="text", text="Создан {{date}}. Строк: {{rows.s1}}. Сумма: {{stat.s1.Сумма.sum}}."),
                Block(id="t", type="table", title="Продажи", source="s1", max_rows=20,
                      columns=[ColumnSpec("Id", total="count"), ColumnSpec("Клиент"), ColumnSpec("Регион"),
                               ColumnSpec("Сумма", total="sum", format="currency"), ColumnSpec("Дата"),
                               ColumnSpec("Кол-во")]),
                Block(id="st", type="stats", title="Статистика", source="s1"),
                Block(id="c", type="chart", title="По регионам", source="s1", x="Регион", y=["Сумма"]),
                Block(id="sm", type="summary", source="s1"), Block(id="n", type="notes")])
    for k, v in kw.items():
        setattr(t, k, v)
    return t


def build(t=None, ds=None, **ctxkw):
    t = t or make_template()
    ds = ds or sales_ds()
    kw = dict(param_values={"year": 2024}, raw_params={"year": "2024"}, server="SRV", database="Sales", max_rows=1000)
    kw.update(ctxkw)
    return build_document(BuildContext(t, {"s1": ds}, **kw))


class DocumentBuildTests(unittest.TestCase):
    def test_variables_title_and_structure(self):
        d = build()
        self.assertEqual(d.title, "Отчёт за 2024")             # целый параметр без разделителя тысяч
        kinds = [type(e).__name__ for e in d.elements]
        self.assertEqual(kinds[:3], ["DocTitlePage", "DocPageBreak", "DocToc"])
        text = " ".join(e.text for e in d.elements if type(e).__name__ == "DocParagraph")
        self.assertIn("Строк: 30", text)
        self.assertIn("Сумма: 4", text)                         # 10,5*(1+…+30) = 4 882,50
        self.assertNotIn("{{", text)
        self.assertEqual(d.warnings[0].startswith("Таблица «Продажи»: показана часть строк (20 из 30)"), True)

    def test_unknown_variable_is_kept_and_reported(self):
        t = make_template()
        t.blocks[1].text = "Привет {{нет_такой}}"
        d = build(t)
        self.assertTrue(any("нет_такой" in w for w in d.warnings))
        self.assertTrue(any("{{нет_такой}}" in getattr(e, "text", "") for e in d.elements))

    def test_table_columns_hidden_sort_totals_computed(self):
        t = make_template()
        blk = t.blocks[2]
        blk.columns = [ColumnSpec("Клиент"), ColumnSpec("Сумма", total="sum", decimals=1),
                       ColumnSpec("Регион", hidden=True), ColumnSpec("Нет такого")]
        blk.sort = [SortSpec("Сумма", True)]
        blk.max_rows = 5
        blk.row_numbers = True
        blk.computed = [ComputedColumn("Сумма / 2", "div", "Сумма", "2", 2), ComputedColumn("Деление на 0", "div", "Сумма", "0")]
        d = build(t)
        table = d.tables()[0]
        self.assertEqual([c.title for c in table.columns], ["№", "Клиент", "Сумма", "Сумма / 2", "Деление на 0"])
        self.assertEqual(table.rows[0][1].text, "Клиент 30")                       # сортировка по убыванию
        self.assertEqual(len(table.rows), 5)
        self.assertEqual(table.rows[0][3].raw, 157.5)                              # 315 / 2
        self.assertTrue(table.rows[0][4].empty)                                    # деление на ноль -> пусто
        self.assertEqual(table.totals[2].raw, 315.0 + 304.5 + 294.0 + 283.5 + 273.0)  # итоги по показанным строкам
        self.assertIn("Итоги рассчитаны по показанным строкам", table.note)
        self.assertTrue(any("Нет такого" in w for w in d.warnings))

    def test_empty_cells_marked_and_alignment(self):
        table = build().tables()[0]
        qty = next(i for i, c in enumerate(table.columns) if c.title == "Кол-во")
        self.assertTrue(table.rows[6][qty].empty)                                  # i == 7 -> None
        self.assertEqual(table.columns[0].align, "right")
        self.assertEqual(table.columns[1].align, "left")

    def test_no_stats_for_unsuitable_types(self):
        stats_tables = [t for t in build().tables() if t.compact]
        titles = " ".join(t.title for t in stats_tables)
        self.assertIn("Числовые поля", titles)
        numeric = next(t for t in stats_tables if t.title == "Числовые поля")
        id_row = next(r for r in numeric.rows if r[0].text == "Id")
        self.assertEqual(id_row[3].text, "—")                                      # у идентификатора нет суммы
        text_table = next(t for t in stats_tables if t.title.startswith("Текстовые"))
        self.assertFalse(any(r[0].text == "Сумма" for r in text_table.rows))

    def test_group_sections_and_limit(self):
        t = make_template()
        t.blocks[2].group_by = "Регион"
        t.blocks[2].columns = []
        d = build(t, max_rows=1000)
        t.blocks[2].max_rows = 0
        d_all = build(t, max_rows=1000)
        heads = [e.text for e in d_all.elements if isinstance(e, DocHeading)]
        self.assertIn("Регион: Север", heads)
        self.assertIn("Регион: Юг", heads)
        groups = [e for e in d_all.elements if isinstance(e, DocTable) and not e.compact]
        self.assertEqual(sorted(len(g.rows) for g in groups), [10, 20])
        self.assertNotIn("Регион", [c.title for c in groups[0].columns])          # столбец группировки не дублируется

    def test_group_limit_is_reported_in_document(self):
        t = make_template()
        t.blocks[2].group_by = "Регион"
        t.blocks[2].columns = []
        t.blocks[2].max_rows = 20                                                   # Север исчерпает лимит целиком
        d = build(t)
        notes = [e.text for e in d.elements if type(e).__name__ == "DocParagraph" and e.style == "note"]
        self.assertTrue(any("не показано групп — 1 из 2" in n for n in notes))
        self.assertTrue(any("не показано групп" in w for w in d.warnings))
        self.assertEqual(sum(len(g.rows) for g in d.tables() if not g.compact), 20)

    def test_chart_skipped_with_reason_for_unsuitable_data(self):
        t = make_template()
        t.blocks[4].chart_kind = "line"                    # ось X — текст
        d = build(t)
        self.assertFalse(any(isinstance(e, DocChart) for e in d.elements))
        self.assertTrue(any("не построена" in w for w in d.warnings))

    def test_metrics_block(self):
        t = make_template()
        t.blocks[3].metrics = [Metric("Всего", "sum", "Сумма"), Metric("Клиентов", "distinct", "Клиент"),
                               Metric("Доля", "ratio", "Сумма", "Кол-во")]
        d = build(t)
        m = next(x for x in d.tables() if x.title == "Показатели")
        self.assertEqual([r[0].text for r in m.rows], ["Всего", "Клиентов", "Доля"])
        self.assertEqual(m.rows[1][1].text, "30")

    def test_parts(self):
        t = make_template(part_rows=8)
        t.blocks[2].max_rows = 0
        d0 = build(t, part=(0, 3))
        d1 = build(t, part=(1, 3))
        self.assertEqual(len(d0.tables()[0].rows), 8)
        self.assertEqual(len(d1.tables()[0].rows), 8)
        self.assertIn("часть 2 из 3", d1.tables()[0].note)
        self.assertTrue(any(isinstance(e, DocChart) for e in d0.elements))
        self.assertFalse(any(isinstance(e, DocChart) for e in d1.elements))       # мелкие блоки — только в первой части

    def test_empty_dataset(self):
        ds = Dataset("s1", "Продажи", ["Id", "Клиент", "Регион", "Сумма", "Дата", "Кол-во"], [])
        d = build(ds=ds)
        self.assertEqual(len(d.tables()[0].rows), 0)
        self.assertIn("Нет данных", d.tables()[0].note)


class DocxRenderTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def render(self, t=None, ds=None, **kw):
        path = self.d / "r.docx"
        warns = render_docx(build(t, ds, **kw), path)
        return Document(str(path)), path, warns

    def test_structure(self):
        doc, path, warns = self.render()
        self.assertEqual(warns, [])
        xml = doc.element.xml
        self.assertIn('TOC \\o "1-3"', xml)                                          # поле оглавления
        heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        for h in ("Введение", "Продажи", "Статистика", "По регионам", "Выводы", "Примечания"):
            self.assertIn(h, heads)
        sec = doc.sections[0]
        self.assertTrue(sec.different_first_page_header_footer)
        self.assertEqual(sec.header.paragraphs[0].text, "Отчёт за 2024")
        foot = sec.footer._element.xml
        self.assertIn("PAGE", foot)
        self.assertIn("NUMPAGES", foot)
        self.assertEqual(doc.core_properties.author, "Иванов")
        self.assertEqual(doc.core_properties.title, "Отчёт за 2024")
        self.assertEqual(len(doc.inline_shapes), 1)                                   # диаграмма

    def test_table_header_repeat_cant_split_zebra_and_totals(self):
        doc, *_ = self.render()
        t = doc.tables[0]
        self.assertIn("w:tblHeader", t.rows[0]._tr.xml)
        self.assertTrue(all("w:cantSplit" in r._tr.xml for r in t.rows))
        self.assertIn("F2F6FA", t.rows[2]._tr.xml.upper())                            # чередование строк
        self.assertNotIn("F2F6FA", t.rows[1]._tr.xml.upper())
        self.assertEqual(t.rows[-1].cells[1].text, "Итого")
        self.assertEqual(t.rows[1].cells[3].text.replace("\u00a0", " "), "10,50")  # денежный формат без символа валюты
        self.assertIn("w:tblLayout", t._tbl.xml)

    def test_empty_values_shown_as_dash(self):
        doc, *_ = self.render()
        texts = [c.text for r in doc.tables[0].rows for c in r.cells]
        self.assertIn("—", texts)

    def test_wide_table_goes_landscape_and_heading_stays_with_it(self):
        cols = ["Id"] + [f"Показатель {i}" for i in range(1, 12)]
        ds = Dataset("s1", "Широкая", cols, [(i,) + tuple(round(i * j * 1.37, 2) for j in range(1, 12)) for i in range(1, 6)])
        t = make_template(toc=False)
        t.blocks = [Block(id="a", type="table", title="Малая", source="s1", columns=[ColumnSpec("Id")]),
                    Block(id="b", type="table", title="Широкая", source="s1"),
                    Block(id="c", type="table", title="После", source="s1", columns=[ColumnSpec("Id")])]
        doc, *_ = self.render(t, ds)
        orient = [s.page_width > s.page_height for s in doc.sections]
        self.assertEqual(orient, [False, True, False])
        body = doc.element.body
        order = []
        for el in body.iterchildren():
            if el.tag == qn("w:p"):
                txt = "".join(x.text or "" for x in el.iter(qn("w:t")))
                if txt in ("Малая", "Широкая", "После"):
                    order.append(txt)
            if el.tag == qn("w:p") and el.find(qn("w:pPr")) is not None and el.find(qn("w:pPr")).find(qn("w:sectPr")) is not None:
                order.append("|")
        self.assertEqual(order, ["Малая", "|", "Широкая", "|", "После"])

    def test_styles_change_appearance(self):
        fonts = {}
        for style in ("business", "academic", "technical", "financial"):
            doc, *_ = self.render(make_template(style=style))
            fonts[style] = doc.styles["Normal"].font.name
            if style == "academic":
                self.assertIn("w:val=\"nil\"", doc.tables[0]._tbl.tblPr.xml.replace("'", '"'))  # без сетки
            if style == "financial":
                self.assertIn("double", doc.tables[0].rows[-1]._tr.xml)
        self.assertEqual(fonts["academic"], "Times New Roman")
        self.assertEqual(len(set(fonts.values())), 4)

    def test_user_font_size_accent_and_margins_override(self):
        doc, *_ = self.render(make_template(font_name="Arial", font_size=9, accent="#AA0000", margin_cm=1.5))
        self.assertEqual(doc.styles["Normal"].font.name, "Arial")
        self.assertEqual(doc.styles["Normal"].font.size.pt, 9)
        self.assertAlmostEqual(doc.sections[0].left_margin.cm, 1.5, places=1)
        self.assertEqual(str(doc.styles["Heading 1"].font.color.rgb), "AA0000")

    def test_landscape_template_and_no_title_page_toc(self):
        doc, *_ = self.render(make_template(orientation="landscape", title_page=False, toc=False, page_numbers=False,
                                            header_text="", footer_text=""))
        self.assertGreater(doc.sections[0].page_width, doc.sections[0].page_height)
        self.assertNotIn("TOC \\o", doc.element.xml)
        self.assertNotIn("NUMPAGES", doc.sections[0].footer._element.xml)
        self.assertFalse(doc.sections[0].different_first_page_header_footer)

    def test_large_table_performance_and_correctness(self):
        n = 3000
        rows = [(i, f"Клиент {i}", "Север", i * 1.5, dt.date(2024, 1, 1), i) for i in range(n)]
        ds = Dataset("s1", "Большая", ["Id", "Клиент", "Регион", "Сумма", "Дата", "Кол-во"], rows)
        t = make_template(title_page=False, toc=False)
        t.blocks = [Block(id="t", type="table", title="Много", source="s1", max_rows=0)]
        started = time.time()
        doc, *_ = self.render(t, ds, max_rows=n)
        elapsed = time.time() - started
        self.assertEqual(len(doc.tables[0].rows), n + 1)
        self.assertLess(elapsed, 15, f"формирование 3000 строк заняло {elapsed:.1f} с")

    def test_hostile_cell_content_does_not_break_document(self):
        rows = [(1, "Плохой\x00символ\x0b\x1f", "<w:p>&amp;</w:p>", 1.0, dt.date(2024, 1, 1), 1),
                (2, "Строка 1\nСтрока 2\tс табом", "x" * 5000, 2.0, dt.date(2024, 1, 2), 2),
                (3, "\ud83d\ude00 эмодзи", "=1+1", 3.0, dt.date(2024, 1, 3), 3)]
        ds = Dataset("s1", "Продажи", ["Id", "Клиент", "Регион", "Сумма", "Дата", "Кол-во"], rows)
        t = make_template(title_page=False, toc=False)
        t.blocks = [Block(id="t", type="table", title="Т<&>\x00", source="s1")]
        doc, path, _w = self.render(t, ds)
        cells = [[c.text for c in r.cells] for r in doc.tables[0].rows]
        self.assertEqual(cells[1][1], "Плохойсимвол")                          # управляющие символы удалены
        self.assertEqual(cells[1][2], "<w:p>&amp;</w:p>")                     # разметка в данных — это просто текст
        self.assertEqual(cells[2][1], "Строка 1\nСтрока 2 с табом")
        self.assertEqual(len(cells[2][2]), 2001)                              # длинное значение усечено с «…»
        self.assertEqual(cells[3][2], "=1+1")
        self.assertEqual(Document(str(path)).paragraphs[0].text, "Т<&>")     # и заголовок очищен
        with zipfile.ZipFile(path) as z:
            from lxml import etree
            etree.fromstring(z.read("word/document.xml"))                       # XML корректен

    def test_user_word_template_docx_and_dotx(self):
        base = Document()
        base.sections[0].header.paragraphs[0].text = "{{report_title}} — конфиденциально ({{author}})"
        base.add_paragraph("Преамбула организации: {{organization}}")
        src = self.d / "base.docx"
        base.save(str(src))
        dotx = self.d / "base.dotx"
        with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dotx, "w") as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.replace(b"wordprocessingml.document.main+xml", b"wordprocessingml.template.main+xml")
                zout.writestr(item, data)
        for path in (src, dotx):
            doc, _p, warns = self.render(make_template(base_docx=str(path), title_page=False, toc=False))
            self.assertEqual(doc.sections[0].header.paragraphs[0].text, "Отчёт за 2024 — конфиденциально (Иванов)")
            self.assertIn("Преамбула организации: ООО «Ромашка»", [p.text for p in doc.paragraphs])
            self.assertTrue(doc.tables)
            self.assertEqual(warns, [])

    def test_missing_user_template_falls_back_with_warning(self):
        doc, _p, warns = self.render(make_template(base_docx=str(self.d / "nope.docx")))
        self.assertTrue(any("не найден" in w for w in warns))
        self.assertTrue(doc.tables)
        bad = self.d / "bad.docx"
        bad.write_bytes(b"not a zip")
        doc, _p, warns = self.render(make_template(base_docx=str(bad)))
        self.assertTrue(any("Не удалось открыть" in w for w in warns))


class OtherFormatsTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)
        rows = [(1, "<script>alert(1)</script>", "a|b", 1.5, dt.date(2024, 1, 1), "=HYPERLINK(\"http://x\")"),
                (2, "Обычный", "x", None, dt.date(2024, 2, 1), None), (3, "Третий", "x", 2.25, dt.date(2024, 3, 1), "-5")]
        self.ds = Dataset("s1", "Продажи", ["Id", "Клиент", "Регион", "Сумма", "Дата", "Заметка"], rows)
        t = make_template()
        t.blocks = [Block(id="h", type="heading", title="Раздел"),
                    Block(id="t", type="table", title="Продажи", source="s1",
                          columns=[ColumnSpec("Id", total="count"), ColumnSpec("Клиент"), ColumnSpec("Регион"),
                                   ColumnSpec("Сумма", total="sum", format="currency"), ColumnSpec("Дата"), ColumnSpec("Заметка")]),
                    Block(id="s", type="stats", source="s1"), Block(id="c", type="chart", title="График", source="s1",
                                                                       x="Клиент", y=["Сумма"]),
                    Block(id="n", type="notes", title="Примечания")]
        t.currency_symbol = "₽"
        self.doc = build(t, self.ds)

    def tearDown(self):
        self.tmp.cleanup()

    def test_html_escapes_everything(self):
        r = render_html(self.doc, self.d, "rep")
        html_text = r.files[0].read_text(encoding="utf-8")
        self.assertNotIn("<script", html_text)
        self.assertIn("&lt;script&gt;alert(1)&lt;/script&gt;", html_text)
        self.assertIn("data:image/png;base64,", html_text)
        self.assertIn("<thead>", html_text)
        self.assertIn("<tfoot>", html_text)
        self.assertIn('class="empty">—', html_text)
        self.assertIn('href="#h1"', html_text)

    def test_markdown_escapes_and_saves_images(self):
        r = render_md(self.doc, self.d, "rep")
        md = r.files[0].read_text(encoding="utf-8")
        self.assertIn("a\\|b", md)
        self.assertIn("&lt;script&gt;", md)
        self.assertNotIn("<script>", md)
        self.assertIn("![График](rep_files/chart_01.png)", md)
        self.assertTrue((self.d / "rep_files" / "chart_01.png").read_bytes().startswith(b"\x89PNG"))
        self.assertIn("---:", md)

    def test_xlsx_types_and_safety(self):
        import openpyxl
        r = render_xlsx(self.doc, self.d, "rep")
        wb = openpyxl.load_workbook(r.files[0])
        self.assertIn("Продажи", wb.sheetnames)
        ws = wb["Продажи"]
        self.assertEqual([c.value for c in ws[1]][:3], ["Id", "Клиент", "Регион"])
        self.assertEqual(ws["D2"].value, 1.5)
        self.assertEqual(ws["D2"].number_format, '#,##0.00 "₽"')
        self.assertIsInstance(ws["E2"].value, dt.datetime)
        self.assertEqual(ws["F2"].value, '=HYPERLINK("http://x")')
        self.assertEqual(ws["F2"].data_type, "s")                                    # формула не выполняется
        self.assertEqual(ws["F4"].value, "-5")
        self.assertIsNone(ws["D3"].value)                                            # пусто остаётся пустым
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertTrue(ws.auto_filter.ref)
        self.assertEqual(ws["D5"].value, 3.75)                                       # строка итогов — число
        self.assertTrue(all(len(n) <= 31 for n in wb.sheetnames))
        self.assertEqual(len(wb["Диаграммы"]._images), 1)
        self.assertIn("Статистика", wb.sheetnames)

    def test_csv_single_and_multiple_tables(self):
        r = render_csv(self.doc, self.d, "rep")
        self.assertEqual([f.name for f in r.files], ["rep.csv"])
        text = r.files[0].read_text(encoding="utf-8-sig")
        self.assertIn("'=HYPERLINK", text)                                           # защита от формул
        self.assertIn("\n3;Третий;x;2.25;2024-03-01;-5", text.replace("\r\n", "\n"))
        self.assertTrue(any("итог" in w for w in r.warnings))
        two = build(None, self.ds)
        two.elements.extend([e for e in two.elements if isinstance(e, DocTable) and not e.compact][:1])
        r2 = render_csv(two, self.d, "two")
        self.assertEqual(len(r2.files), 2)
        self.assertTrue(r2.files[0].name.startswith("two_01_"))
        empty = build(None, self.ds)
        empty.elements = [e for e in empty.elements if not isinstance(e, DocTable)]
        self.assertEqual(render_csv(empty, self.d, "none").files, [])


@unittest.skipUnless(SOFFICE, "LibreOffice не установлен")
class LibreOfficeTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)
        self.docx = self.d / "r.docx"
        render_docx(build(), self.docx)

    def tearDown(self):
        self.tmp.cleanup()

    def test_docx_to_pdf(self):
        pdf = convert(self.docx, "pdf", self.d, soffice=SOFFICE)
        self.assertEqual(pdf.name, "r.pdf")
        self.assertTrue(pdf.read_bytes().startswith(b"%PDF-"))
        if HAVE_PDFINFO:
            info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
            self.assertGreaterEqual(int(info.split("Pages:")[1].split()[0]), 3)
            text = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True).stdout
            for needle in ("Отчёт за 2024", "Введение", "Клиент 1", "Итого", "Страница 2 из"):
                self.assertIn(needle, text)

    def test_source_untouched_and_no_profile_leak(self):
        before = self.docx.read_bytes()
        profiles = set(glob.glob(os.path.join(tempfile.gettempdir(), "dbwp-lo-*")))
        convert(self.docx, "pdf", self.d, soffice=SOFFICE)
        self.assertEqual(self.docx.read_bytes(), before)
        self.assertEqual(set(glob.glob(os.path.join(tempfile.gettempdir(), "dbwp-lo-*"))), profiles)

    def test_cancel_kills_process_and_cleans_up(self):
        tok = CancelToken()
        profiles = set(glob.glob(os.path.join(tempfile.gettempdir(), "dbwp-lo-*")))
        threading.Timer(0.4, tok.cancel).start()
        with self.assertRaises(ConversionCancelled):
            convert(self.docx, "pdf", self.d, soffice=SOFFICE, cancel_token=tok)
        self.assertFalse((self.d / "r.pdf").exists())
        self.assertEqual(set(glob.glob(os.path.join(tempfile.gettempdir(), "dbwp-lo-*"))), profiles)

    def test_timeout(self):
        with self.assertRaises(ConversionError) as ctx:
            convert(self.docx, "pdf", self.d, soffice=SOFFICE, timeout=0.05)
        self.assertIn("тайм-аут", str(ctx.exception))
        self.assertFalse((self.d / "r.pdf").exists())


class LibreOfficeErrorsTests(unittest.TestCase):
    def test_missing_program_and_files(self):
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "a.docx"
            src.write_bytes(b"x")
            with mock.patch("converter.libreoffice.find_soffice", return_value=None):
                with self.assertRaises(ConversionError) as ctx:
                    convert(src, "pdf", d)
            self.assertIn("libreoffice.org", str(ctx.exception))
            with self.assertRaises(ConversionError):
                convert(src, "pdf", d, soffice=str(Path(d) / "no-such-soffice"))
            with self.assertRaises(ConversionError):
                convert(Path(d) / "missing.docx", "pdf", d, soffice="/bin/true")
            from converter.libreoffice import locate_soffice
            wrong = locate_soffice(str(Path(d) / "nope"))
            self.assertEqual(wrong.path, find_soffice(""))                # неверный путь — автоматический поиск
            self.assertIn("не подходит", wrong.note)

    def test_verify_output(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "a.pdf"
            with self.assertRaises(ConversionError):
                verify_output(p, "pdf")
            p.write_bytes(b"")
            with self.assertRaises(ConversionError):
                verify_output(p, "pdf")
            p.write_bytes(b"<html>")
            with self.assertRaises(ConversionError):
                verify_output(p, "pdf")
            p.write_bytes(b"%PDF-1.7 ...")
            verify_output(p, "pdf")


if __name__ == "__main__":
    unittest.main()
