import csv
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from utils.export import (ExportError, export_rows, guard_formula, rows_to_tsv, unique_names)
from utils.files import read_text_auto, write_sql_file
from utils.tasks import CancelToken

COLS = ["id", "name", "note"]
ROWS = [(1, "Анна", None), (2, "=SUM(A1)", "x;y"), (3, "-5", "line\nbreak")]


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_csv_roundtrip_with_guard(self):
        r = export_rows(self.d / "a.csv", "csv", COLS, ROWS)
        self.assertEqual(r.rows, 3)
        with open(self.d / "a.csv", encoding="utf-8-sig", newline="") as fh:
            data = list(csv.reader(fh, delimiter=";"))
        self.assertEqual(data[0], COLS)
        self.assertEqual(data[1], ["1", "Анна", ""])
        self.assertEqual(data[2][1], "'=SUM(A1)")     # формула обезврежена
        self.assertEqual(data[3][1], "-5")            # число-как-текст не портится
        self.assertEqual(data[3][2], "line\nbreak")
        self.assertTrue(r.warnings)
        self.assertEqual(os.listdir(self.d), ["a.csv"])  # временных файлов нет

    def test_csv_without_guard_and_comma(self):
        export_rows(self.d / "b.csv", "csv", COLS, ROWS, formula_guard=False, delimiter=",")
        with open(self.d / "b.csv", encoding="utf-8-sig", newline="") as fh:
            self.assertEqual(list(csv.reader(fh))[2][1], "=SUM(A1)")

    def test_json(self):
        import datetime as dt
        import decimal
        rows = [(1, decimal.Decimal("1.50"), dt.datetime(2024, 1, 2, 3, 4), b"\x01", float("nan"))]
        r = export_rows(self.d / "a.json", "json", ["a", "a", "c", "d", "e"], rows)
        data = json.loads((self.d / "a.json").read_text(encoding="utf-8"))
        self.assertEqual(data, [{"a": 1, "a_2": "1.50", "c": "2024-01-02T03:04:00", "d": "0x01", "e": None}])
        self.assertTrue(any("Повторяющиеся" in w for w in r.warnings))

    def test_json_empty(self):
        export_rows(self.d / "e.json", "json", ["a"], [])
        self.assertEqual(json.loads((self.d / "e.json").read_text()), [])

    def test_xlsx_content_and_safety(self):
        import openpyxl
        rows = [(1, "=1+1", "ok\x00bad"), (2 ** 62, "x" * 40000, None)]
        r = export_rows(self.d / "a.xlsx", "xlsx", COLS, rows)
        ws = openpyxl.load_workbook(self.d / "a.xlsx").active
        self.assertEqual([c.value for c in ws[1]], COLS)
        self.assertEqual(ws["B2"].value, "=1+1")
        self.assertEqual(ws["B2"].data_type, "s")           # не формула
        self.assertEqual(ws["C2"].value, "okbad")            # недопустимые символы удалены
        self.assertEqual(ws["A3"].value, str(2 ** 62))       # bigint — текстом
        self.assertEqual(len(ws["B3"].value), 32767)
        self.assertEqual(len(r.warnings), 2)

    def test_xlsx_splits_sheets(self):
        import openpyxl
        r = export_rows(self.d / "big.xlsx", "xlsx", ["n"], [(i,) for i in range(25)], max_rows_per_sheet=10)
        wb = openpyxl.load_workbook(self.d / "big.xlsx")
        self.assertEqual(wb.sheetnames, ["Результат", "Результат (2)", "Результат (3)"])
        self.assertEqual([ws.max_row for ws in wb.worksheets], [11, 11, 6])  # + строка заголовка на каждом листе
        self.assertEqual(r.sheets, 3)
        self.assertTrue(any("листов" in w for w in r.warnings))

    def test_xlsx_too_many_columns(self):
        with self.assertRaises(ExportError):
            export_rows(self.d / "w.xlsx", "xlsx", [f"c{i}" for i in range(16385)], [])

    def test_cancel_removes_temp_and_keeps_old_file(self):
        target = self.d / "keep.csv"
        target.write_text("old")
        tok = CancelToken()
        tok.cancel()
        with self.assertRaises(InterruptedError):
            export_rows(target, "csv", ["n"], [(i,) for i in range(10)], cancel_token=tok)
        self.assertEqual(target.read_text(), "old")
        self.assertEqual(os.listdir(self.d), ["keep.csv"])

    def test_disk_full_is_reported_and_cleaned(self):
        with mock.patch("utils.export.ensure_free_space", side_effect=__import__("utils.files", fromlist=["x"]).InsufficientDiskSpaceError("нет места")):
            with self.assertRaises(ExportError) as ctx:
                export_rows(self.d / "x.csv", "csv", ["a"], [(1,)])
        self.assertIn("нет места", str(ctx.exception))
        self.assertEqual(os.listdir(self.d), [])

    def test_replace_failure_keeps_old(self):
        target = self.d / "k.csv"
        target.write_text("old")
        with mock.patch("os.replace", side_effect=OSError("denied")):
            with self.assertRaises(ExportError):
                export_rows(target, "csv", ["a"], [(1,)])
        self.assertEqual(target.read_text(), "old")
        self.assertEqual(os.listdir(self.d), ["k.csv"])

    def test_unknown_format_and_bad_delimiter(self):
        with self.assertRaises(ExportError):
            export_rows(self.d / "x.pdf", "pdf", ["a"], [])
        with self.assertRaises(ExportError):
            export_rows(self.d / "x.csv", "csv", ["a"], [], delimiter=";;")
        self.assertEqual(os.listdir(self.d), [])

    def test_missing_directory_is_created(self):
        export_rows(self.d / "sub" / "dir" / "a.csv", "csv", ["a"], [(1,)])
        self.assertTrue((self.d / "sub" / "dir" / "a.csv").exists())

    def test_guard_formula_and_helpers(self):
        for bad in ("=1", "@x", "+cmd", "-cmd", "\tx"):
            self.assertTrue(guard_formula(bad).startswith("'"), bad)
        for ok in ("abc", "-5", "+7", "-1.5", "", "12"):
            self.assertEqual(guard_formula(ok), ok)
        self.assertEqual(unique_names(["a", "a", "", "a"]), ["a", "a_2", "column3", "a_3"])

    def test_tsv(self):
        self.assertEqual(rows_to_tsv(["a", "b"], [(1, "x\ty"), (None, "l1\nl2")], True), "a\tb\n1\tx y\n\tl1 l2")

    def test_large_table_streams_in_chunks(self):
        rows = [(i, "v" * 20) for i in range(30000)]
        r = export_rows(self.d / "big.csv", "csv", ["i", "v"], rows, chunk=1000)
        self.assertEqual(r.rows, 30000)
        with open(self.d / "big.csv", encoding="utf-8-sig") as fh:
            self.assertEqual(sum(1 for _ in fh), 30001)


class SqlFileTests(unittest.TestCase):
    def test_encodings_and_roundtrip(self):
        with tempfile.TemporaryDirectory() as d:
            d = Path(d)
            (d / "u8.sql").write_bytes("SELECT 'Привет'".encode("utf-8"))
            (d / "bom.sql").write_bytes(b"\xef\xbb\xbf" + "SELECT 'Я'".encode("utf-8"))
            (d / "u16.sql").write_bytes("SELECT 'Я'".encode("utf-16"))
            (d / "cp.sql").write_bytes("SELECT 'Привет'".encode("cp1251"))
            self.assertEqual(read_text_auto(d / "u8.sql"), ("SELECT 'Привет'", "utf-8"))
            self.assertEqual(read_text_auto(d / "bom.sql"), ("SELECT 'Я'", "utf-8-sig"))
            self.assertEqual(read_text_auto(d / "u16.sql"), ("SELECT 'Я'", "utf-16"))
            self.assertEqual(read_text_auto(d / "cp.sql"), ("SELECT 'Привет'", "cp1251"))
            write_sql_file(d / "out.sql", "a\nb")
            raw = (d / "out.sql").read_bytes()
            self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
            self.assertEqual(raw[3:], b"a\r\nb")
            self.assertEqual(read_text_auto(d / "out.sql")[0], "a\nb")

    def test_too_large_file_refused(self):
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "big.sql"
            p.write_bytes(b"x" * 2048)
            with self.assertRaises(ValueError):
                read_text_auto(p, max_bytes=1024)


if __name__ == "__main__":
    unittest.main()
