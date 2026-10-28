"""Дымовая проверка интерфейса.

* Если PySide6 установлен — реальный запуск главного окна в режиме offscreen.
* Если нет — запуск с заглушкой Qt (tests/qt_stub.py): проверяет импорты и конструкторы.
"""
import importlib.util
import os
import tempfile
import unittest
from unittest import mock

from bootstrap import create_services
from security.credentials import InMemoryCredentialStore
from tests import qt_stub
from tests.fakes import FakeServer

HAVE_QT = importlib.util.find_spec("PySide6") is not None


def no_modal_dialogs() -> list:
    """На настоящем Qt модальное окно в тесте ждало бы нажатия вечно. Здесь все окна сразу «отвечают»
    безопасно: «Нет» / «Отмена» / пусто. Тесты, которым нужен другой ответ, подменяют его явно."""
    from PySide6.QtGui import QColor
    from PySide6.QtWidgets import QColorDialog, QDialog, QFileDialog, QInputDialog, QMessageBox
    no, ok = QMessageBox.StandardButton.No, QMessageBox.StandardButton.Ok
    static = lambda value: staticmethod(lambda *a, **k: value)
    return [
        mock.patch.object(QMessageBox, "question", static(no)),
        mock.patch.object(QMessageBox, "warning", static(no)),
        mock.patch.object(QMessageBox, "information", static(ok)),
        mock.patch.object(QMessageBox, "critical", static(ok)),
        mock.patch.object(QMessageBox, "exec", lambda self: 0),
        mock.patch.object(QDialog, "exec", lambda self: 0),
        mock.patch.object(QInputDialog, "getText", static(("", False))),
        mock.patch.object(QInputDialog, "getMultiLineText", static(("", False))),
        mock.patch.object(QInputDialog, "getItem", static(("", False))),
        mock.patch.object(QFileDialog, "getOpenFileName", static(("", ""))),
        mock.patch.object(QFileDialog, "getOpenFileNames", static(([], ""))),
        mock.patch.object(QFileDialog, "getSaveFileName", static(("", ""))),
        mock.patch.object(QFileDialog, "getExistingDirectory", static("")),
        mock.patch.object(QColorDialog, "getColor", static(QColor())),
    ]
EXPECTED_PAGES = {"home", "connections", "explorer", "editor", "reports", "templates", "converter",
                  "automation", "monitoring", "compare", "docs", "history", "settings", "help"}


class UiSmokeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.saved = None
        if HAVE_QT:
            os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
            if os.name == "nt":                     # offscreen не видит системные шрифты Windows
                os.environ.setdefault("QT_QPA_FONTDIR", os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts"))
            from PySide6.QtWidgets import QApplication
            cls.app = QApplication.instance() or QApplication([])
            cls.dialog_patches = no_modal_dialogs()
            for patch in cls.dialog_patches:
                patch.start()
        else:
            cls.saved = qt_stub.install()

    @classmethod
    def tearDownClass(cls):
        if cls.saved is not None:
            qt_stub.uninstall(cls.saved)
        for patch in getattr(cls, "dialog_patches", []):
            patch.stop()

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.server = FakeServer()
        self.svc = create_services(self.tmp.name, InMemoryCredentialStore(), self.server.make_connection)
        self.modals: list[str] = []
        if HAVE_QT:
            self._guard_modals()

    def _guard_modals(self) -> None:
        """На настоящем Qt (offscreen) модальное окно ждало бы нажатия вечно. Все модальные вызовы сразу
        возвращают безопасный ответ («Нет» / «Отмена»); тест, которому нужно «Да», задаёт это явно."""
        from PySide6.QtGui import QColor
        from PySide6.QtWidgets import QColorDialog, QDialog, QFileDialog, QInputDialog, QMessageBox
        no = QMessageBox.StandardButton.No

        def rec(name, value):
            def f(*_a, **_k):
                self.modals.append(name)
                return value
            return f

        targets = [(QMessageBox, n, no) for n in ("information", "warning", "critical", "question")]
        targets += [(QMessageBox, "exec", 0), (QDialog, "exec", 0),
                    (QInputDialog, "getText", ("", False)), (QInputDialog, "getMultiLineText", ("", False)),
                    (QInputDialog, "getItem", ("", False)), (QInputDialog, "getInt", (0, False)),
                    (QFileDialog, "getOpenFileName", ("", "")), (QFileDialog, "getOpenFileNames", ([], "")),
                    (QFileDialog, "getSaveFileName", ("", "")), (QFileDialog, "getExistingDirectory", ""),
                    (QColorDialog, "getColor", QColor())]
        for cls, name, value in targets:
            p = mock.patch.object(cls, name, side_effect=rec(f"{cls.__name__}.{name}", value))
            p.start()
            self.addCleanup(p.stop)

    def tearDown(self):
        self.svc.shutdown()
        self.tmp.cleanup()

    def test_main_window_builds_all_pages(self):
        from ui.main_window import SECTIONS, MainWindow
        win = MainWindow(self.svc)
        win.ctx.automation.stop()                                         # планировщик запускается вместе с окном
        self.assertEqual({k for k, *_ in SECTIONS}, EXPECTED_PAGES)
        self.assertEqual(set(win._pages), EXPECTED_PAGES)
        from ui.placeholder_page import PlaceholderPage
        stubs = [k for k, page in win._pages.items() if isinstance(page, PlaceholderPage)]
        self.assertEqual(stubs, [], "в окне остались заглушки вместо готовых страниц")
        self.assertTrue(all(stage is None for _k, _t, stage, _d in SECTIONS))
        expected_types = {"home": "HomePage", "connections": "ConnectionsPage", "explorer": "ExplorerPage",
                          "editor": "EditorPage", "reports": "ReportsPage", "templates": "TemplatesPage",
                          "converter": "ConverterPage", "automation": "AutomationPage", "monitoring": "MonitoringPage",
                          "compare": "ComparePage", "docs": "DocumentationPage", "history": "HistoryPage",
                          "settings": "SettingsPage", "help": "HelpPage"}
        for key, cls in expected_types.items():
            if key in win._pages:
                self.assertEqual(type(win._pages[key]).__name__, cls, key)
        win.show_page("settings")
        win.show_page("home")
        win.ctx.theme_changed.emit("dark")
        win.ctx.connections_changed.emit()

    def test_pages_refresh_with_data(self):
        from database.profile import ConnectionProfile
        from ui.context import AppContext
        from ui.connections_page import ConnectionsPage
        from ui.explorer_page import ExplorerPage
        from ui.history_page import HistoryPage
        from ui.home_page import HomePage
        p = self.svc.profiles.save(ConnectionProfile.new("Dev", "srv"))
        self.svc.profiles.connect(p)
        self.svc.oplog.add("task", "demo", "ok", "x", 5)
        ctx = AppContext(self.svc)
        for cls in (HomePage, ConnectionsPage, ExplorerPage, HistoryPage):
            page = cls(ctx)
            if hasattr(page, "refresh"):
                page.refresh()
        ctx.connections_changed.emit()

    @unittest.skipIf(HAVE_QT, "ручной ввод в поля имитируется только заглушкой")
    def test_connections_form_save_and_connect_logic(self):
        from database.profile import AUTH_SQL
        from ui.context import AppContext
        from ui.connections_page import ConnectionsPage
        ctx = AppContext(self.svc)
        page = ConnectionsPage(ctx)
        page._new()
        page.f_server.setText(r"host\SQLEXPRESS")
        page.f_auth.setCurrentIndex(page.f_auth.findData(AUTH_SQL))
        page.f_user.setText("sa")
        saved = page._save()
        self.assertIsNotNone(saved)
        self.assertEqual((saved.server, saved.instance), ("host", "SQLEXPRESS"))
        self.assertEqual(len(self.svc.profiles.repo.list()), 1)
        page.f_server.setText("")  # некорректный ввод не сохраняется
        self.assertIsNone(page._save())
        self.assertEqual(len(self.svc.profiles.repo.list()), 1)

    def _editor(self):
        from database.profile import ConnectionProfile
        from ui.context import AppContext
        from ui.editor_page import EditorPage
        p = self.svc.profiles.save(ConnectionProfile.new("Dev", "srv"))
        self.svc.profiles.connect(p)
        page = EditorPage(AppContext(self.svc))
        return page, page.tab()

    def test_editor_runs_query_in_background_task(self):
        from utils.tasks import TaskState
        page, tab = self._editor()
        self.assertIsNotNone(tab)
        tab.editor.sql_to_run = lambda: ("SELECT 1 AS x", 0)
        page.run()
        self.assertIsNotNone(tab.task)
        self.assertTrue(tab.task.wait(3))
        self.assertEqual(tab.task.state, TaskState.SUCCEEDED)
        report = tab.task.result
        self.assertEqual(report.results[0].rows, [(1,)])
        self.assertEqual(len(self.svc.manager.connections()), 2)    # основное + собственное соединение вкладки
        for t in page._all_tabs():
            page._drop_session(t)
        self.assertEqual(len(self.svc.manager.connections()), 1)

    def test_editor_blocks_modifying_sql_in_read_only_mode(self):
        page, tab = self._editor()
        for sql in ("DELETE FROM t WHERE id = 1", "DROP TABLE t", "UPDATE t SET a = 1"):
            tab.editor.sql_to_run = lambda sql=sql: (sql, 0)
            with mock.patch("ui.editor_page.QMessageBox"):
                page.run()
            self.assertIsNone(tab.task, sql)
        self.assertFalse(any(q.startswith(("DELETE", "DROP", "UPDATE")) for q, _ in self.server.executed))

    def test_editor_rejects_syntax_errors_without_contacting_server(self):
        page, tab = self._editor()
        tab.editor.sql_to_run = lambda: ("SELECT (1", 0)
        page.run()
        self.assertIsNone(tab.task)
        self.assertEqual(len(self.svc.manager.connections()), 1)

    def test_dialogs_construct(self):
        from ui.export_dialog import ExportDialog
        from ui.query_dialogs import FavoritesDialog, HistoryDialog
        self.svc.history.add("SELECT 1", "ok", connection_name="x")
        self.svc.favorites.add("Fav", "SELECT 2")
        ExportDialog(None, settings=self.svc.settings, shown_rows=5, truncated=True, full_allowed=True)
        ExportDialog(None, settings=self.svc.settings, shown_rows=5, truncated=False, full_allowed=False,
                     full_reason="только чтение")
        HistoryDialog(None, self.svc, lambda sql, new: None, lambda sql: None)
        FavoritesDialog(None, self.svc, lambda sql, new: None)

    def _ctx_with_connection(self):
        from database.profile import ConnectionProfile
        from ui.context import AppContext
        p = self.svc.profiles.save(ConnectionProfile.new("Dev", "srv"))
        self.svc.profiles.connect(p)
        return AppContext(self.svc), p

    def test_template_builder_roundtrips_every_builtin_template(self):
        from reports.model import validate_template
        from ui.templates_page import TemplatesPage
        ctx, _ = self._ctx_with_connection()
        page = TemplatesPage(ctx)
        builtin = [i for i in page.store.list() if i.builtin]
        self.assertGreaterEqual(len(builtin), 3)
        for info in builtin:
            t = page.store.load(info.id)
            page.load(t)
            c = page.collect()
            self.assertEqual(validate_template(c), [], info.name)
            self.assertEqual(c.to_dict(), t.to_dict(), info.name)       # ни одно поле не потерялось

    def test_template_builder_edit_operations_and_save(self):
        from reports.model import DataSource, validate_template
        from ui.templates_page import TemplatesPage
        ctx, _ = self._ctx_with_connection()
        page = TemplatesPage(ctx)
        page.load(page.store.load("builtin-db-overview"))
        n = len(page.blocks)
        page._select_block(1)
        page.add_block("text")
        self.assertEqual((len(page.blocks), page.blocks[2].type), (n + 1, "text"))
        page.move_block(-1)
        self.assertEqual(page.blocks[1].type, "text")
        page.remove_block()
        self.assertEqual(len(page.blocks), n)
        page.add_source()
        self.assertEqual(page.sources[-1].id, "s1")                     # первый свободный идентификатор
        page.sources.append(DataSource(id="t2", name="Табличный", kind="table", table="Orders", columns=["Id"]))
        page._refresh_sources()
        page._select_source(len(page.sources) - 1)
        page.show_sql()
        self.assertIn("SELECT TOP (100) [Id] FROM [dbo].[Orders]", page.s_preview.toPlainText())
        self.assertTrue(page._dirty)
        # встроенный шаблон не перезаписывается: сохраняется копия
        with mock.patch("ui.templates_page.QMessageBox"):
            self.assertTrue(page.save())
        ids = {i.id: i for i in page.store.list()}
        self.assertNotEqual(page.t.id, "builtin-db-overview")
        self.assertFalse(ids[page.t.id].builtin)
        self.assertEqual(validate_template(page.store.load(page.t.id)), [])
        self.assertFalse(page._dirty)

    def test_reports_page_generates_files_in_background(self):
        import tempfile as _tf
        from pathlib import Path
        from reports.model import Block, DataSource, ReportTemplate
        from ui.home_page import HomePage
        from ui.reports_page import ReportsPage
        ctx, _ = self._ctx_with_connection()
        t = ReportTemplate(id="ui-t1", name="UI-отчёт", title="UI",
                           sources=[DataSource(id="s1", name="Данные", kind="sql", sql="SELECT a, b FROM t1")],
                           blocks=[Block(id="t", type="table", title="Таблица", source="s1")])
        ctx.reports.templates.save(t)
        self.server.when("FROM t1", columns=["a", "b"], rows=[(1, 2.5), (2, 5.0)])
        page = ReportsPage(ctx)
        page.c_tpl.setCurrentIndex(page.c_tpl.findData("ui-t1"))
        page._on_template(0)
        self.assertEqual(page.template.id, "ui-t1")
        with _tf.TemporaryDirectory() as out:
            page.e_dir.setText(out)
            page.e_name.setText("ui-report")
            for f, cb in page.k_fmt.items():
                cb.setChecked(f in ("docx", "html"))
            with mock.patch("ui.reports_page.QMessageBox"):
                page.run()
            self.assertIsNotNone(page._task)
            self.assertTrue(page._task.wait(30))
            res = page._task.result
            self.assertTrue(res.ok, getattr(res, "errors", res))
            self.assertEqual(sorted(Path(f.path).name for f in res.files), ["ui-report.docx", "ui-report.html"])
            self.assertEqual(self.svc.oplog.recent(1, "report")[0].status, "ok")
            home = HomePage(ctx)
            home.refresh()
            self.assertEqual(home._recent[0].name, "UI-отчёт")
            self.assertEqual(len(self.svc.manager.connections()), 1)        # временное соединение закрыто

    def test_reports_page_refuses_invalid_template_and_no_formats(self):
        from reports.model import DataSource, ReportTemplate
        from ui.reports_page import ReportsPage
        ctx, _ = self._ctx_with_connection()
        bad = ReportTemplate(id="ui-bad", name="Плохой", sources=[DataSource(id="s1", kind="sql", sql="DELETE FROM t")])
        ctx.reports.templates.save(bad)
        page = ReportsPage(ctx)
        page.c_tpl.setCurrentIndex(page.c_tpl.findData("ui-bad"))
        page._on_template(0)
        with mock.patch("ui.reports_page.QMessageBox") as box:
            page.run()
        self.assertIsNone(page._task)
        self.assertTrue(box.warning.called or box.information.called)
        self.assertFalse(any(q.startswith("DELETE") for q, _ in self.server.executed))

    def test_converter_page_queue_and_batch(self):
        import tempfile as _tf
        from pathlib import Path
        from ui.context import AppContext
        from ui.converter_page import CompatDialog, ConverterPage
        ctx = AppContext(self.svc)
        page = ConverterPage(ctx)
        with _tf.TemporaryDirectory() as tmp:
            src, out = Path(tmp) / "in", Path(tmp) / "out"
            (src / "sub").mkdir(parents=True)
            (src / "a.md").write_text("# Заголовок\n\nтекст", encoding="utf-8")
            (src / "sub" / "b.md").write_text("# B", encoding="utf-8")
            (src / "t.csv").write_text("x;y\n1;2\n", encoding="utf-8")
            (src / "doc.pdf").write_bytes(b"%PDF-1.4\n")
            (src / "run.exe").write_bytes(b"MZ")
            page._add_paths([str(src), str(src / "a.md")])                 # повторное добавление не дублирует
            self.assertEqual(sorted(j.src.name for j in page.jobs), ["a.md", "b.md", "doc.pdf", "t.csv"])
            page._add_paths([str(src)])
            self.assertEqual(len(page.jobs), 4)
            targets = [page.c_target.itemData(i) for i in range(page.c_target.count())]
            self.assertIn("html", targets)
            self.assertIn("xlsx", targets)                                  # csv -> xlsx доступен
            page.c_target.setCurrentIndex(page.c_target.findData("html"))
            page._refresh_notes()
            self.assertIn("Markdown (.md) → HTML", page.notes.toPlainText())
            page.e_dir.setText(str(out))
            page.k_same.setChecked(False)
            page.run()
            self.assertIsNotNone(page._task)
            self.assertTrue(page._task.wait(30))
            summary = page._task.result
            self.assertEqual((summary.ok, summary.skipped, summary.failed), (2, 2, 0))   # csv и pdf -> html не поддерживаются
            self.assertTrue((out / "a.html").exists())
            self.assertTrue((out / "sub" / "b.html").exists())               # структура папок сохранена
            self.assertEqual(summary.check_consistency(), [])
            self.assertEqual(len(ctx.converter.recent()), 4)
        CompatDialog(None, ctx.converter.caps())

    def test_task_dialog_roundtrips_every_task_kind(self):
        from automation.schedule import Schedule
        from automation.scripts import Scenario
        from automation.store import ScheduledTask
        from reports.model import Block, DataSource, ReportTemplate
        from ui.automation_page import TaskDialog
        ctx, p = self._ctx_with_connection()
        ctx.reports.templates.save(ReportTemplate(id="t1", name="Т", sources=[DataSource(id="s", kind="sql", sql="SELECT 1")],
                                                  blocks=[Block(id="b", type="table", source="s")]))
        scn = Scenario.new("Скрипт", "SELECT 1")
        scn.profile_id = p.id
        ctx.automation.scenarios.save(scn)
        configs = {
            "health": {"latency_warn_ms": 750},
            "table_stats": {"database": "Sales", "top": 12},
            "export": {"database": "Sales", "tables": ["dbo.A", "hr.B"], "format": "xlsx", "output_dir": "/tmp/x", "max_rows": 100},
            "backup": {"databases": ["Sales", "HR"], "folder": "D:\\bak", "copy_only": True, "checksum": False, "verify": True,
                       "compression": False, "cleanup": {"enabled": True, "local_folder": "\\\\srv\\bak", "keep_last": 3,
                                                         "older_than_days": 30, "acknowledged": True}},
            "report": {"template_id": "t1", "database": "Sales", "formats": ["docx", "pdf"], "output_dir": "/tmp/r",
                       "policy": "skip", "params": {"year": "2026"}, "keep_history": True, "compare_keys": {"s": ["id", "y"]}},
            "script": {"scenario_id": scn.id, "params": {"a": "1"}},
        }
        schedules = [Schedule("weekly", time="09:15", weekdays=[1, 3]), Schedule("monthly", time="06:00", day=31),
                     Schedule("interval", interval_minutes=45), Schedule("daily", time="23:30"),
                     Schedule("once", at="2027-01-02 03:04"), Schedule("on_connect")]
        for (kind, cfg), sched in zip(configs.items(), schedules):
            cfg = dict(cfg, notify="always")
            if kind != "script":
                cfg["profile_id"] = p.id
            t = ScheduledTask.new(f"Задача {kind}", kind, sched, cfg)
            t.catch_up = kind == "health"
            if kind == "backup":
                t.schedule = Schedule("daily", time="02:00")
            self.assertEqual(ctx.automation.save_task(t), [], kind)
            dlg = TaskDialog(ctx, None, ctx.automation.tasks.get(t.id))
            got = dlg.collect()
            self.assertEqual((got.kind, got.name, got.catch_up), (t.kind, t.name, t.catch_up), kind)
            self.assertEqual(got.schedule.kind, t.schedule.kind, kind)
            for key, val in cfg.items():
                self.assertEqual(got.config.get(key), val, (kind, key))
            self.assertEqual(ctx.automation.validate(got), [], kind)

    def test_automation_page_scenario_flow(self):
        from ui.automation_page import AutomationPage, WindowsTaskDialog
        ctx, p = self._ctx_with_connection()
        page = AutomationPage(ctx)
        page.new_scenario()
        scn = ctx.automation.scenarios.list()[0]
        page._load_scn(scn)
        page.s_sql.setPlainText("UPDATE t SET a = 1 WHERE id = 1")
        yes = mock.MagicMock()
        yes.warning.return_value = yes.StandardButton.Yes                  # «Разрешить изменения?» — «Да»
        with mock.patch("ui.automation_page.QMessageBox", yes):
            page.k_mod.setChecked(True)                                     # на настоящем Qt вызывает предупреждение
        with mock.patch("ui.automation_page.QMessageBox"), \
                mock.patch("ui.automation_page.QInputDialog.getMultiLineText", return_value=("ВЫПОЛНИТЬ", True)):
            page.approve_scenario()
        stored = ctx.automation.scenarios.get(scn.id)
        self.assertTrue(stored.allow_modifications and stored.is_approved)
        self.server.when("UPDATE t", sets=[{"rowcount": 1}])
        box = mock.MagicMock()
        box.question.return_value = box.StandardButton.Yes                  # пользователь отвечает «Да»
        with mock.patch("ui.automation_page.QMessageBox", box):
            page.run_scenario_now()
        for t in self.svc.tasks.tasks():
            t.wait(5)
        self.assertEqual(ctx.automation.scenarios.runs(scn.id)[0]["status"], "ok")
        from automation.schedule import Schedule
        from automation.store import ScheduledTask
        task = ScheduledTask.new("H", "health", Schedule("daily"), {"profile_id": p.id})
        ctx.automation.save_task(task)
        page.refresh_tasks()
        WindowsTaskDialog(None, task)
        page.refresh_notes()

    def test_monitoring_page_collects_and_stores_metrics(self):
        from ui.monitoring_page import MonitoringPage
        ctx, p = self._ctx_with_connection()
        self.server.when("HAS_PERMS_BY_NAME", columns=["p"], rows=[(0,)])
        self.server.when("sys.master_files", columns=["d", "a", "b"], rows=[("Sales", 10.0, 1.0)])
        page = MonitoringPage(ctx)
        page.e_db.setText("Sales")
        page.refresh()
        tasks = [t for t in self.svc.tasks.tasks() if t.name.startswith("Мониторинг")]
        self.assertTrue(tasks and tasks[-1].wait(5))
        snap = tasks[-1].result
        page.show_snapshot(snap)
        self.assertTrue(snap.probe.ok)
        self.assertEqual(ctx.automation.metrics.latest(p.id, "db_size_mb")["Sales"][1], 11.0)
        self.assertTrue(any("VIEW SERVER STATE" in n for n in snap.notes))

    def _two_databases(self):
        cols = lambda *extra: [(1, 1, "Id", "int", 4, 10, 0, False, True, False, None, None, None, 1, 1, None)] + list(extra)
        self.server.when("FROM [A].sys.tables t JOIN", columns=["o", "s", "n"], rows=[(1, "dbo", "T"), (2, "dbo", "OnlyA")])
        self.server.when("FROM [B].sys.tables t JOIN", columns=["o", "s", "n"], rows=[(1, "dbo", "T")])
        self.server.when("FROM [A].sys.columns c JOIN", columns=list("abcdefghijklmnop"),
                         rows=cols((1, 2, "Name", "nvarchar", 100, 0, 0, True, False, False, None, None, None, None, None, None),
                                   (2, 1, "Id", "int", 4, 10, 0, False, False, False, None, None, None, None, None, None)))
        self.server.when("FROM [B].sys.columns c JOIN", columns=list("abcdefghijklmnop"), rows=cols())

    def test_compare_page_schema_script_and_confirmation(self):
        from ui.compare_page import ComparePage, SyncScriptDialog
        from comparison.sync_script import sync_schema_script
        ctx, p = self._ctx_with_connection()
        self._two_databases()
        page = ComparePage(ctx)
        for side, db in ((page.s_src, "A"), (page.s_tgt, "B")):
            side["profile"].setCurrentIndex(side["profile"].findData(p.id))
            side["db"].setText(db)
        page.compare_schema()
        task = [t for t in self.svc.tasks.tasks() if t.name == "Сравнение структуры"][-1]
        self.assertTrue(task.wait(10))
        page._schema_done(task.result)
        objs = {(d.object, d.status) for d in page.diff.items}
        self.assertIn(("dbo.OnlyA", "only_source"), objs)
        self.assertIn(("dbo.T.Name", "only_source"), objs)
        script = sync_schema_script(page.diff)
        dlg = SyncScriptDialog(ctx, None, script, p.id, "B")
        with mock.patch("ui.compare_page.QInputDialog.getText", return_value=("A", True)), \
                mock.patch("ui.compare_page.QMessageBox"):
            dlg._run()                                                      # неверное имя базы — не выполняется
        self.assertFalse(any("ALTER TABLE" in q or "CREATE TABLE" in q for q, _ in self.server.executed))
        with mock.patch("ui.compare_page.QInputDialog.getText", return_value=("B", True)), \
                mock.patch("ui.compare_page.QMessageBox"):
            dlg._run()
        sync = [t for t in self.svc.tasks.tasks() if t.name.startswith("Синхронизация")][-1]
        self.assertTrue(sync.wait(10))
        self.assertTrue(sync.result.ok, sync.result.error)
        self.assertTrue(any("CREATE TABLE [dbo].[OnlyA]" in q for q, _ in self.server.executed))
        self.assertIn("USE [B]", [q for q, _ in self.server.executed])

    def test_compare_page_data(self):
        from ui.compare_page import ComparePage
        ctx, p = self._ctx_with_connection()
        self.server.when("FROM [A].[dbo].[T]", columns=["Id", "V"], rows=[(1, "a"), (2, "x")])
        self.server.when("FROM [B].[dbo].[T]", columns=["Id", "V"], rows=[(1, "a"), (3, "c")])
        page = ComparePage(ctx)
        for side, db in ((page.d_src, "A"), (page.d_tgt, "B")):
            side["profile"].setCurrentIndex(side["profile"].findData(p.id))
            side["db"].setText(db)
            side["table"].setText("dbo.T")
        page.e_keys.setText("Id")
        page.compare_data()
        task = [t for t in self.svc.tasks.tasks() if t.name == "Сравнение данных"][-1]
        self.assertTrue(task.wait(10))
        page._data_done(task.result)
        self.assertEqual((page.data.diff.added_total, page.data.diff.removed_total), (1, 1))

    def test_documentation_page_and_open_in_editor(self):
        from ui.documentation_page import DocumentationPage
        from ui.main_window import MainWindow
        import tempfile as _tf
        from pathlib import Path
        ctx, p = self._ctx_with_connection()
        self._two_databases()
        page = DocumentationPage(ctx)
        page.c_profile.setCurrentIndex(page.c_profile.findData(p.id))
        page.e_db.setText("A")
        with _tf.TemporaryDirectory() as out:
            page.e_dir.setText(out)
            page.generate()
            task = [t for t in self.svc.tasks.tasks() if t.name == "Документация: A"][-1]
            self.assertTrue(task.wait(20))
            page._done(task.result)
            self.assertEqual(len(page._files), 1)
            self.assertTrue(Path(page._files[0]).exists())
        win = MainWindow(self.svc)
        win.ctx.automation.stop()
        editor = win._pages["editor"]
        before = editor.tabs.count()
        win._open_sql("SELECT 42")
        self.assertEqual(editor.tabs.count(), before + 1)

    def test_reports_page_saved_query_and_code_appendix(self):
        import tempfile as _tf
        from pathlib import Path
        from docx import Document
        from reports.model import Block, DataSource, ReportTemplate
        from reports.quick import FAV_PREFIX
        from ui.reports_page import ReportsPage
        ctx, _p = self._ctx_with_connection()
        fav_id = self.svc.favorites.add("Остатки", "SELECT Item, Qty FROM dbo.Stock WHERE Qty > :min")
        self.server.when("FROM dbo.Stock", columns=["Item", "Qty"], rows=[("Болт", 120), ("Гайка", 75)])
        page = ReportsPage(ctx)
        idx = page.c_tpl.findData(f"{FAV_PREFIX}{fav_id}")
        self.assertGreaterEqual(idx, 0)                                      # сохранённый запрос в списке шаблонов
        page.c_tpl.setCurrentIndex(idx)
        page._on_template(0)
        self.assertEqual([b.type for b in page.template.blocks], ["query", "notes"])
        self.assertIn("min", page._param_edits)
        page._param_edits["min"].setText("10")
        with _tf.TemporaryDirectory() as out:
            page.e_dir.setText(out)
            page.e_name.setText("stock")
            for f, cb in page.k_fmt.items():
                cb.setChecked(f == "docx")
            with mock.patch("ui.reports_page.QMessageBox"):
                page.run()
            self.assertTrue(page._task.wait(30))
            res = page._task.result
            self.assertTrue(res.ok, getattr(res, "errors", res))
            doc = Document(str(Path(out) / "stock.docx"))
            self.assertIn("WHERE Qty > :min", "\n".join(par.text for par in doc.paragraphs))
            self.assertIn("Болт", [c.text for t in doc.tables for r in t.rows for c in r.cells])
        ctx.reports.templates.save(ReportTemplate(id="plain", name="Обычный", sources=[DataSource(id="s", name="S", kind="sql", sql="SELECT 1 AS a")],
                                                  blocks=[Block(id="b", type="table", source="s")]))
        page._reload_templates()
        page.c_tpl.setCurrentIndex(page.c_tpl.findData("plain"))
        page._on_template(0)
        page.k_code.setChecked(True)
        self.assertEqual([b.type for b in page._prepared_template().blocks], ["table", "heading", "query"])
        page.k_code.setChecked(False)
        self.assertEqual([b.type for b in page._prepared_template().blocks], ["table"])
        page.c_tpl.setCurrentIndex(page.c_tpl.findData(f"{FAV_PREFIX}{fav_id}"))
        page._on_template(0)
        page._edit_template()                                                # «Изменить» → сохраняется как шаблон
        self.assertTrue(any(i.name == "Запрос: Остатки" for i in ctx.reports.templates.list()))

    def test_template_builder_query_block_and_source_from_favorite(self):
        from reports.model import Block, DataSource, ReportTemplate
        from ui.templates_page import TemplatesPage
        ctx, _p = self._ctx_with_connection()
        self.svc.favorites.add("Клиенты", "SELECT * FROM dbo.Clients WHERE City = :city")
        page = TemplatesPage(ctx)
        t = ReportTemplate(id="qt", name="С запросом", sources=[DataSource(id="s", name="S", kind="sql", sql="SELECT 1 AS a")],
                           blocks=[Block(id="b1", type="query", source="s", title="Код", show_code=False, max_rows=25)])
        page.load(t)
        self.assertEqual(page.collect().to_dict(), t.to_dict())             # блок проходит через редактор без потерь
        with mock.patch("ui.templates_page.QInputDialog.getItem", return_value=("Клиенты", True)):
            page.add_source_from_favorite()
        c = page.collect()
        self.assertEqual(c.sources[-1].sql, "SELECT * FROM dbo.Clients WHERE City = :city")
        self.assertIn("city", [p.name for p in c.params])
        page.add_block("query")
        self.assertEqual(page.blocks[page._blk_idx].type, "query")

    def test_settings_libreoffice_status(self):
        import tempfile as _tf
        from pathlib import Path
        from ui.context import AppContext
        from ui.settings_page import SettingsPage
        page = SettingsPage(AppContext(self.svc))
        with _tf.TemporaryDirectory() as d:
            (Path(d) / "program").mkdir()
            (Path(d) / "program" / "soffice").write_bytes(b"x")
            page._save_text("libreoffice_path", d)                            # указана папка LibreOffice
            self.assertEqual(self.svc.settings.get("libreoffice_path"), d)
            page._auto_lo()
            self.assertEqual(self.svc.settings.get("libreoffice_path"), "")

    def test_theme_stylesheets_are_complete(self):
        from ui import theme
        for name in ("light", "dark"):
            qss = theme.build_stylesheet(name)
            self.assertNotIn("$", qss)  # все плейсхолдеры подставлены
            self.assertIn(theme.colors(name)["accent"], qss)
        self.assertEqual(set(theme.PALETTES["light"]), set(theme.PALETTES["dark"]))


if __name__ == "__main__":
    unittest.main()
