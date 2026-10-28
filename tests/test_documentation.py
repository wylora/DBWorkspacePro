import tempfile
import unittest
from pathlib import Path

from docx import Document

from bootstrap import create_services
from comparison.schema import Column, ForeignKey, Index, Module, Param, SchemaModel, Table
from comparison.service import CompareService
from database.profile import ConnectionProfile
from documentation.diagram import er_diagram_png, relationships
from documentation.generator import DocOptions, build_documentation
from documentation.service import DocumentationService
from documentation.store import SchemaSnapshotStore
from reports.adhoc import save_doc
from security.credentials import InMemoryCredentialStore
from tests.fakes import FakeOdbcError, FakeServer


def sample(n_tables=6, with_fk=True):
    m = SchemaModel("SRV", "Sales", "2026-10-20T10:00:00")
    for i in range(n_tables):
        t = Table("dbo", f"T{i}", [Column("Id", "int", False, True, 1, 1, ordinal=1),
                                   Column("ParentId", "int", True, ordinal=2, description="Ссылка <на> родителя & т.п.")],
                  [Index(f"PK_T{i}", "CLUSTERED", True, True, False, ["Id"])], rows=i * 10, description=f"Таблица {i}")
        if with_fk and i:
            t.fks.append(ForeignKey(f"FK_T{i}", ["ParentId"], "dbo", f"T{i - 1}", ["Id"], "CASCADE"))
        m.tables[SchemaModel.key("dbo", t.name)] = t
    m.modules[SchemaModel.key("dbo", "v", "view")] = Module("dbo", "v", "view", "VIEW", "CREATE VIEW dbo.v AS SELECT 1 AS a",
                                                             columns=[Column("a", "int", False)])
    m.modules[SchemaModel.key("dbo", "p", "procedure")] = Module("dbo", "p", "procedure", "SQL_STORED_PROCEDURE", None,
                                                                  params=[Param("@id", "int", False, 1), Param("@out", "int", True, 2)])
    m.modules[SchemaModel.key("dbo", "f", "function")] = Module("dbo", "f", "function", "SQL_SCALAR_FUNCTION",
                                                                 "CREATE FUNCTION dbo.f() RETURNS int AS BEGIN RETURN 1 END",
                                                                 params=[Param("", "int", False, 0)])
    m.modules[SchemaModel.key("dbo", "tr", "trigger")] = Module("dbo", "tr", "trigger", "SQL_TRIGGER", "CREATE TRIGGER tr ...",
                                                                 "dbo.T1")
    m.warnings.append("Описания объектов недоступны: нет прав")
    return m


class DocumentationTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def render(self, model, opts=None, previous=None, formats=("docx",)):
        res = save_doc(build_documentation(model, opts or DocOptions(), previous), self.d, "doc", list(formats), policy="overwrite")
        self.assertEqual(res.errors, {})
        return res

    def test_docx_structure(self):
        self.render(sample())
        doc = Document(str(self.d / "doc.docx"))
        heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        for h in ("Обзор", "Схема связей", "Таблицы", "dbo.T0", "dbo.T5", "Представления", "Хранимые процедуры", "Функции",
                  "Триггеры", "Сведения о формировании"):
            self.assertIn(h, heads)
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("таблиц — 6", text)
        self.assertIn("связей — 5", text)
        self.assertIn("Определение недоступно", text)                   # процедура без VIEW DEFINITION
        self.assertIn("Возвращает: int", text)
        self.assertIn("На эту таблицу ссылаются: dbo.T1 (FK_T1)", text)
        self.assertIn("Триггеры: tr.", text)
        self.assertEqual(len(doc.inline_shapes), 1)                       # схема связей
        cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
        self.assertIn("Ссылка <на> родителя & т.п.", cells)
        self.assertIn("FK → T0", cells)
        self.assertIn("CASCADE / NO ACTION", cells)

    def test_options_and_no_relationships(self):
        self.render(sample(with_fk=False), DocOptions(include_definitions=False, include_indexes=False))
        doc = Document(str(self.d / "doc.docx"))
        heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
        self.assertNotIn("Схема связей", heads)
        self.assertNotIn("CREATE FUNCTION", "\n".join(p.text for p in doc.paragraphs))
        self.assertEqual(len(doc.inline_shapes), 0)
        self.assertEqual(er_diagram_png(sample(with_fk=False))[0], None)

    def test_changes_since_previous_version(self):
        prev = sample()
        cur = sample()
        cur.tables["dbo.t5"].columns.append(Column("New", "int", True, ordinal=3))
        del cur.tables["dbo.t0"]
        self.render(cur, previous=prev)
        doc = Document(str(self.d / "doc.docx"))
        rows = [[c.text for c in r.cells] for t in doc.tables for r in t.rows]
        self.assertIn(["Столбец", "dbo.T5.New", "добавлено", "int, NULL"], rows)
        self.assertTrue(any(r[1] == "dbo.T0" and r[2] == "удалено" for r in rows if len(r) == 4))
        self.render(cur, previous=cur)
        self.assertIn("структура не изменилась", "\n".join(p.text for p in Document(str(self.d / "doc.docx")).paragraphs))

    def test_big_diagram_is_capped_and_html_pdf(self):
        png, note = er_diagram_png(sample(80), max_tables=30)
        self.assertTrue(png.startswith(b"\x89PNG"))
        self.assertIn("30 наиболее связанных таблиц из 80", note)
        self.assertEqual(len(relationships(sample(80))), 79)
        from converter.libreoffice import find_soffice
        formats = ["html", "md"] + (["pdf"] if find_soffice() else [])
        res = self.render(sample(), formats=formats)
        html = (self.d / "doc.html").read_text(encoding="utf-8")
        self.assertIn("Ссылка &lt;на&gt; родителя &amp; т.п.", html)
        self.assertIn("data:image/png;base64", html)
        self.assertEqual(len(res.files), len(formats))

    def test_snapshot_store(self):
        st = SchemaSnapshotStore(self.d / "snap", keep=2)
        self.assertIsNone(st.latest("p", "db"))
        for i in range(3):
            m = sample(i + 1)
            st.save("p", "db", m)
        self.assertEqual(len(st.latest("p", "db").tables), 3)
        self.assertEqual(len(list((self.d / "snap").rglob("*.json.gz"))), 2)
        broken = self.d / "snap" / "p" / "db" / "99999999_999999_999999.json.gz"
        broken.write_bytes(b"not gzip")
        self.assertIsNone(st.latest("p", "db"))


class ServicesTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.server = FakeServer()
        self.svc = create_services(Path(self.tmp.name) / "h", InMemoryCredentialStore(), self.server.make_connection)
        self.p = self.svc.profiles.save(ConnectionProfile.new("P", "srv"))
        self.server.when("sys.tables t JOIN", columns=["o", "s", "n"], rows=[(1, "dbo", "Orders")])
        self.server.when("FROM [Sales].sys.columns c JOIN", columns=list("abcdefghijklmnop"),
                         rows=[(1, 1, "Id", "int", 4, 10, 0, False, True, False, None, None, None, 1, 1, None)])

    def tearDown(self):
        self.svc.shutdown()
        self.tmp.cleanup()

    def test_documentation_service_snapshots_and_changes(self):
        ds = DocumentationService(self.svc)
        out = Path(self.tmp.name) / "out"
        r1 = ds.generate(self.p.id, "Sales", DocOptions(), ["docx"], output_dir=out)
        self.assertEqual((len(r1.files), r1.changes), (1, None))
        self.server.handlers[1]["rows"] = self.server.handlers[1]["rows"] + [
            (1, 2, "Total", "decimal", 9, 18, 2, True, False, False, None, None, None, None, None, None)]
        r2 = ds.generate(self.p.id, "Sales", DocOptions(), ["docx"], output_dir=out)
        self.assertEqual(r2.changes, 1)
        self.assertEqual(len(list(out.glob("*.docx"))), 2)
        self.assertEqual(self.svc.oplog.recent(1, "documentation")[0].status, "ok")
        self.assertEqual(self.svc.manager.count(), 0)

    def test_sync_script_runs_in_one_transaction_and_rolls_back(self):
        cs = CompareService(self.svc)
        self.server.when("ALTER TABLE [dbo].[Bad]", errors=[FakeOdbcError("42000", "[42000] Invalid object name (4902)")])
        script = ("-- заголовок\nALTER TABLE [dbo].[A] ADD [x] int NULL;\nGO\n-- только комментарий\nGO\n"
                  "ALTER TABLE [dbo].[Bad] ADD [y] int NULL;\nGO\nALTER TABLE [dbo].[C] ADD [z] int NULL;\nGO\n")
        r = cs.execute_script(self.p.id, "Target", script)
        self.assertFalse(r.ok)
        self.assertEqual((r.batches_done, r.batches_total, r.rolled_back), (1, 3, True))
        self.assertIn("Пакет 2", r.error)
        executed = [q for q, _ in self.server.executed]
        self.assertIn("USE [Target]", executed)
        self.assertFalse(any("[dbo].[C]" in q for q in executed))          # после ошибки не выполняется
        ok = cs.execute_script(self.p.id, "Target", "ALTER TABLE [dbo].[A] ADD [x] int NULL;\nGO\n")
        self.assertTrue(ok.ok)
        self.assertEqual(self.svc.oplog.recent(1, "compare")[0].status, "ok")


if __name__ == "__main__":
    unittest.main()
