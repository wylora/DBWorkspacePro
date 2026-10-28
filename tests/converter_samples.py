"""Образцы файлов всех поддерживаемых форматов для тестов конвертера."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from docx import Document

from converter.formats import PRIMARY_EXT
from converter.libreoffice import convert


def build_samples(folder: Path, soffice: str | None) -> dict[str, Path]:
    folder.mkdir(parents=True, exist_ok=True)
    s: dict[str, Path] = {}
    doc = Document()
    doc.add_heading("Отчёт", 1)
    doc.add_paragraph("Текст с кириллицей и жирным")
    doc.add_paragraph("пункт", style="List Bullet")
    t = doc.add_table(rows=3, cols=2)
    for i, (a, b) in enumerate([("Имя", "Сумма"), ("Анна", "10"), ("Борис", "20")]):
        t.cell(i, 0).text, t.cell(i, 1).text = a, b
    doc.save(folder / "sample.docx")
    s["docx"] = folder / "sample.docx"
    (folder / "sample.html").write_text(
        "<html><body><h1>Отчёт</h1><p>Текст <b>жирный</b></p><ul><li>пункт</li></ul><table><tr><th>Имя</th><th>Сумма"
        "</th></tr><tr><td>Анна</td><td>10</td></tr></table></body></html>", encoding="utf-8")
    (folder / "sample.md").write_text("# Отчёт\n\nТекст **жирный**\n\n- пункт\n\n| Имя | Сумма |\n|---|---|\n| Анна | 10 |\n",
                                      encoding="utf-8")
    (folder / "sample.txt").write_text("Отчёт\n\nТекст с кириллицей", encoding="utf-8")
    (folder / "sample.sql").write_text("SELECT 'привет';\nGO\n", encoding="utf-8")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Лист1"
    ws.append(["Имя", "Сумма"])
    ws.append(["Анна", 10])
    ws.append(["Борис", 20])
    w2 = wb.create_sheet("Лист2")
    w2.append(["x"])
    w2.append([1])
    wb.save(folder / "sample.xlsx")
    (folder / "sample.csv").write_text("Имя;Сумма\nАнна;10\nБорис;20\n", encoding="utf-8")
    (folder / "sample.tsv").write_text("Имя\tСумма\nАнна\t10\n", encoding="utf-8")
    (folder / "sample.json").write_text(json.dumps([{"Имя": "Анна", "Сумма": 10}, {"Имя": "Борис", "Сумма": 20}],
                                                   ensure_ascii=False), encoding="utf-8")
    for f in ("html", "md", "txt", "sql", "xlsx", "csv", "tsv", "json"):
        s[f] = folder / f"sample{PRIMARY_EXT[f]}"
    if soffice:
        convert(folder / "sample.docx", "odt", folder, soffice=soffice)
        convert(folder / "sample.docx", "rtf", folder, soffice=soffice)
        convert(folder / "sample.odt", "doc:MS Word 97", folder, soffice=soffice)
        convert(folder / "sample.xlsx", "xls:MS Excel 97", folder, soffice=soffice)
        for f in ("odt", "rtf", "doc", "xls"):
            s[f] = folder / f"sample{PRIMARY_EXT[f]}"
    return s
