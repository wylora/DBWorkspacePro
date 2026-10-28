import datetime as dt
import tempfile
import unittest
from pathlib import Path

from docx import Document

from bootstrap import create_services
from database.profile import ConnectionProfile
from reports.data import Dataset
from reports.document import BuildContext, DocKeyValues, DocParagraph, DocTable, build_document
from reports.generator import GenerationRequest, generate_report
from reports.model import Block, DataSource, Filter, Param, ReportTemplate, validate_template
from reports.quick import FAV_PREFIX, appendix_blocks, template_from_query
from security.credentials import InMemoryCredentialStore
from tests.fakes import FakeServer

SQL = "SELECT Region AS Регион, SUM(Total) AS Сумма\nFROM dbo.Orders\nWHERE Year = :year\nGROUP BY Region"


def ds(rows=None, **kw):
    kw.setdefault("sql", "SELECT ...")
    return Dataset("q", "Продажи", ["Регион", "Сумма"], rows if rows is not None else [("Север", 1500.5), ("Юг", 900)],
                   elapsed=0.25, **kw)


def tpl(block=None, **src):
    source = DataSource(id="q", name="Продажи", kind="sql", sql=SQL)
    for k, v in src.items():
        setattr(source, k, v)
    return ReportTemplate(name="T", title_page=False, toc=False, params=[Param("year", "Год", "int")], sources=[source],
                          blocks=[block or Block(id="b", type="query", source="q")])


def texts(doc):
    return [e for e in doc.elements]


class QueryBlockTests(unittest.TestCase):
    def build(self, t, data=None, **ctx):
        return build_document(BuildContext(t, {"q": data or ds()}, param_values={"year": 2026}, **ctx))

    def test_code_params_and_result(self):
        doc = self.build(tpl())
        els = texts(doc)
        self.assertEqual(els[0].text, "Запрос «Продажи»")
        code = next(e for e in els if isinstance(e, DocParagraph) and e.style == "code")
        self.assertEqual(code.text, SQL)                                          # исходный текст с :year
        kv = next(e for e in els if isinstance(e, DocKeyValues))
        self.assertEqual(kv.items, [(":year", "2026")])                            # без разделителя тысяч
        note = next(e for e in els if isinstance(e, DocParagraph) and e.text.startswith("Результат:"))
        self.assertIn("строк — 2", note.text)
        self.assertIn("0.25 с", note.text)
        table = next(e for e in els if isinstance(e, DocTable))
        self.assertEqual([c.title for c in table.columns], ["Регион", "Сумма"])
        self.assertEqual(table.rows[0][0].text, "Север")
        self.assertFalse(table.show_title)
        self.assertEqual(validate_template(tpl()), [])

    def test_switches_limits_and_empty_results(self):
        no_code = self.build(tpl(Block(id="b", type="query", source="q", show_code=False)))
        self.assertFalse(any(isinstance(e, DocParagraph) and e.style == "code" for e in no_code.elements))
        self.assertTrue(any(isinstance(e, DocTable) for e in no_code.elements))
        no_res = self.build(tpl(Block(id="b", type="query", source="q", show_result=False)))
        self.assertFalse(any(isinstance(e, DocTable) for e in no_res.elements))
        limited = self.build(tpl(Block(id="b", type="query", source="q", max_rows=1)), ds(truncated=True))
        note = next(e.text for e in limited.elements if isinstance(e, DocParagraph) and e.text.startswith("Результат:"))
        self.assertIn("показаны первые 1", note)
        self.assertIn("только часть строк", note)
        self.assertEqual(len(next(e for e in limited.elements if isinstance(e, DocTable)).rows), 1)
        empty = self.build(tpl(), Dataset("q", "Продажи", [], [], elapsed=0.01))
        self.assertTrue(any(isinstance(e, DocParagraph) and "не вернул набор строк" in e.text for e in empty.elements))

    def test_table_source_shows_generated_sql_and_values(self):
        t = tpl(kind="table", sql="", schema="dbo", table="Orders", columns=["Region"], filters=[Filter("Year", "=", ":year")])
        doc = self.build(t, ds(sql="SELECT TOP (5000) [Region] FROM [dbo].[Orders] WHERE [Year] = ?"))
        code = next(e.text for e in doc.elements if isinstance(e, DocParagraph) and e.style == "code")
        self.assertIn("FROM [dbo].[Orders]", code)
        kv = next(e for e in doc.elements if isinstance(e, DocKeyValues))
        self.assertEqual(kv.items, [("? № 1", "2026")])

    def test_secrets_in_code_are_masked(self):
        t = tpl()
        t.sources[0].sql = "SELECT 1 -- connect: PWD=hunter2"
        code = next(e.text for e in self.build(t).elements if isinstance(e, DocParagraph) and e.style == "code")
        self.assertNotIn("hunter2", code)

    def test_appendix_blocks_and_json_roundtrip(self):
        t = tpl(Block(id="b", type="table", source="q"))
        t.sources.append(DataSource(id="z", name="Второй", kind="sql", sql="SELECT 1 AS a"))
        extra = appendix_blocks(t)
        self.assertEqual([b.type for b in extra], ["heading", "query", "query"])
        self.assertEqual([b.source for b in extra[1:]], ["q", "z"])
        b = Block(id="b", type="query", source="q", show_code=False, max_rows=7)
        again = ReportTemplate.from_dict(tpl(b).to_dict()).blocks[0]
        self.assertEqual((again.type, again.show_code, again.show_result, again.max_rows), ("query", False, True, 7))
        old = tpl(Block(id="b", type="table", source="q")).to_dict()
        for blk in old["blocks"]:
            blk.pop("show_code", None)
            blk.pop("show_result", None)
        self.assertTrue(ReportTemplate.from_dict(old).blocks[0].show_code)       # старые шаблоны читаются


class QuickTemplateTests(unittest.TestCase):
    def test_template_from_saved_query(self):
        t = template_from_query("Продажи", SQL, key="7")
        self.assertEqual(t.id, f"{FAV_PREFIX}7")
        self.assertEqual([p.name for p in t.params], ["year"])
        self.assertEqual([b.type for b in t.blocks], ["query", "notes"])
        self.assertEqual(validate_template(t), [])
        self.assertTrue(any("изменяет" in e for e in validate_template(template_from_query("x", "DELETE FROM t WHERE id=1"))))
        self.assertTrue(validate_template(template_from_query("x", "SELECT 1\nGO\nSELECT 2")))

    def test_end_to_end_docx_contains_code_and_output(self):
        with tempfile.TemporaryDirectory() as d:
            server = FakeServer()
            server.when("FROM dbo.Orders", columns=["Регион", "Сумма"], rows=[("Север", 1500.5), ("Юг", 900)])
            svc = create_services(Path(d) / "h", InMemoryCredentialStore(), server.make_connection)
            p = svc.profiles.save(ConnectionProfile.new("P", "srv"))
            req = GenerationRequest(template_from_query("Продажи", SQL), {"year": "2026"}, ["docx", "html"], Path(d) / "out",
                                    lambda: svc.profiles.lease(p.id, "r"), base_name="q")
            res = generate_report(req)
            self.assertEqual(res.errors, {})
            doc = Document(str(Path(d) / "out" / "q.docx"))
            text = "\n".join(par.text for par in doc.paragraphs)
            self.assertIn("WHERE Year = :year", text)
            self.assertIn("строк — 2", text)
            cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
            self.assertIn("Север", cells)
            self.assertIn(":year", cells)
            html = (Path(d) / "out" / "q.html").read_text(encoding="utf-8")
            self.assertIn("GROUP BY Region", html)
            sql, params = next((s, pr) for s, pr in server.executed if "FROM dbo.Orders" in s)
            self.assertIn("Year = ?", sql)                                        # значение — параметром
            self.assertEqual(params, ("2026",))
            svc.shutdown()


if __name__ == "__main__":
    unittest.main()
