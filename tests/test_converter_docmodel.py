import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree

from converter import docmodel as dm


def kinds(blocks):
    return [type(b).__name__ for b in blocks]


class MarkdownParseTests(unittest.TestCase):
    def test_blocks(self):
        md = ("# Заголовок\nАбзац **жирный** *курсив* `код` [ссылка](https://example.com).\n\n- один\n  - вложенный\n- два\n\n"
              "1. первый\n2. второй\n\n> цитата\n\n```sql\nSELECT 1;\n```\n\n| a | b |\n|---|--:|\n| 1 | x\\|y |\n\n---\n")
        r = dm.parse_markdown(md)
        self.assertEqual(kinds(r.blocks), ["Heading", "Para", "ListItem", "ListItem", "ListItem", "ListItem", "ListItem",
                                           "Para", "Code", "TableB", "Rule"])
        para = r.blocks[1].inlines
        self.assertEqual([(i.text.strip(), i.bold, i.italic, i.code, i.url) for i in para if i.text.strip()][:4],
                         [("Абзац", False, False, False, ""), ("жирный", True, False, False, ""),
                          ("курсив", False, True, False, ""), ("код", False, False, True, "")])
        self.assertEqual(para[-2].url, "https://example.com")
        items = [b for b in r.blocks if isinstance(b, dm.ListItem)]
        self.assertEqual([(i.level, i.ordered, i.number) for i in items], [(0, False, 1), (1, False, 1), (0, False, 1),
                                                                           (0, True, 1), (0, True, 2)])
        self.assertTrue(r.blocks[7].quote)
        self.assertEqual((r.blocks[8].lang, r.blocks[8].text), ("sql", "SELECT 1;"))
        self.assertEqual(dm.plain(r.blocks[9].rows[1][1]), "x|y")

    def test_hard_break_setext_escapes_and_underscores(self):
        r = dm.parse_markdown("Строка один  \nстрока два\nмягкий перенос\n\nЗаголовок\n=====\n\nПодзаголовок\n---\n\n"
                              "snake_case_word и \\*не курсив\\* и 2*3*4\n")
        self.assertEqual(kinds(r.blocks), ["Para", "Heading", "Heading", "Para"])
        self.assertEqual(dm.plain(r.blocks[0].inlines), "Строка один\nстрока два мягкий перенос")
        self.assertEqual((r.blocks[1].level, r.blocks[2].level), (1, 2))
        text = dm.plain(r.blocks[3].inlines)
        self.assertIn("snake_case_word", text)
        self.assertIn("*не курсив*", text)
        self.assertFalse(any(i.italic for i in r.blocks[3].inlines if "snake" in i.text))

    def test_unsafe_links_are_dropped_images_become_text(self):
        r = dm.parse_markdown("[x](javascript:alert(1)) [ok](https://e.com/a_(b)) ![схема](pic.png) [m](mailto:a@b.c)")
        urls = [i.url for i in r.blocks[0].inlines if i.url]
        self.assertEqual(urls, ["https://e.com/a_(b)", "mailto:a@b.c"])
        self.assertIn("[изображение: схема]", dm.plain(r.blocks[0].inlines))
        self.assertNotIn("javascript", dm.to_html(r.blocks))
        self.assertNotIn("javascript", dm.to_markdown(r.blocks))

    def test_empty_and_degenerate_inputs(self):
        self.assertEqual(dm.parse_markdown("").blocks, [])
        self.assertEqual(dm.parse_markdown("\n\n  \n").blocks, [])
        self.assertEqual(kinds(dm.parse_markdown("| a |\n").blocks), ["Para"])        # таблица без разделителя — абзац
        self.assertEqual(kinds(dm.parse_markdown("```\nнезакрытый").blocks), ["Code"])
        dm.parse_markdown("**" * 500 + "_" * 500 + "[" * 300)                          # не зависает


class HtmlTextParseTests(unittest.TestCase):
    def test_html_structure(self):
        html = ("<html><head><title>T</title><style>p{}</style><script>alert(1)</script></head><body>"
                "<h1>Заголовок</h1><p>Текст <b>жирный</b> и <i>курсив</i> <a href='https://e.com'>ссылка</a><br>перенос</p>"
                "<ul><li>раз</li><li>два<ul><li>вложенный</li></ul></li></ul><ol start='3'><li>три</li></ol>"
                "<pre>a\n  b</pre><blockquote>цитата</blockquote><table><tr><th>A</th><th>B</th></tr>"
                "<tr><td>1</td><td>2 &amp; 3</td></tr></table><img src='x.png' alt='рис'><hr></body></html>")
        r = dm.parse_html(html)
        names = kinds(r.blocks)
        self.assertEqual(names[0], "Heading")
        self.assertNotIn("alert", " ".join(dm.plain(getattr(b, "inlines", [])) for b in r.blocks))
        para = next(b for b in r.blocks if isinstance(b, dm.Para) and not b.quote)
        self.assertEqual(dm.plain(para.inlines), "Текст жирный и курсив ссылка\nперенос")
        items = [b for b in r.blocks if isinstance(b, dm.ListItem)]
        self.assertEqual([(i.level, i.ordered, i.number) for i in items], [(0, False, 1), (0, False, 2), (1, False, 1), (0, True, 3)])
        self.assertEqual(next(b for b in r.blocks if isinstance(b, dm.Code)).text, "a\n  b")
        tb = next(b for b in r.blocks if isinstance(b, dm.TableB))
        self.assertEqual([[dm.plain(c) for c in row] for row in tb.rows], [["A", "B"], ["1", "2 & 3"]])
        self.assertTrue(any(isinstance(b, dm.Rule) for b in r.blocks))
        self.assertTrue(any("Изображений" in w for w in r.warnings))
        self.assertTrue(next(b for b in r.blocks if isinstance(b, dm.Para) and b.quote))

    def test_broken_html_does_not_crash(self):
        r = dm.parse_html("<div><p>не закрыт <b>жирный<table><tr><td>x<tr><td>y</div> текст")
        self.assertTrue(r.blocks)
        self.assertEqual(dm.parse_html("").blocks, [])

    def test_text_paragraphs(self):
        r = dm.parse_text("Первый\nабзац\n\n\nВторой")
        self.assertEqual([dm.plain(b.inlines) for b in r.blocks], ["Первый\nабзац", "Второй"])


class WritersTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)
        self.blocks = dm.parse_markdown("# Н\n\nТекст **ж** [с](https://e.com)\n\n- а\n- б\n\n1. в\n2. г\n\n| x | y |\n|---|---|\n| 1 | 2 |\n").blocks

    def tearDown(self):
        self.tmp.cleanup()

    def test_html_escapes_everything(self):
        evil = [dm.Para([dm.Inline("<script>alert(1)</script>", url="https://e.com/?a=1&b=\"2\"")]),
                dm.Code("<b>x</b>"), dm.TableB([[[dm.Inline("<img onerror=x>")]]])]
        html = dm.to_html(evil, "<T>")
        self.assertNotIn("<script", html)
        self.assertNotIn("<img", html)
        self.assertIn("&lt;script&gt;", html)
        self.assertIn("&amp;b=&quot;2&quot;", html)
        self.assertIn("&lt;T&gt;", html)
        parsed = dm.parse_html(html)
        self.assertEqual(kinds(parsed.blocks)[:1], ["Para"])

    def test_markdown_writer_escapes_and_roundtrips(self):
        md = dm.to_markdown(self.blocks)
        self.assertIn("**ж**", md)
        self.assertIn("[с](https://e.com)", md)
        self.assertIn("1. в\n2. г", md)
        self.assertIn("- а\n- б\n\n1. в", md)                              # между списками разного вида — пустая строка
        back = dm.parse_markdown(md)
        self.assertEqual(kinds(back.blocks), kinds(self.blocks))
        tricky = dm.to_markdown([dm.Para([dm.Inline("*не* _курсив_ [x](y) <b> `c` | #")])])
        again = dm.parse_markdown(tricky)
        self.assertEqual(dm.plain(again.blocks[0].inlines), "*не* _курсив_ [x](y) <b> `c` | #")
        self.assertFalse(any(i.italic or i.bold or i.url or i.code for i in again.blocks[0].inlines))

    def test_text_writer(self):
        txt = dm.to_text(self.blocks)
        self.assertIn("Н\n=", txt)
        self.assertIn("Текст ж с (https://e.com)", txt)
        self.assertIn("- а", txt)
        self.assertIn("1. в", txt)
        self.assertIn("x | y", txt)

    def test_docx_roundtrip_keeps_structure(self):
        path = self.d / "t.docx"
        dm.to_docx(self.blocks + dm.parse_markdown("> цитата\n\n```\ncode\n```\n").blocks, path, "Т")
        r = dm.read_docx(path)
        self.assertEqual(kinds(r.blocks), ["Heading", "Para", "ListItem", "ListItem", "ListItem", "ListItem", "TableB", "Para", "Code"])
        self.assertEqual([(i.ordered, i.number) for i in r.blocks if isinstance(i, dm.ListItem)],
                         [(False, 1), (False, 1), (True, 1), (True, 2)])
        link = next(i for i in r.blocks[1].inlines if i.url)
        self.assertEqual((link.text, link.url), ("с", "https://e.com"))
        self.assertTrue(r.blocks[7].quote)
        self.assertEqual(Document(str(path)).core_properties.title, "Т")

    def test_docx_writer_survives_hostile_text(self):
        path = self.d / "h.docx"
        blocks = [dm.Heading(1, [dm.Inline("З\x00аг\x0b")]), dm.Para([dm.Inline("Текст\x1f с\tтабом\nи переносом", True)]),
                  dm.Code("a\x00\n\tb"), dm.TableB([[[dm.Inline("я\x01чейка")], []]]),
                  dm.Para([dm.Inline("ссылка", url="https://e.com/?a=1&b=2")])]
        dm.to_docx(blocks, path)
        with zipfile.ZipFile(path) as z:
            etree.fromstring(z.read("word/document.xml"))
        r = dm.read_docx(path)
        self.assertEqual(dm.plain(r.blocks[0].inlines), "Заг")
        self.assertIn("Текст с", dm.plain(r.blocks[1].inlines))

    def test_docx_reader_warns_about_images_and_headers(self):
        from docx.shared import Cm
        d = Document()
        d.sections[0].header.paragraphs[0].text = "Верхний колонтитул"
        d.add_paragraph("текст")
        png = self.d / "p.png"
        from PIL import Image
        Image.new("RGB", (10, 10), "red").save(png)
        d.add_picture(str(png), width=Cm(1))
        d.save(self.d / "i.docx")
        r = dm.read_docx(self.d / "i.docx")
        self.assertTrue(any("Изображений" in w for w in r.warnings))
        self.assertTrue(any("Колонтитулы" in w for w in r.warnings))

    def test_extract_tables_names_and_header(self):
        blocks = dm.parse_markdown("# Продажи\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n| c |\n|---|\n| 3 |\n")
        tables = dm.extract_tables(blocks.blocks)
        self.assertEqual([(n, h, r) for n, h, r in tables], [("Продажи", ["a", "b"], [("1", "2")]), ("Продажи", ["c"], [("3",)])])
        self.assertEqual(dm.extract_tables(dm.parse_text("нет таблиц").blocks), [])


if __name__ == "__main__":
    unittest.main()
