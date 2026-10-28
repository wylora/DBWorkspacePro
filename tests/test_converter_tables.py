import datetime as dt
import json
import tempfile
import time
import unittest
from decimal import Decimal
from pathlib import Path

import openpyxl

from converter.libreoffice import ConversionCancelled
from converter.tables import (CsvReader, JsonReader, MemReader, TableError, detect_encoding, infer_cell, open_reader,
                              sniff_delimiter, write_tables)
from utils.tasks import CancelToken


class TableTestBase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def read(self, path, fmt):
        with open_reader(path, fmt) as r:
            return {n: (r.open(n).columns, list(r.open(n).rows)) for n in r.names}, r.warnings


class CsvTests(TableTestBase):
    def test_encodings(self):
        text = "Имя;Сумма\nАнна;1\n"
        for enc, guessed in (("utf-8", False), ("utf-8-sig", False), ("utf-16", False), ("cp1251", True)):
            p = self.d / f"{enc}.csv"
            p.write_bytes(text.encode(enc))
            self.assertEqual(detect_encoding(p), (enc if enc != "utf-8" else "utf-8", guessed) if enc != "utf-16" else ("utf-16", False))
            data, warns = self.read(p, "csv")
            self.assertEqual(data[p.stem][0], ["Имя", "Сумма"], enc)
            self.assertEqual(data[p.stem][1], [("Анна", "1")], enc)
            self.assertEqual(bool(warns), guessed, enc)

    def test_delimiters_and_quotes(self):
        for delim in (",", ";", "\t", "|"):
            p = self.d / "d.csv"
            p.write_text(f'a{delim}b{delim}c\n"x{delim}y"{delim}2{delim}"q""r"\n', encoding="utf-8")
            self.assertEqual(sniff_delimiter(p, "utf-8"), delim)
            (cols, rows), = self.read(p, "csv")[0].values()
            self.assertEqual((cols, rows), (["a", "b", "c"], [(f"x{delim}y", "2", 'q"r')]))

    def test_ragged_blank_header_and_embedded_newlines(self):
        p = self.d / "r.csv"
        p.write_text('a;;c\n\n1;2\n3;4;5;6\n"мульти\nстрока";x;y\n;;\n', encoding="utf-8")
        with open_reader(p, "csv") as r:
            sh = r.open("r")
            rows = list(sh.rows)
            self.assertEqual(sh.columns, ["a", "Столбец 2", "c"])
            self.assertEqual(rows, [("1", "2", ""), ("3", "4", "5"), ("мульти\nстрока", "x", "y")])
            self.assertEqual(r.skipped_long, 1)

    def test_empty_file_and_huge_cell(self):
        p = self.d / "e.csv"
        p.write_text("\n\n", encoding="utf-8")
        with open_reader(p, "csv") as r:
            self.assertEqual(r.open("e").columns, [])
        p.write_text("a\n" + "x" * 500_000 + "\n", encoding="utf-8")
        self.assertEqual(len(self.read(p, "csv")[0]["e"][1][0][0]), 500_000)

    def test_no_open_file_handles_left(self):
        import gc
        import warnings
        p = self.d / "h.csv"
        p.write_text("a;b\n1;2\n3;4\n", encoding="utf-8")
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            with open_reader(p, "csv") as r:
                r.open("h")                                    # лист открыт, но строки не прочитаны
                sh = r.open("h")
                next(sh.rows)                                  # строки читаются частично
                del sh
            gc.collect()
        leaks = [w for w in caught if issubclass(w.category, ResourceWarning)]
        self.assertEqual(leaks, [], [str(w.message) for w in leaks])
        p.unlink()                                             # на Windows открытый файл удалить нельзя

    def test_tsv_forces_tab(self):
        p = self.d / "t.tsv"
        p.write_text("a,b\tc\n1,2\t3\n", encoding="utf-8")
        self.assertEqual(self.read(p, "tsv")[0]["t"][0], ["a,b", "c"])

    def test_streaming_large_file(self):
        p = self.d / "big.csv"
        with open(p, "w", encoding="utf-8") as fh:
            fh.write("id;name;val\n")
            fh.writelines(f"{i};Имя {i};{i * 1.5}\n" for i in range(60_000))
        with open_reader(p, "csv") as r:
            sh = r.open("big")
            self.assertTrue(hasattr(sh.rows, "__next__"))                  # итератор, а не список
            started = time.time()
            res = write_tables(r, "xlsx", self.d, "big", infer=True)
            self.assertEqual(res.rows, 60_000)
        self.assertLess(time.time() - started, 60)
        wb = openpyxl.load_workbook(self.d / "big.xlsx", read_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
        self.assertEqual((len(rows), rows[1], rows[-1][0]), (60_001, (0, "Имя 0", 0), 59_999))


class InferTests(unittest.TestCase):
    def test_conservative_inference(self):
        cases = {"42": 42, "-7": -7, "3.14": Decimal("3.14"), "2,5": Decimal("2.5"), "2024-03-05": dt.date(2024, 3, 5),
                 "2024-03-05 10:30": dt.datetime(2024, 3, 5, 10, 30), "2024-03-05T10:30:15": dt.datetime(2024, 3, 5, 10, 30, 15)}
        for s, v in cases.items():
            self.assertEqual(infer_cell(s), v, s)
        for keep in ("00123", "0123", "+7 999", "1e5", "1 000", "12345678901234567890", "2024-13-45", "abc", "1.2.3", "=1+1"):
            self.assertEqual(infer_cell(keep), keep, keep)
        self.assertIsNone(infer_cell("   "))
        self.assertEqual(infer_cell(5), 5)


class XlsxJsonTests(TableTestBase):
    def make_book(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Продажи"
        ws.append(["Имя", "Сумма", "Дата", None])
        ws.append(["Анна", 10.5, dt.datetime(2024, 1, 2), None])
        ws.append([None, None, None, None])
        ws.append(["Борис", 20, dt.datetime(2024, 1, 3, 12, 30), None])
        ws.append(["=1+1", "=SUM(B2:B4)", None, None])                  # формулы без сохранённых значений
        ws.cell(row=7, column=1, value="=HYPERLINK(\"http://x\")").data_type = "s"      # строка, похожая на формулу
        ws.cell(row=7, column=2, value=5)
        e = wb.create_sheet("Пустой")
        w3 = wb.create_sheet("Данные 2")
        w3.append(["x"])
        w3.append([1])
        p = self.d / "b.xlsx"
        wb.save(p)
        return p

    def test_xlsx_read_trims_blank_columns_and_rows(self):
        data, warns = self.read(self.make_book(), "xlsx")
        self.assertTrue(any("2 ячеек с формулами без сохранённых значений" in w for w in warns))
        cols, rows = data["Продажи"]
        self.assertEqual(cols, ["Имя", "Сумма", "Дата"])
        self.assertEqual(rows[0], ("Анна", 10.5, dt.datetime(2024, 1, 2)))
        self.assertEqual(len(rows), 3)                       # пустая строка и строка с пустыми формулами пропущены
        self.assertEqual(data["Пустой"], ([], []))

    def test_xlsx_to_csv_per_sheet_with_safe_names_and_guard(self):
        with open_reader(self.make_book(), "xlsx") as r:
            res = write_tables(r, "csv", self.d, "out", delimiter=";", formula_guard=True)
            self.assertTrue(any("формулами без сохранённых" in w for w in r.warnings))
        names = sorted(f.name for f in res.files)
        self.assertEqual(names, ["out_Данные 2.csv", "out_Продажи.csv"])
        self.assertTrue(any("Лист «Пустой» пуст" in w for w in res.warnings))
        text = (self.d / "out_Продажи.csv").read_text(encoding="utf-8-sig")
        self.assertIn("Анна;10.5;2024-01-02", text.replace("\r\n", "\n"))
        self.assertIn("'=HYPERLINK", text)                                          # строка-«формула» обезврежена

    def test_xlsx_to_xlsx_style_multisheet_via_json_and_back(self):
        with open_reader(self.make_book(), "xlsx") as r:
            write_tables(r, "json", self.d, "j")
        self.assertEqual(json.loads((self.d / "j_Продажи.json").read_text(encoding="utf-8"))[0],
                         {"Имя": "Анна", "Сумма": 10.5, "Дата": "2024-01-02T00:00:00"})
        with open_reader(self.d / "j_Данные 2.json", "json") as r:
            res = write_tables(r, "xlsx", self.d, "back")
        self.assertEqual(openpyxl.load_workbook(res.files[0]).active.title, "j_Данные 2")

    def test_json_shapes(self):
        cases = {
            "list_of_dicts": ([{"a": 1, "b": {"x": 1}}, {"a": 2, "c": [1, 2]}], ["a", "b", "c"],
                              [(1, '{"x": 1}', None), (2, None, "[1, 2]")]),
            "list_of_lists": ([[1, 2], [3]], ["Столбец 1", "Столбец 2"], [(1, 2), (3, None)]),
            "single_object": ({"a": 1, "b": "x"}, ["a", "b"], [(1, "x")]),
        }
        for name, (data, cols, rows) in cases.items():
            p = self.d / f"{name}.json"
            p.write_text(json.dumps(data), encoding="utf-8")
            got = self.read(p, "json")[0]
            self.assertEqual(got[name], (cols, rows), name)
        p = self.d / "multi.json"
        p.write_text(json.dumps({"A": [{"x": 1}], "B": [{"y": 2}]}), encoding="utf-8")
        self.assertEqual(list(self.read(p, "json")[0]), ["A", "B"])

    def test_json_errors(self):
        for name, text, msg in (("bad.json", "{oops", "повреждён"), ("scalar.json", "42", "массив записей"),
                                ("mixed.json", '[{"a":1}, 5]', "либо объектами")):
            p = self.d / name
            p.write_text(text, encoding="utf-8")
            with self.assertRaises(TableError) as ctx:
                open_reader(p, "json")
            self.assertIn(msg, str(ctx.exception), name)
        p = self.d / "big.json"
        p.write_text("[]", encoding="utf-8")
        with self.assertRaises(TableError) as ctx:
            JsonReader(p, max_mb=0)
        self.assertIn("слишком большой", str(ctx.exception))
        (self.d / "latin.json").write_bytes('[{"а": "б"}]'.encode("cp1251"))
        with self.assertRaises(TableError):
            open_reader(self.d / "latin.json", "json")

    def test_corrupt_xlsx(self):
        p = self.d / "x.xlsx"
        p.write_bytes(b"PK\x03\x04garbage")
        with self.assertRaises(TableError):
            open_reader(p, "xlsx")


class WriteTests(TableTestBase):
    def test_nothing_to_write_and_unique_sheet_names(self):
        with self.assertRaises(TableError):
            write_tables(MemReader([("Пусто", [], [])]), "csv", self.d, "x")
        mem = MemReader([("Т", ["a"], [(1,)]), ("Т", ["a"], [(2,)])])
        self.assertEqual(mem.names, ["Т", "Т (2)"])
        res = write_tables(mem, "xlsx", self.d, "m")
        self.assertEqual(openpyxl.load_workbook(res.files[0]).sheetnames, ["Т", "Т (2)"])

    def test_cancel_removes_partial_files(self):
        tok = CancelToken()
        rows = [(i,) for i in range(10_000)]

        class Reader(MemReader):
            def open(self, name):
                sh = super().open(name)
                original = sh.rows

                def gen():
                    for n, r in enumerate(original):
                        if n == 3000:
                            tok.cancel()
                        yield r
                sh.rows = gen()
                return sh

        with self.assertRaises(ConversionCancelled):
            write_tables(Reader([("a", ["x"], rows)]), "csv", self.d, "c", cancel_token=tok)
        self.assertEqual(list(self.d.iterdir()), [])

    def test_tsv_and_delimiter(self):
        mem = lambda: MemReader([("a", ["x", "y"], [("1;2", "3")])])
        res = write_tables(mem(), "tsv", self.d, "t")
        self.assertEqual(res.files[0].read_bytes().decode("utf-8-sig"), "x\ty\r\n1;2\t3\r\n")
        res = write_tables(mem(), "csv", self.d, "c", delimiter=",")
        self.assertEqual(res.files[0].read_bytes().decode("utf-8-sig"), 'x,y\r\n1;2,3\r\n')


if __name__ == "__main__":
    unittest.main()
