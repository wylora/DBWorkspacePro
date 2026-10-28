import hashlib
import os
import shutil
import subprocess
import tempfile
import threading
import unittest
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest import mock

import openpyxl
from docx import Document

from converter import docmodel as dm
from converter.engine import (BatchSummary, ConvertOptions, Job, JobResult, choose_output_stem, convert_file,
                              expand_paths, run_batch)
from converter.formats import PRIMARY_EXT, SOURCE_FORMATS, TARGET_FORMATS, plan
from converter.libreoffice import find_soffice
from converter.verify import VerifyError, verify_output
from tests.converter_samples import build_samples
from utils.files import InsufficientDiskSpaceError
from utils.tasks import CancelToken

SOFFICE = find_soffice()
HAVE_PDFTEXT = shutil.which("pdftotext") is not None
SLOW = os.environ.get("DBWP_SKIP_SLOW") == "1"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class Base(unittest.TestCase):
    samples_dir = None
    samples: dict = {}

    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        cls.root = Path(cls._tmp.name)
        cls.samples = build_samples(cls.root / "samples", SOFFICE)

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def setUp(self):
        self._t = tempfile.TemporaryDirectory()
        self.t = Path(self._t.name)
        self.out, self.work = self.t / "out", self.t / "work"
        self.work.mkdir()

    def tearDown(self):
        self._t.cleanup()

    def opts(self, **kw):
        base = dict(out_dir=self.out, temp_root=self.work, soffice_path=SOFFICE or "")
        base.update(kw)
        return ConvertOptions(**base)

    def conv(self, fmt, dst, **kw):
        return convert_file(Job(self.samples[fmt]), dst, self.opts(**kw))

    def ok(self, fmt, dst, **kw):
        r = self.conv(fmt, dst, **kw)
        self.assertEqual((r.status, r.error), ("ok", ""), (fmt, dst))
        return r


class ContentTests(Base):
    def test_markdown_and_docx(self):
        r = self.ok("md", "docx")
        blocks = dm.read_docx(r.outputs[0]).blocks
        self.assertEqual(blocks[0].level, 1)
        self.assertEqual(dm.plain(blocks[0].inlines), "Отчёт")
        self.assertTrue(any(isinstance(b, dm.TableB) for b in blocks))
        md = self.ok("docx", "md").outputs[0].read_text(encoding="utf-8")
        self.assertIn("# Отчёт", md)
        self.assertIn("| Анна | 10 |", md)
        self.assertIn("- пункт", md)

    def test_html_and_text_routes(self):
        md = self.ok("html", "md").outputs[0].read_text(encoding="utf-8")
        self.assertIn("**жирный**", md)
        self.assertIn("| Анна | 10 |", md)
        self.assertIn("Отчёт\n=", self.ok("docx", "txt").outputs[0].read_text(encoding="utf-8"))
        html = self.ok("md", "html").outputs[0].read_text(encoding="utf-8")
        self.assertIn("<h1>Отчёт</h1>", html)
        self.assertIn("<strong>жирный</strong>", html)

    def test_table_routes(self):
        r = self.ok("xlsx", "csv")
        self.assertEqual(sorted(p.name for p in r.outputs), ["sample_Лист1.csv", "sample_Лист2.csv"])
        self.assertTrue(any("несколько листов" in w for w in r.warnings))
        self.assertEqual((self.out / "sample_Лист1.csv").read_bytes().decode("utf-8-sig").splitlines(), ["Имя;Сумма", "Анна;10", "Борис;20"])
        data = self.ok("json", "xlsx").outputs[0]
        self.assertEqual([[c.value for c in row] for row in openpyxl.load_workbook(data).active.iter_rows()],
                         [["Имя", "Сумма"], ["Анна", 10], ["Борис", 20]])
        tsv = self.ok("csv", "tsv").outputs[0].read_bytes().decode("utf-8-sig")
        self.assertEqual(tsv, "Имя\tСумма\r\nАнна\t10\r\nБорис\t20\r\n")

    def test_cp1251_sources(self):
        src = self.t / "old.csv"
        src.write_bytes("Имя;Город\nАнна;Москва\n".encode("cp1251"))
        r = convert_file(Job(src), "xlsx", self.opts())
        self.assertEqual(r.status, "ok")
        self.assertTrue(any("cp1251" in w for w in r.warnings))
        self.assertEqual(openpyxl.load_workbook(r.outputs[0]).active["B2"].value, "Москва")
        txt = self.t / "old.txt"
        txt.write_bytes("Привет, мир".encode("cp1251"))
        r = convert_file(Job(txt), "docx", self.opts())
        self.assertEqual(Document(str(r.outputs[0])).paragraphs[0].text, "Привет, мир")

    def test_extract_tables_from_documents(self):
        for fmt in ("docx", "md", "html"):
            r = self.ok(fmt, "csv", out_dir=self.out / fmt)
            self.assertEqual(r.outputs[0].read_bytes().decode("utf-8-sig").splitlines()[0], "Имя;Сумма", fmt)
        no_tables = self.conv("txt", "csv")
        self.assertEqual(no_tables.status, "unsupported")
        plain = self.t / "p.md"
        plain.write_text("# Только текст", encoding="utf-8")
        r = convert_file(Job(plain), "csv", self.opts())
        self.assertEqual(r.status, "failed")
        self.assertIn("нет таблиц", r.error)

    def test_sql_txt_copy(self):
        r = self.ok("sql", "txt")
        self.assertEqual(r.outputs[0].read_text(encoding="utf-8"), "SELECT 'привет';\nGO\n")
        back = convert_file(Job(r.outputs[0]), "sql", self.opts(out_dir=self.out / "b"))
        self.assertEqual(back.outputs[0].read_text(encoding="utf-8"), "SELECT 'привет';\nGO\n")

    @unittest.skipUnless(SOFFICE, "LibreOffice не установлен")
    def test_libreoffice_routes_keep_content(self):
        pdf = self.ok("docx", "pdf").outputs[0]
        if HAVE_PDFTEXT:
            text = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True).stdout
            for needle in ("Отчёт", "Текст с кириллицей", "Анна", "Борис"):
                self.assertIn(needle, text)
        src = self.t / "cp.txt"
        src.write_bytes("Кириллица в cp1251".encode("cp1251"))
        pdf2 = convert_file(Job(src), "pdf", self.opts()).outputs[0]
        if HAVE_PDFTEXT:
            self.assertIn("Кириллица в cp1251", subprocess.run(["pdftotext", str(pdf2), "-"], capture_output=True, text=True).stdout)
        html = self.ok("odt", "html").outputs
        self.assertTrue(any(p.suffix == ".html" for p in html))

    @unittest.skipUnless(SOFFICE, "LibreOffice не установлен")
    def test_doc_output_keeps_tables(self):
        doc = self.ok("docx", "doc").outputs[0]                          # через ODT: прямой путь терял таблицы
        back = convert_file(Job(doc), "csv", self.opts(out_dir=self.out / "back"))
        self.assertEqual(back.status, "ok", back.error)
        self.assertEqual(back.outputs[0].read_bytes().decode("utf-8-sig").splitlines(), ["Имя;Сумма", "Анна;10", "Борис;20"])
        md = self.ok("md", "doc", out_dir=self.out / "md").outputs[0]
        self.assertEqual(convert_file(Job(md), "csv", self.opts(out_dir=self.out / "md2")).status, "ok")


class SafetyAndPolicyTests(Base):
    def test_policies(self):
        names = lambda: sorted(p.name for p in self.out.iterdir())
        for _ in range(3):
            self.ok("md", "html")
        self.assertEqual(names(), ["sample (2).html", "sample (3).html", "sample.html"])
        before = (self.out / "sample.html").read_bytes()
        (self.out / "sample.html").write_bytes(b"OLD")
        r = self.conv("md", "html", policy="skip")
        self.assertEqual(r.status, "skipped")
        self.assertEqual((self.out / "sample.html").read_bytes(), b"OLD")
        self.ok("md", "html", policy="overwrite")
        self.assertEqual((self.out / "sample.html").read_bytes(), before)
        self.assertEqual(len(names()), 3)

    def test_multifile_families_count_as_existing(self):
        self.ok("xlsx", "csv")
        r = self.conv("xlsx", "csv", policy="skip")
        self.assertEqual(r.status, "skipped")
        r2 = self.ok("xlsx", "csv")
        self.assertTrue(all("(2)" in p.name for p in r2.outputs))

    def test_source_is_never_overwritten(self):
        out = self.t / "same"
        out.mkdir()
        src = out / "a.md"
        src.write_text("# x", encoding="utf-8")
        self.assertEqual(choose_output_stem(out, "a", "md", "overwrite", src), ("a (2)", False))
        self.assertEqual(choose_output_stem(out, "a", "md", "skip", src), ("a", True))
        self.assertEqual(choose_output_stem(out, "b", "md", "overwrite", src), ("b", False))

    def test_sources_stay_byte_identical(self):
        before = {k: sha(p) for k, p in self.samples.items()}
        for fmt in ("docx", "md", "html", "xlsx", "csv", "json", "txt"):
            for dst in TARGET_FORMATS:
                if plan(fmt, dst) and not plan(fmt, dst).needs_lo:
                    self.conv(fmt, dst, out_dir=self.out / f"{fmt}-{dst}")
        self.assertEqual({k: sha(p) for k, p in self.samples.items()}, before)

    def test_same_dir_and_folder_structure(self):
        folder = self.t / "docs"
        (folder / "a" / "b").mkdir(parents=True)
        shutil.copy(self.samples["md"], folder / "top.md")
        shutil.copy(self.samples["md"], folder / "a" / "b" / "deep.md")
        jobs, ignored, _ = expand_paths([folder])
        run_batch(jobs, "html", self.opts())
        self.assertTrue((self.out / "top.html").exists())
        self.assertTrue((self.out / "a" / "b" / "deep.html").exists())
        run_batch(jobs, "txt", self.opts(same_dir=True))
        self.assertTrue((folder / "a" / "b" / "deep.txt").exists())

    def test_statuses_for_bad_inputs(self):
        pdf = self.t / "a.pdf"
        pdf.write_bytes(b"%PDF-1.4\n")
        empty = self.t / "e.docx"
        empty.write_bytes(b"")
        bad = self.t / "b.docx"
        bad.write_bytes(b"not a zip at all")
        cases = [(pdf, "docx", "unsupported", "не поддерживается как исходный"), (empty, "pdf", "failed", "пустой"),
                 (bad, "md", "failed", "повреждён"), (self.t / "nope.docx", "md", "failed", "не найден"),
                 (self.samples["sql"], "docx", "unsupported", "только в TXT"), (self.samples["xlsx"], "docx", "unsupported", "Таблицу")]
        for src, dst, status, text in cases:
            r = convert_file(Job(src), dst, self.opts())
            self.assertEqual(r.status, status, (src.name, dst, r.error))
            self.assertIn(text, r.error, (src.name, dst))
        self.assertFalse(self.out.exists() and any(self.out.iterdir()))

    def test_libreoffice_missing_only_affects_lo_routes(self):
        with mock.patch("converter.engine.find_soffice", return_value=None):
            r = self.conv("docx", "pdf")
            self.assertEqual(r.status, "failed")
            self.assertIn("LibreOffice не найден", r.error)
            r2 = self.conv("docx", "pdf", soffice_path="X:/no/soffice.exe")
            self.assertIn("Указанный путь", r2.error)
            self.assertEqual(self.conv("docx", "md").status, "ok")

    def test_temp_files_always_cleaned(self):
        self.ok("md", "docx")
        self.conv("sql", "docx")
        convert_file(Job(self.t / "missing.md"), "html", self.opts())
        with mock.patch("converter.engine._run_route", side_effect=RuntimeError("boom")):
            r = self.conv("md", "html")
        self.assertEqual((r.status, r.error), ("failed", "RuntimeError: boom"))
        self.assertEqual(list(self.work.iterdir()), [])

    def test_disk_full_and_partial_copy_rollback(self):
        with mock.patch("converter.engine.ensure_free_space", side_effect=InsufficientDiskSpaceError("Недостаточно места на диске")):
            r = self.conv("md", "html")
        self.assertEqual(r.status, "failed")
        self.assertIn("Недостаточно места", r.error)
        self.assertFalse(self.out.exists() and any(self.out.iterdir()))
        from utils.files import atomic_copy
        calls = []

        def flaky(src, dest):
            calls.append(dest)
            if len(calls) == 2:
                raise OSError(28, "No space left on device")
            return atomic_copy(src, dest)

        with mock.patch("converter.engine.atomic_copy", side_effect=flaky):
            r = self.conv("xlsx", "csv")                                # два файла: на втором сбой
        self.assertEqual(r.status, "failed")
        self.assertEqual(list(self.out.iterdir()), [])                  # первый файл откатан
        self.assertEqual(list(self.work.iterdir()), [])

    def test_corrupt_result_is_caught_by_quality_control(self):
        def fake(route, fmt, src, stem, stage, *a, **k):
            p = stage / f"{stem}.docx"
            p.write_bytes(b"PK\x03\x04 broken")
            return [p]
        with mock.patch("converter.engine._run_route", side_effect=fake):
            r = self.conv("md", "docx")
        self.assertEqual(r.status, "failed")
        self.assertIn("повреждён", r.error)
        self.assertFalse(self.out.exists() and any(self.out.iterdir()))

    def test_cancel_before_and_between(self):
        tok = CancelToken()
        tok.cancel()
        self.assertEqual(convert_file(Job(self.samples["md"]), "html", self.opts(), tok).status, "cancelled")
        self.assertFalse(self.out.exists() and any(self.out.iterdir()))


class VerifyTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.d = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def w(self, name, data):
        p = self.d / name
        p.write_bytes(data)
        return p

    def test_broken_outputs_are_rejected(self):
        cases = [("a.docx", b"hello", "docx"), ("b.xlsx", b"PK\x03\x04junk", "xlsx"), ("c.pdf", b"%PDF-1.4 cut", "pdf"),
                 ("d.pdf", b"<html>", "pdf"), ("e.odt", b"PK\x03\x04junk", "odt"), ("f.rtf", b"plain", "rtf"),
                 ("g.doc", b"PK....", "doc"), ("h.xls", b"PK....", "xls"), ("i.html", b"just text", "html"),
                 ("j.json", b"{oops", "json"), ("k.txt", b"\xff\xfe\xfa", "txt"), ("l.csv", "Привет".encode("cp1251"), "csv"),
                 ("m.md", b"", "md")]
        for name, data, fmt in cases:
            with self.assertRaises(VerifyError, msg=name):
                verify_output(self.w(name, data), fmt)
        with self.assertRaises(VerifyError):
            verify_output(self.d / "missing.pdf", "pdf")

    def test_large_valid_json_is_not_rejected(self):
        p = self.d / "big.json"
        row = '{"a": "' + "x" * 90 + '"}'
        p.write_text("[" + ",\n".join([row] * 230_000) + "]", encoding="utf-8")        # ≈ 22 МБ, больше старого порога чтения
        self.assertGreater(p.stat().st_size, 20 * 1048576)
        verify_output(p, "json")
        p.write_bytes(p.read_bytes()[:-1000])                                            # а обрезанный — по-прежнему ошибка
        with self.assertRaises(VerifyError):
            verify_output(p, "json")

    def test_good_outputs_pass(self):
        d = Document()
        d.add_paragraph("x")
        d.save(self.d / "g.docx")
        wb = openpyxl.Workbook()
        wb.save(self.d / "g.xlsx")
        self.w("g.pdf", b"%PDF-1.4\n" + b"x" * 100 + b"\n%%EOF\n")
        self.w("g.rtf", b"{\\rtf1 hi}")
        self.w("g.html", "<p>Привет</p>".encode("utf-8"))
        self.w("g.json", b'[{"a": 1}]')
        self.w("g.csv", "Имя;1\n".encode("utf-8-sig"))
        for name, fmt in (("g.docx", "docx"), ("g.xlsx", "xlsx"), ("g.rtf", "rtf"), ("g.html", "html"),
                          ("g.json", "json"), ("g.csv", "csv")):
            verify_output(self.d / name, fmt)
        try:
            import pypdf  # noqa: F401
        except ImportError:
            verify_output(self.d / "g.pdf", "pdf")


class ExpandAndBatchTests(Base):
    def test_expand_paths(self):
        root = self.t / "in"
        (root / "sub" / ".hidden").mkdir(parents=True)
        (root / ".git").mkdir()
        for name in ("a.docx", "b.MD", "~$a.docx", ".secret.md", "c.exe", "sub/d.csv", "sub/.hidden/e.md", ".git/f.md"):
            (root / name).write_bytes(b"x")
        jobs, ignored, truncated = expand_paths([root, root / "a.docx", self.t / "missing.txt"])
        self.assertEqual(sorted(str(j.src.relative_to(root)) for j in jobs), ["a.docx", "b.MD", "sub/d.csv"])
        self.assertEqual({str(j.rel_dir) for j in jobs}, {".", "sub"})
        self.assertFalse(truncated)
        reasons = {p.name: r for p, r in ignored}
        self.assertEqual(reasons["c.exe"], "формат не поддерживается")
        self.assertEqual(reasons["missing.txt"], "файл не найден")
        flat, _, _ = expand_paths([root], recursive=False)
        self.assertEqual(sorted(j.src.name for j in flat), ["a.docx", "b.MD"])
        capped, _, trunc = expand_paths([root], max_files=2)
        self.assertEqual((len(capped), trunc), (2, True))

    @unittest.skipUnless(hasattr(os, "symlink"), "нет символических ссылок")
    def test_symlink_loops_are_not_followed(self):
        root = self.t / "loop"
        (root / "a").mkdir(parents=True)
        (root / "a" / "x.md").write_text("# x", encoding="utf-8")
        try:
            os.symlink(root, root / "a" / "back", target_is_directory=True)
        except OSError:
            self.skipTest("нельзя создать ссылку")
        jobs, _, _ = expand_paths([root])
        self.assertEqual([j.src.name for j in jobs], ["x.md"])

    def make_jobs(self, n):
        files = []
        for i in range(n):
            p = self.t / "src" / f"f{i:02d}.md"
            p.parent.mkdir(exist_ok=True)
            p.write_text(f"# Файл {i}\n\nтекст", encoding="utf-8")
            files.append(Job(p))
        return files

    def test_batch_counts_progress_and_consistency(self):
        jobs = self.make_jobs(4) + [Job(self.samples["sql"]), Job(self.t / "missing.md")]
        seen, stages = [], []
        summary = run_batch(jobs, "docx", self.opts(), progress=lambda s, f: stages.append((s, f)),
                            on_result=lambda i, r: seen.append(i))
        self.assertEqual((summary.total, summary.ok, summary.failed, summary.skipped, summary.cancelled), (6, 4, 1, 1, 0))
        self.assertEqual(sorted(seen), list(range(6)))
        self.assertEqual(summary.check_consistency(), [])
        self.assertEqual(stages[0][0], "Файл 1 из 6: f00.md")
        self.assertEqual(stages[-1], ("Готово", 1.0))
        summary.results[0].outputs[0].unlink()
        self.assertTrue(any("отсутствует" in p for p in summary.check_consistency()))
        bad = BatchSummary(summary.results[:-1] + [])
        bad.results.append(JobResult(Path("x"), "docx", "weird"))
        self.assertTrue(any("не совпадает" in p for p in bad.check_consistency()))

    def test_batch_cancel_stops_remaining_files(self):
        tok = CancelToken()
        jobs = self.make_jobs(6)
        summary = run_batch(jobs, "docx", self.opts(), cancel_token=tok,
                            on_result=lambda i, r: tok.cancel() if i == 1 else None)
        self.assertEqual((summary.ok, summary.cancelled), (2, 4))
        self.assertEqual(len(list(self.out.iterdir())), 2)
        self.assertEqual(summary.check_consistency(), [])
        self.assertEqual(list(self.work.iterdir()), [])

    def test_parallel_workers_keep_order_and_results(self):
        jobs = self.make_jobs(12)
        summary = run_batch(jobs, "html", self.opts(), workers=3)
        self.assertEqual([r.src.name for r in summary.results], [j.src.name for j in jobs])
        self.assertEqual(summary.ok, 12)
        self.assertEqual(len(list(self.out.iterdir())), 12)

    def test_name_collisions_in_batch(self):
        a, b = self.t / "x" / "same.md", self.t / "y" / "same.md"
        for p in (a, b):
            p.parent.mkdir()
            p.write_text("# t", encoding="utf-8")
        summary = run_batch([Job(a), Job(b)], "html", self.opts())
        self.assertEqual(sorted(p.name for r in summary.results for p in r.outputs), ["same (2).html", "same.html"])

    @unittest.skipUnless(SOFFICE, "LibreOffice не установлен")
    def test_cancel_kills_libreoffice_and_leaves_nothing(self):
        import glob
        tok = CancelToken()
        before = set(glob.glob(os.path.join(tempfile.gettempdir(), "dbwp-lo-*")))
        threading.Timer(0.5, tok.cancel).start()
        r = convert_file(Job(self.samples["docx"]), "pdf", self.opts(), tok)
        self.assertEqual(r.status, "cancelled")
        self.assertFalse(self.out.exists() and any(self.out.iterdir()))
        self.assertEqual(list(self.work.iterdir()), [])
        self.assertEqual(set(glob.glob(os.path.join(tempfile.gettempdir(), "dbwp-lo-*"))), before)

    @unittest.skipUnless(SOFFICE, "LibreOffice не установлен")
    def test_libreoffice_timeout_is_reported(self):
        r = self.conv("docx", "pdf", timeout=0.05)
        self.assertEqual(r.status, "failed")
        self.assertIn("тайм-аут", r.error)

    def conv(self, fmt, dst, **kw):
        return convert_file(Job(self.samples[fmt]), dst, self.opts(**kw))


@unittest.skipUnless(SOFFICE and not SLOW, "нужен LibreOffice; отключается DBWP_SKIP_SLOW=1")
class FullMatrixTests(Base):
    """Каждая пара «исходный → целевой», которую обещает планировщик, должна реально работать."""

    def test_every_planned_route_converts_and_verifies(self):
        pairs = [(s, d) for s in SOURCE_FORMATS for d in TARGET_FORMATS if plan(s, d)]
        self.assertGreaterEqual(len(pairs), 100)

        def run(pair):
            s, d = pair
            return pair, convert_file(Job(self.samples[s]), d, self.opts(out_dir=self.out / f"{s}-{d}"))

        with ThreadPoolExecutor(max_workers=3) as ex:
            results = list(ex.map(run, pairs))
        failed = [(p, r.status, r.error) for p, r in results if r.status != "ok"]
        self.assertEqual(failed, [])
        for (s, d), r in results:
            self.assertTrue(r.outputs and all(o.is_file() and o.stat().st_size > 0 for o in r.outputs), (s, d))
            self.assertEqual(r.outputs[0].suffix, PRIMARY_EXT[d], (s, d))
        self.assertEqual(list(self.work.iterdir()), [])


if __name__ == "__main__":
    unittest.main()
