"""Упрощённая модель документа для преобразований Markdown / HTML / TXT / DOCX без внешних программ.

Сохраняются: заголовки, абзацы, списки, таблицы, блоки кода, цитаты, полужирный/курсив/код, ссылки.
Не сохраняются: изображения, колонтитулы, сноски, поля, точное оформление (об этом сообщают предупреждения).
"""
from __future__ import annotations

import html as _html
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from reports.formatting import clean_xml_text

SAFE_SCHEMES = ("http://", "https://", "mailto:", "#", "/", "./", "../")


@dataclass
class Inline:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False
    url: str = ""


@dataclass
class Heading:
    level: int
    inlines: list[Inline]


@dataclass
class Para:
    inlines: list[Inline]
    quote: bool = False


@dataclass
class ListItem:
    inlines: list[Inline]
    level: int = 0
    ordered: bool = False
    number: int = 1


@dataclass
class Code:
    text: str
    lang: str = ""


@dataclass
class TableB:
    rows: list[list[list[Inline]]]
    header: bool = True


@dataclass
class Rule:
    pass


Block = Heading | Para | ListItem | Code | TableB | Rule


@dataclass
class ReadResult:
    blocks: list
    warnings: list[str] = field(default_factory=list)


def safe_url(url: str) -> str:
    u = (url or "").strip()
    low = u.lower()
    return u if low.startswith(SAFE_SCHEMES) or ("://" not in low and ":" not in low.split("/")[0]) else ""


def plain(inlines: list[Inline]) -> str:
    return "".join(i.text for i in inlines)


def _trim(inlines: list[Inline]) -> list[Inline]:
    out = [Inline(i.text, i.bold, i.italic, i.code, i.url) for i in inlines if i.text != ""]
    while out and out[0].text.strip() == "" and "\n" not in out[0].text:
        out.pop(0)
    while out and out[-1].text.strip() == "":
        out.pop()
    if out:
        out[0].text = out[0].text.lstrip(" ")
        out[-1].text = out[-1].text.rstrip(" ")
    return out


# =============================================================== Markdown
_ESC_CHARS = "!\"#$%&'()*+,-./:;<=>?@[\\]^_`{|}~"
_ESC = {c: chr(0xE000 + i) for i, c in enumerate(_ESC_CHARS)}
_UNESC = {v: k for k, v in _ESC.items()}
_INLINE = re.compile(r"""
  (?P<code>(?P<fence>`+)(?P<ctext>.+?)(?<!`)(?P=fence)(?!`))
 |(?P<img>!\[(?P<alt>[^\]]*)\]\((?P<isrc>[^)\s]*)(?:\s+"[^"]*")?\))
 |(?P<link>\[(?P<ltext>[^\]]+)\]\((?P<lurl>(?:[^()\s]|\([^()\s]*\))*)(?:\s+"[^"]*")?\))
 |(?P<bold>(?P<bm>\*\*|__)(?=\S)(?P<btext>.+?)(?<=\S)(?P=bm))
 |(?P<ital>(?<![\w*])\*(?=[^\s*])(?P<itext>.+?)(?<=[^\s*])\*(?!\*)|(?<!\w)_(?=[^\s_])(?P<utext>.+?)(?<=[^\s_])_(?!\w))
""", re.X | re.S)


def _unescape(s: str) -> str:
    return "".join(_UNESC.get(c, c) for c in s)


def parse_inline(text: str, bold=False, italic=False, url="", _top=True) -> list[Inline]:
    if _top:
        text = re.sub(r"\\([!-/:-@\[-`{-~])", lambda m: _ESC[m.group(1)], text)
    out: list[Inline] = []
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            out.append(Inline(text[pos:m.start()], bold, italic, False, url))
        if m.group("code") is not None:
            out.append(Inline(m.group("ctext").strip(), bold, italic, True, url))
        elif m.group("img") is not None:
            alt = m.group("alt") or "изображение"
            out.append(Inline(f"[изображение: {alt}]", bold, italic, False, url))
        elif m.group("link") is not None:
            out.extend(parse_inline(m.group("ltext"), bold, italic, safe_url(m.group("lurl")), False))
        elif m.group("bold") is not None:
            out.extend(parse_inline(m.group("btext"), True, italic, url, False))
        else:
            out.extend(parse_inline(m.group("itext") or m.group("utext"), bold, True, url, False))
        pos = m.end()
    if pos < len(text):
        out.append(Inline(text[pos:], bold, italic, False, url))
    if _top:
        for i in out:
            i.text = _unescape(i.text)
            i.url = _unescape(i.url)
    return out


_HR = re.compile(r"^\s{0,3}([-*_])(\s*\1){2,}\s*$")
_ATX = re.compile(r"^\s{0,3}(#{1,6})\s+(.*?)\s*#*\s*$")
_LIST = re.compile(r"^(\s*)([-*+]|\d{1,9}[.)])\s+(.*)$")
_FENCE = re.compile(r"^\s{0,3}(```+|~~~+)\s*([\w+-]*)\s*$")
_TSEP = re.compile(r"^\s*\|?\s*:?-{1,}:?\s*(\|\s*:?-{1,}:?\s*)*\|?\s*$")


def _split_row(line: str) -> list[str]:
    s = line.strip()
    s = s[1:] if s.startswith("|") else s
    s = s[:-1] if s.endswith("|") and not s.endswith("\\|") else s
    return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", s)]


def parse_markdown(text: str) -> ReadResult:
    lines = text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", "    ").split("\n")
    blocks: list = []
    para: list[str] = []
    counters: dict[int, int] = {}

    def flush_para(quote=False):
        if para:
            joined = ""
            for k, ln in enumerate(para):
                if k:
                    prev = para[k - 1]
                    joined += "\n" if (prev.endswith("  ") or prev.endswith("\\")) else " "
                joined += ln.strip().rstrip("\\") if ln.endswith("\\") and not ln.endswith("\\\\") else ln.strip()
            blocks.append(Para(_trim(parse_inline(joined)), quote))
            para.clear()

    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        m = _FENCE.match(line)
        if m:
            flush_para()
            fence, lang = m.group(1), m.group(2)
            body = []
            i += 1
            while i < n and not re.match(rf"^\s{{0,3}}{re.escape(fence[0])}{{{len(fence)},}}\s*$", lines[i]):
                body.append(lines[i])
                i += 1
            blocks.append(Code("\n".join(body), lang))
            i += 1
            continue
        if not line.strip():
            flush_para()
            counters.clear()
            i += 1
            continue
        m = _ATX.match(line)
        if m:
            flush_para()
            blocks.append(Heading(len(m.group(1)), _trim(parse_inline(m.group(2)))))
            i += 1
            continue
        if para and len(para) == 1 and re.match(r"^\s{0,3}(=+|-+)\s*$", line):
            blocks.append(Heading(1 if "=" in line else 2, _trim(parse_inline(para[0].strip()))))
            para.clear()
            i += 1
            continue
        if _HR.match(line):
            flush_para()
            blocks.append(Rule())
            i += 1
            continue
        if line.lstrip().startswith(">"):
            flush_para()
            quote_lines = []
            while i < n and lines[i].lstrip().startswith(">"):
                quote_lines.append(re.sub(r"^\s{0,3}>\s?", "", lines[i]))
                i += 1
            inner = parse_markdown("\n".join(quote_lines)).blocks
            for b in inner:
                if isinstance(b, Para):
                    b.quote = True
                blocks.append(b)
            continue
        if "|" in line and i + 1 < n and _TSEP.match(lines[i + 1]) and "-" in lines[i + 1]:
            head = _split_row(line)
            sep = _split_row(lines[i + 1])
            if len(head) == len(sep):
                flush_para()
                rows = [[_trim(parse_inline(c)) for c in head]]
                i += 2
                while i < n and lines[i].strip() and "|" in lines[i]:
                    cells = _split_row(lines[i])
                    cells = (cells + [""] * len(head))[:len(head)]
                    rows.append([_trim(parse_inline(c)) for c in cells])
                    i += 1
                blocks.append(TableB(rows, True))
                continue
        m = _LIST.match(line)
        if m:
            flush_para()
            indent = len(m.group(1).replace("\t", "    "))
            level = min(indent // 2, 6) if indent < 4 else min(indent // 4 + (1 if indent % 4 else 0), 6)
            ordered = m.group(2)[0].isdigit()
            body = m.group(3)
            i += 1
            while i < n and lines[i].strip() and not _LIST.match(lines[i]) and lines[i].startswith(" ") \
                    and not _FENCE.match(lines[i]):
                body += " " + lines[i].strip()
                i += 1
            for deeper in [k for k in counters if k > level]:
                del counters[deeper]
            counters[level] = counters.get(level, 0) + 1 if ordered else 0
            num = counters[level] if ordered else 1
            if ordered and re.match(r"\d+", m.group(2)) and counters[level] == 1:
                counters[level] = num = int(re.match(r"\d+", m.group(2)).group())
            blocks.append(ListItem(_trim(parse_inline(body)), level, ordered, num))
            continue
        para.append(line)
        i += 1
    flush_para()
    return ReadResult(blocks)


# ==================================================================== HTML
_SKIP = {"script", "style", "head", "title", "noscript", "template", "svg", "iframe"}
_BLOCK_FLUSH = {"p", "div", "section", "article", "header", "footer", "nav", "aside", "main", "figure", "figcaption",
                "form", "address", "dl", "dt", "dd", "fieldset", "details", "summary"}


class _HtmlReader(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks: list = []
        self.inl: list[Inline] = []
        self.skip = 0
        self.bold = self.italic = self.code = 0
        self.links: list[str] = []
        self.heading = 0
        self.lists: list[list] = []        # [ordered, counter]
        self.li = None                     # (level, ordered, number) при открытом <li>
        self.pre: list[str] | None = None
        self.quote = 0
        self.table: list | None = None
        self.row: list | None = None
        self.cell: list[Inline] | None = None
        self.images = 0

    # --- вспомогательное
    def _add(self, text: str):
        if text:
            self.inl.append(Inline(text, bool(self.bold), bool(self.italic), bool(self.code), self.links[-1] if self.links else ""))

    def _flush(self):
        items = _trim(self.inl)
        self.inl = []
        if self.cell is not None:
            return
        if not items:
            return
        if self.heading:
            self.blocks.append(Heading(self.heading, items))
        elif self.li is not None:
            lvl, ordered, num = self.li
            self.blocks.append(ListItem(items, lvl, ordered, num))
            self.li = (lvl, ordered, num)
        else:
            self.blocks.append(Para(items, self.quote > 0))

    # --- события
    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        if tag in _SKIP:
            self.skip += 1
            return
        if self.skip:
            return
        if tag in _BLOCK_FLUSH:
            self._flush()
        elif re.fullmatch(r"h[1-6]", tag):
            self._flush()
            self.heading = int(tag[1])
        elif tag == "br":
            self._add("\n")
        elif tag == "hr":
            self._flush()
            self.blocks.append(Rule())
        elif tag in ("ul", "ol"):
            self._flush()
            self.lists.append([tag == "ol", int(a.get("start", 1) or 1) - 1 if a.get("start", "1").isdigit() else 0])
        elif tag == "li":
            self._flush()
            if not self.lists:
                self.lists.append([False, 0])
            self.lists[-1][1] += 1
            self.li = (len(self.lists) - 1, self.lists[-1][0], self.lists[-1][1])
        elif tag == "pre":
            self._flush()
            self.pre = []
        elif tag == "blockquote":
            self._flush()
            self.quote += 1
        elif tag == "table":
            self._flush()
            self.table = []
        elif tag == "tr" and self.table is not None:
            self.row = []
        elif tag in ("td", "th") and self.row is not None:
            self.cell = []
            self.inl = []
        elif tag in ("b", "strong"):
            self.bold += 1
        elif tag in ("i", "em", "cite"):
            self.italic += 1
        elif tag in ("code", "kbd", "samp", "tt") and self.pre is None:
            self.code += 1
        elif tag == "a":
            self.links.append(safe_url(a.get("href", "")))
        elif tag == "img":
            self.images += 1
            self._add(f"[изображение: {a.get('alt') or 'без подписи'}]")

    def handle_endtag(self, tag):
        if tag in _SKIP:
            self.skip = max(0, self.skip - 1)
            return
        if self.skip:
            return
        if tag in _BLOCK_FLUSH or re.fullmatch(r"h[1-6]", tag):
            self._flush()
            if tag.startswith("h") and len(tag) == 2:
                self.heading = 0
        elif tag == "li":
            self._flush()
            self.li = None
        elif tag in ("ul", "ol"):
            self._flush()
            if self.lists:
                self.lists.pop()
            self.li = None
        elif tag == "pre" and self.pre is not None:
            self.blocks.append(Code("".join(self.pre).strip("\n")))
            self.pre = None
        elif tag == "blockquote":
            self._flush()
            self.quote = max(0, self.quote - 1)
        elif tag in ("td", "th") and self.cell is not None and self.row is not None:
            self.row.append(_trim(self.inl))
            self.inl = []
            self.cell = None
        elif tag == "tr" and self.row is not None and self.table is not None:
            if self.row:
                self.table.append(self.row)
            self.row = None
        elif tag == "table" and self.table is not None:
            if self.table:
                width = max(len(r) for r in self.table)
                self.blocks.append(TableB([r + [[] for _ in range(width - len(r))] for r in self.table], True))
            self.table = None
        elif tag in ("b", "strong"):
            self.bold = max(0, self.bold - 1)
        elif tag in ("i", "em", "cite"):
            self.italic = max(0, self.italic - 1)
        elif tag in ("code", "kbd", "samp", "tt"):
            self.code = max(0, self.code - 1)
        elif tag == "a" and self.links:
            self.links.pop()
        elif tag == "p":
            self._flush()

    def handle_data(self, data):
        if self.skip:
            return
        if self.pre is not None:
            self.pre.append(data)
        else:
            self._add(re.sub(r"\s+", " ", data))

    def close(self):
        super().close()
        self._flush()


def parse_html(text: str) -> ReadResult:
    r = _HtmlReader()
    r.feed(text)
    r.close()
    warns = []
    if r.images:
        warns.append(f"Изображений в документе: {r.images} — они не переносятся (заменены текстовой пометкой).")
    return ReadResult(r.blocks, warns)


def parse_text(text: str) -> ReadResult:
    blocks = []
    for chunk in re.split(r"\n\s*\n", text.replace("\r\n", "\n").replace("\r", "\n")):
        if chunk.strip():
            blocks.append(Para([Inline(chunk.strip("\n"))]))
    return ReadResult(blocks)


# ==================================================================== DOCX (чтение)
def _numbering_is_ordered(doc, num_id: str, ilvl: str) -> bool:
    try:
        root = doc.part.numbering_part.element
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        num = root.xpath(f"w:num[@w:numId='{num_id}']", namespaces=ns)[0]
        abstract = num.xpath("w:abstractNumId/@w:val", namespaces=ns)[0]
        fmt = root.xpath(f"w:abstractNum[@w:abstractNumId='{abstract}']/w:lvl[@w:ilvl='{ilvl}']/w:numFmt/@w:val",
                         namespaces=ns)
        return bool(fmt) and fmt[0] != "bullet"
    except Exception:  # noqa: BLE001
        return False


def _para_inlines(p) -> list[Inline]:
    out: list[Inline] = []
    try:
        content = list(p.iter_inner_content())
    except AttributeError:
        content = list(p.runs)
    for item in content:
        if item.__class__.__name__ == "Hyperlink":
            url = safe_url(item.url or "")
            for r in item.runs:
                out.append(Inline(r.text, bool(r.bold), bool(r.italic), False, url))
            continue
        name = (item.font.name or "") if hasattr(item, "font") else ""
        out.append(Inline(item.text, bool(item.bold), bool(item.italic), name.lower() in ("consolas", "courier new")))
    return _trim(out)


def read_docx(path: str | Path) -> ReadResult:
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = Document(str(path))
    blocks: list = []
    warns: list[str] = []
    images = 0
    counters: dict[int, int] = {}
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:p"):
            p = Paragraph(el, doc)
            if el.xpath(".//w:drawing | .//w:pict"):
                images += 1
            inl = _para_inlines(p)
            if not inl:
                counters.clear()
                continue
            try:
                style = (p.style.name or "") if p.style is not None else ""
            except Exception:  # noqa: BLE001
                style = ""
            m = re.match(r"^(?:heading|заголовок)\s*(\d)", style, re.I)
            if style.lower() in ("title", "название"):
                blocks.append(Heading(1, inl))
            elif m:
                blocks.append(Heading(min(6, int(m.group(1))), inl))
            else:
                ppr = el.pPr
                numpr = ppr.numPr if ppr is not None else None
                is_list = numpr is not None or style.lower().startswith(("list bullet", "list number", "список"))
                if is_list:
                    ilvl = str(numpr.ilvl.val) if numpr is not None and numpr.ilvl is not None else "0"
                    num_id = str(numpr.numId.val) if numpr is not None and numpr.numId is not None else ""
                    lvl = int(ilvl) if ilvl.isdigit() else 0
                    sm = re.search(r"(\d)$", style)
                    if sm and numpr is None:
                        lvl = max(0, int(sm.group(1)) - 1)
                    ordered = "number" in style.lower() or (num_id != "" and _numbering_is_ordered(doc, num_id, ilvl))
                    for k in [k for k in counters if k > lvl]:
                        del counters[k]
                    counters[lvl] = counters.get(lvl, 0) + 1
                    blocks.append(ListItem(inl, lvl, ordered, counters[lvl] if ordered else 1))
                else:
                    pf = p.paragraph_format
                    mnum = re.match(r"^(\d{1,4})[.)]\s+", plain(inl))
                    if mnum and pf.first_line_indent is not None and pf.first_line_indent < 0:
                        lvl = max(0, round(((pf.left_indent.cm if pf.left_indent is not None else 0.9) - 0.9) / 0.8))
                        first = inl[0]
                        first.text = first.text[mnum.end():] if len(plain([first])) >= mnum.end() else first.text
                        blocks.append(ListItem(_trim(inl), lvl, True, int(mnum.group(1))))
                        continue
                    counters.clear()
                    quote = style.lower() in ("quote", "intense quote", "цитата")
                    mono = all(i.code for i in inl)
                    blocks.append(Code(plain(inl)) if mono and not quote else Para(inl, quote))
        elif el.tag == qn("w:tbl"):
            counters.clear()
            t = Table(el, doc)
            rows = []
            for r in t.rows:
                row = []
                seen = set()
                for c in r.cells:
                    key = id(c._tc)
                    row.append([Inline("\n".join(pp.text for pp in c.paragraphs).strip())] if key not in seen else [])
                    seen.add(key)
                rows.append(row)
            if rows:
                width = max(len(r) for r in rows)
                blocks.append(TableB([r + [[] for _ in range(width - len(r))] for r in rows], True))
    if images:
        warns.append(f"Изображений в документе: {images} — они не переносятся.")
    try:
        for sec in doc.sections:
            if any(pp.text.strip() for pp in sec.header.paragraphs + sec.footer.paragraphs):
                warns.append("Колонтитулы не переносятся.")
                break
    except Exception:  # noqa: BLE001
        pass
    return ReadResult(blocks, warns)


# ==================================================================== запись
def _md_text(s: str) -> str:
    return re.sub(r"([\\`*_\[\]<>|])", r"\\\1", s).replace("\n", "  \n")


def _md_inline(inls: list[Inline]) -> str:
    out = []
    for i in inls:
        if not i.text:
            continue
        if i.code:
            t = "`" + i.text.replace("`", "'") + "`"
        else:
            t = _md_text(i.text)
            if i.bold and i.italic:
                t = f"***{t}***"
            elif i.bold:
                t = f"**{t}**"
            elif i.italic:
                t = f"*{t}*"
        if i.url:
            t = f"[{t}]({i.url})"
        out.append(t)
    return "".join(out)


def to_markdown(blocks: list) -> str:
    out: list[str] = []
    prev_list: ListItem | None = None
    for b in blocks:
        is_list = isinstance(b, ListItem)
        if out and not (is_list and prev_list is not None and prev_list.ordered == b.ordered):
            out.append("")
        prev_list = b if is_list else None
        if isinstance(b, Heading):
            out.append("#" * b.level + " " + _md_inline(b.inlines).replace("  \n", " "))
        elif isinstance(b, Para):
            txt = _md_inline(b.inlines)
            out.append("\n".join(("> " + ln) for ln in txt.split("\n")) if b.quote else txt)
        elif isinstance(b, ListItem):
            mark = f"{b.number}." if b.ordered else "-"
            out.append("  " * b.level + f"{mark} " + _md_inline(b.inlines).replace("  \n", " "))
        elif isinstance(b, Code):
            fence = "```" if "```" not in b.text else "~~~~"
            out.append(f"{fence}{b.lang}\n{b.text}\n{fence}")
        elif isinstance(b, Rule):
            out.append("---")
        elif isinstance(b, TableB):
            def cell(c):
                return _md_inline(c).replace("\n", "<br>").replace("  <br>", "<br>") or " "
            head, *body = b.rows
            out.append("| " + " | ".join(cell(c) for c in head) + " |")
            out.append("| " + " | ".join("---" for _ in head) + " |")
            out += ["| " + " | ".join(cell(c) for c in r) + " |" for r in body]
    return "\n".join(out).strip() + "\n"


def to_text(blocks: list) -> str:
    def tl(inls):
        parts = []
        for i in inls:
            parts.append(i.text if not i.url or i.url == i.text else f"{i.text} ({i.url})")
        return "".join(parts)

    out: list[str] = []
    prev_list: ListItem | None = None
    for b in blocks:
        is_list = isinstance(b, ListItem)
        if out and not (is_list and prev_list is not None and prev_list.ordered == b.ordered):
            out.append("")
        prev_list = b if is_list else None
        if isinstance(b, Heading):
            text = tl(b.inlines).replace("\n", " ")
            out.append(text)
            if b.level <= 2:
                out.append(("=" if b.level == 1 else "-") * min(len(text), 80))
        elif isinstance(b, Para):
            out.append(("    " if b.quote else "") + tl(b.inlines))
        elif isinstance(b, ListItem):
            out.append("  " * b.level + (f"{b.number}. " if b.ordered else "- ") + tl(b.inlines).replace("\n", " "))
        elif isinstance(b, Code):
            out.append(b.text)
        elif isinstance(b, Rule):
            out.append("-" * 40)
        elif isinstance(b, TableB):
            out += [" | ".join(plain(c).replace("\n", " ") for c in r) for r in b.rows]
    return "\n".join(out).strip() + "\n"


def to_html(blocks: list, title: str = "Документ") -> str:
    esc = _html.escape

    def inl(items):
        parts = []
        for i in items:
            t = esc(i.text).replace("\n", "<br>")
            if i.code:
                t = f"<code>{t}</code>"
            if i.bold:
                t = f"<strong>{t}</strong>"
            if i.italic:
                t = f"<em>{t}</em>"
            if i.url:
                t = f'<a href="{esc(i.url, quote=True)}">{t}</a>'
            parts.append(t)
        return "".join(parts)

    body: list[str] = []
    stack: list[str] = []

    def close_lists(to=0):
        while len(stack) > to:
            body.append(f"</{stack.pop()}>")

    for b in blocks:
        if isinstance(b, ListItem):
            tag = "ol" if b.ordered else "ul"
            while len(stack) > b.level + 1:
                body.append(f"</{stack.pop()}>")
            while len(stack) < b.level + 1:
                stack.append(tag)
                body.append(f"<{tag}>")
            if stack[-1] != tag:
                body.append(f"</{stack.pop()}>")
                stack.append(tag)
                body.append(f"<{tag}>")
            body.append(f"<li>{inl(b.inlines)}</li>")
            continue
        close_lists()
        if isinstance(b, Heading):
            body.append(f"<h{b.level}>{inl(b.inlines)}</h{b.level}>")
        elif isinstance(b, Para):
            body.append(f"<blockquote><p>{inl(b.inlines)}</p></blockquote>" if b.quote else f"<p>{inl(b.inlines)}</p>")
        elif isinstance(b, Code):
            body.append(f"<pre><code>{esc(b.text)}</code></pre>")
        elif isinstance(b, Rule):
            body.append("<hr>")
        elif isinstance(b, TableB):
            rows = []
            for ri, r in enumerate(b.rows):
                tag = "th" if ri == 0 and b.header else "td"
                rows.append("<tr>" + "".join(f"<{tag}>{inl(c)}</{tag}>" for c in r) + "</tr>")
            body.append("<table>\n" + "\n".join(rows) + "\n</table>")
    close_lists()
    css = ("body{font-family:Segoe UI,Arial,sans-serif;max-width:900px;margin:2em auto;padding:0 1em;line-height:1.55}"
           "table{border-collapse:collapse}th,td{border:1px solid #bbb;padding:4px 8px;vertical-align:top}"
           "th{background:#f0f0f0}pre{background:#f5f5f5;padding:10px;overflow:auto}code{font-family:Consolas,monospace}"
           "blockquote{border-left:3px solid #ccc;margin-left:0;padding-left:1em;color:#555}")
    return (f'<!DOCTYPE html>\n<html lang="ru"><head><meta charset="utf-8"><title>{esc(title)}</title>'
            f"<style>{css}</style></head><body>\n" + "\n".join(body) + "\n</body></html>\n")


def _add_hyperlink(paragraph, text: str, url: str, bold: bool, italic: bool) -> None:
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    if bold:
        rpr.append(OxmlElement("w:b"))
    if italic:
        rpr.append(OxmlElement("w:i"))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = clean_xml_text(text).replace("\n", " ")
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def _fill_paragraph(p, inlines: list[Inline]) -> None:
    for i in inlines:
        if not i.text:
            continue
        if i.url:
            _add_hyperlink(p, i.text, i.url, i.bold, i.italic)
            continue
        run = p.add_run(clean_xml_text(i.text).replace("\t", "    "))
        run.bold = i.bold or None
        run.italic = i.italic or None
        if i.code:
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)


def to_docx(blocks: list, path: str | Path, title: str = "") -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.0)
    for b in blocks:
        if isinstance(b, Heading):
            p = doc.add_heading(level=min(b.level, 9))
            _fill_paragraph(p, b.inlines)
        elif isinstance(b, Para):
            try:
                p = doc.add_paragraph(style="Quote") if b.quote else doc.add_paragraph()
            except KeyError:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1)
            _fill_paragraph(p, b.inlines)
        elif isinstance(b, ListItem):
            if b.ordered:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.9 + 0.8 * b.level)
                p.paragraph_format.first_line_indent = Cm(-0.6)
                p.paragraph_format.space_after = Pt(2)
                p.add_run(f"{b.number}. ")
            else:
                style = "List Bullet" if b.level == 0 else f"List Bullet {min(b.level + 1, 3)}"
                p = doc.add_paragraph(style=style)
            _fill_paragraph(p, b.inlines)
        elif isinstance(b, Code):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            for k, line in enumerate(b.text.split("\n")):
                run = p.add_run(clean_xml_text(line).replace("\t", "    "))
                run.font.name, run.font.size = "Consolas", Pt(9)
                if k < b.text.count("\n"):
                    run.add_break()
        elif isinstance(b, Rule):
            p = doc.add_paragraph()
            ppr = p._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "999999")):
                bottom.set(qn(k), v)
            bd.append(bottom)
            ppr.append(bd)
        elif isinstance(b, TableB):
            cols = max(len(r) for r in b.rows)
            t = doc.add_table(rows=len(b.rows), cols=cols)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, r in enumerate(b.rows):
                for ci, c in enumerate(r):
                    cell = t.cell(ri, ci)
                    par = cell.paragraphs[0]
                    items = c if not (ri == 0 and b.header) else [Inline(i.text, True, i.italic, i.code, i.url) for i in c]
                    _fill_paragraph(par, items)
            if b.header:
                trpr = t.rows[0]._tr.get_or_add_trPr()
                trpr.append(OxmlElement("w:tblHeader"))
            doc.add_paragraph()
    if title:
        doc.core_properties.title = clean_xml_text(title)
    doc.core_properties.comments = "Создано в DB Workspace Pro"
    doc.save(str(path))


def extract_tables(blocks: list) -> list[tuple[str, list[str], list[tuple]]]:
    """Таблицы документа: (имя, заголовки, строки). Имя — ближайший предыдущий заголовок или «Таблица N»."""
    out = []
    last_heading = ""
    n = 0
    for b in blocks:
        if isinstance(b, Heading):
            last_heading = plain(b.inlines).strip()
        elif isinstance(b, TableB) and b.rows:
            n += 1
            head = [plain(c).strip() or f"Столбец {i + 1}" for i, c in enumerate(b.rows[0])]
            rows = [tuple(plain(c).strip() for c in r) for r in b.rows[1:]]
            out.append((last_heading or f"Таблица {n}", head, rows))
    return out
