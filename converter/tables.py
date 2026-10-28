"""Чтение и запись таблиц: CSV, TSV, XLSX, JSON (потоково там, где формат это позволяет)."""
from __future__ import annotations

import codecs
import csv
import datetime as dt
import itertools
import json
import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Iterable, Iterator

from utils.export import ExportError, open_writer, unique_names
from utils.files import sanitize_filename

csv.field_size_limit(2 ** 31 - 1)
BATCH = 2000
_DELIMS = [",", ";", "\t", "|"]


class TableError(Exception):
    pass


@dataclass
class Sheet:
    name: str
    columns: list[str]
    rows: Iterator[tuple]


def _col_names(header: Iterable) -> list[str]:
    names = []
    for i, v in enumerate(header):
        s = "" if v is None else str(v).strip()
        names.append(s or f"Столбец {i + 1}")
    return names


# ============================================================ определение кодировки
def detect_encoding(path: Path) -> tuple[str, bool]:
    """(кодировка, это_догадка). BOM и UTF-8 — надёжно; иначе предполагается cp1251."""
    with open(path, "rb") as fh:
        sample = fh.read(262144)
    if sample.startswith(codecs.BOM_UTF8):
        return "utf-8-sig", False
    if sample.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return "utf-16", False
    try:
        codecs.getincrementaldecoder("utf-8")().decode(sample, final=False)
        return "utf-8", False
    except UnicodeDecodeError:
        return "cp1251", True


def sniff_delimiter(path: Path, encoding: str) -> str:
    with open(path, "r", encoding=encoding, errors="replace", newline="") as fh:
        sample = fh.read(65536)
    try:
        return csv.Sniffer().sniff(sample, delimiters="".join(_DELIMS)).delimiter
    except csv.Error:
        first = sample.split("\n", 1)[0]
        best = max(_DELIMS, key=lambda d: first.count(d))
        return best if first.count(best) else ","


# ===================================================================== читатели
class _Reader:
    names: list[str]
    warnings: list[str]

    def open(self, name: str) -> Sheet:  # pragma: no cover
        raise NotImplementedError

    def close(self) -> None:
        pass

    def __enter__(self):
        return self

    def __exit__(self, *a):
        self.close()


class MemReader(_Reader):
    """Таблицы, уже извлечённые из документа."""

    def __init__(self, tables: list[tuple[str, list[str], list[tuple]]]):
        self._t = {}
        self.names = []
        for name, cols, rows in tables:
            base, n = name, 1
            while name in self._t:
                n += 1
                name = f"{base} ({n})"
            self._t[name] = (cols, rows)
            self.names.append(name)
        self.warnings = []

    def open(self, name):
        cols, rows = self._t[name]
        return Sheet(name, list(cols), iter(rows))


class CsvReader(_Reader):
    def __init__(self, path: Path, fmt: str):
        self.path = Path(path)
        self.encoding, guessed = detect_encoding(self.path)
        self.delimiter = "\t" if fmt == "tsv" else sniff_delimiter(self.path, self.encoding)
        self.names = [self.path.stem]
        self.warnings = ["Кодировка файла определена как cp1251 (предположение) — проверьте кириллицу."] if guessed else []
        self.skipped_long = 0

    def _reader(self, fh):
        return csv.reader(fh, delimiter=self.delimiter)

    def open(self, name):
        """Заголовок читается сразу; файл для строк открывается только на время их чтения (без висящих дескрипторов)."""
        header, consumed = None, 0
        with open(self.path, "r", encoding=self.encoding, errors="replace", newline="") as fh:
            for row in self._reader(fh):
                consumed += 1
                if any(c.strip() for c in row):
                    header = row
                    break
        if header is None:
            return Sheet(name, [], iter(()))
        columns = _col_names(header)
        width = len(columns)

        def rows():
            with open(self.path, "r", encoding=self.encoding, errors="replace", newline="") as fh:
                reader = self._reader(fh)
                for _ in range(consumed):
                    next(reader, None)
                for row in reader:
                    if not any(c.strip() for c in row):
                        continue
                    if len(row) > width:
                        self.skipped_long += 1
                    yield tuple((row + [""] * width)[:width])
        return Sheet(name, columns, rows())


def count_uncached_formulas(path: Path, limit_bytes: int = 300 * 1048576) -> int | None:
    """Сколько ячеек содержат формулу, но не имеют сохранённого значения (None — проверка пропущена)."""
    import zipfile
    from xml.etree.ElementTree import iterparse
    try:
        with zipfile.ZipFile(path) as z:
            sheets = [i for i in z.infolist() if i.filename.startswith("xl/worksheets/") and i.filename.endswith(".xml")]
            if sum(i.file_size for i in sheets) > limit_bytes:
                return None
            total = 0
            for info in sheets:
                with z.open(info) as fh:
                    for _ev, el in iterparse(fh, events=("end",)):
                        if el.tag.endswith("}c"):
                            kids = {k.tag.rsplit("}", 1)[-1]: k for k in el}
                            if "f" in kids and ("v" not in kids or not (kids["v"].text or "").strip()):
                                total += 1
                            el.clear()
            return total
    except Exception:  # noqa: BLE001
        return None


class XlsxReader(_Reader):
    def __init__(self, path: Path):
        try:
            from openpyxl import load_workbook
            self.wb = load_workbook(str(path), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise TableError(f"Не удалось открыть книгу Excel: {exc.__class__.__name__}") from exc
        self.names = list(self.wb.sheetnames)
        self.warnings = []
        missing = count_uncached_formulas(Path(path))
        if missing:
            self.warnings.append(
                f"В книге {missing} ячеек с формулами без сохранённых значений (файл создан не в Excel или не "
                "пересчитан): они перенесены как пустые. Откройте файл в Excel и сохраните его — тогда значения появятся.")

    def close(self):
        try:
            self.wb.close()
        except Exception:  # noqa: BLE001
            pass

    def open(self, name):
        ws = self.wb[name]
        it = ws.iter_rows(values_only=True)
        header = None
        for row in it:
            if any(v is not None and str(v).strip() != "" for v in row):
                header = row
                break
        if header is None:
            return Sheet(name, [], iter(()))
        peek = list(itertools.islice(it, 100))
        last = max([max((i for i, v in enumerate(r) if v is not None and v != ""), default=-1) for r in [header] + peek])
        width = last + 1
        columns = _col_names(list(header)[:width])
        columns += [f"Столбец {i + 1}" for i in range(len(columns), width)]

        def rows():
            for r in itertools.chain(peek, it):
                if all(v is None or v == "" for v in r):
                    continue
                yield tuple((list(r) + [None] * width)[:width])
        return Sheet(name, columns, rows())


class JsonReader(_Reader):
    def __init__(self, path: Path, max_mb: int = 200):
        size = Path(path).stat().st_size
        if size > max_mb * 1048576:
            raise TableError(f"JSON-файл слишком большой ({size / 1048576:.0f} МБ, максимум {max_mb} МБ): его нельзя "
                             "прочитать потоково. Разделите файл или используйте CSV.")
        try:
            with open(path, "r", encoding="utf-8-sig") as fh:
                data = json.load(fh)
        except UnicodeDecodeError as exc:
            raise TableError("JSON должен быть в кодировке UTF-8.") from exc
        except json.JSONDecodeError as exc:
            raise TableError(f"Файл JSON повреждён: {exc.msg} (строка {exc.lineno}).") from exc
        self.warnings = []
        self._sheets: dict[str, tuple[list[str], list[tuple]]] = {}
        if isinstance(data, list):
            self._add(Path(path).stem, data)
        elif isinstance(data, dict):
            if data and all(isinstance(v, list) for v in data.values()):
                for k, v in data.items():
                    self._add(str(k), v)
            else:
                self._add(Path(path).stem, [data])
        else:
            raise TableError("JSON должен содержать массив записей или объект.")
        self.names = list(self._sheets)

    @staticmethod
    def _val(v):
        return json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v

    def _add(self, name: str, items: list) -> None:
        if not items:
            self._sheets[name] = ([], [])
            return
        if all(isinstance(x, dict) for x in items):
            cols: list[str] = []
            seen = set()
            for d in items:
                for k in d:
                    if k not in seen:
                        seen.add(k)
                        cols.append(k)
            rows = [tuple(self._val(d.get(c)) for c in cols) for d in items]
        elif all(isinstance(x, (list, tuple)) for x in items):
            width = max(len(x) for x in items)
            cols = [f"Столбец {i + 1}" for i in range(width)]
            rows = [tuple(self._val(v) for v in (list(x) + [None] * width)[:width]) for x in items]
        else:
            raise TableError("Элементы JSON-массива должны быть либо объектами, либо массивами одной структуры.")
        self._sheets[name] = ([str(c) for c in cols], rows)

    def open(self, name):
        cols, rows = self._sheets[name]
        return Sheet(name, list(cols), iter(rows))


def open_reader(path: Path, fmt: str, *, max_json_mb: int = 200) -> _Reader:
    if fmt in ("csv", "tsv"):
        return CsvReader(path, fmt)
    if fmt == "xlsx":
        return XlsxReader(path)
    if fmt == "json":
        return JsonReader(path, max_json_mb)
    raise TableError(f"Формат {fmt} не читается этим модулем.")


# ======================================================== определение типов (CSV -> Excel)
_INT = re.compile(r"^-?(0|[1-9]\d{0,14})$")
_DEC_DOT = re.compile(r"^-?(0|[1-9]\d{0,14})\.\d{1,10}$")
_DEC_COMMA = re.compile(r"^-?(0|[1-9]\d{0,14}),\d{1,10}$")
_DATE = re.compile(r"^(\d{4})-(\d{2})-(\d{2})$")
_DATETIME = re.compile(r"^(\d{4})-(\d{2})-(\d{2})[ T](\d{2}):(\d{2})(?::(\d{2}))?$")


def infer_cell(s):
    """Осторожное распознавание: числа без ведущих нулей, ISO-даты. Остальное остаётся текстом."""
    if not isinstance(s, str):
        return s
    t = s.strip()
    if not t:
        return None
    try:
        if _INT.match(t):
            return int(t)
        if _DEC_DOT.match(t):
            return Decimal(t)
        if _DEC_COMMA.match(t):
            return Decimal(t.replace(",", "."))
        m = _DATE.match(t)
        if m:
            return dt.date(int(m[1]), int(m[2]), int(m[3]))
        m = _DATETIME.match(t)
        if m:
            return dt.datetime(int(m[1]), int(m[2]), int(m[3]), int(m[4]), int(m[5]), int(m[6] or 0))
    except (ValueError, ArithmeticError):
        return s
    return s


# ======================================================================== запись
def _batches(it: Iterator[tuple], n: int = BATCH):
    while True:
        chunk = list(itertools.islice(it, n))
        if not chunk:
            return
        yield chunk


@dataclass
class WriteResult:
    files: list[Path] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    rows: int = 0


def write_tables(reader: _Reader, dst: str, stage: Path, stem: str, *, delimiter: str = ";", formula_guard: bool = True,
                 infer: bool = False, cancel_token=None) -> WriteResult:
    """Записывает все непустые листы reader в формат dst (xlsx — одна книга; csv/tsv/json — файл на лист)."""
    res = WriteResult()
    sheets = []
    for name in reader.names:
        sh = reader.open(name)
        if sh.columns:
            sheets.append((name, sh))
        else:
            res.warnings.append(f"Лист «{name}» пуст — пропущен.")
    if not sheets:
        raise TableError("В исходном файле нет данных для преобразования.")
    conv = (lambda row: tuple(infer_cell(v) for v in row)) if infer else (lambda row: row)

    def check():
        if cancel_token is not None and cancel_token.is_cancelled:
            from .libreoffice import ConversionCancelled
            raise ConversionCancelled("Преобразование отменено.")

    try:
        if dst == "xlsx":
            path = stage / f"{stem}.xlsx"
            w = open_writer(path, "xlsx", sheets[0][1].columns, sheet_name=sheets[0][0], formula_guard=formula_guard)
            try:
                for i, (name, sh) in enumerate(sheets):
                    if i:
                        w.begin_sheet(name, sh.columns)
                    for chunk in _batches(sh.rows):
                        check()
                        w.write_rows([conv(r) for r in chunk])
                r = w.close()
            except BaseException:
                w.abort()
                raise
            res.rows, res.files = r.rows, [path]
            res.warnings += r.warnings
        else:
            used: set[str] = set()
            for name, sh in sheets:
                if len(sheets) == 1:
                    fname = f"{stem}.{dst}"
                else:
                    base = sanitize_filename(name, "sheet")[:60]
                    cand, n = base, 1
                    while cand.lower() in used:
                        n += 1
                        cand = f"{base} ({n})"
                    used.add(cand.lower())
                    fname = f"{stem}_{cand}.{dst}"
                path = stage / fname
                w = open_writer(path, dst if dst != "tsv" else "csv", sh.columns, delimiter="\t" if dst == "tsv" else delimiter,
                                formula_guard=formula_guard)
                try:
                    for chunk in _batches(sh.rows):
                        check()
                        w.write_rows(chunk)
                    r = w.close()
                except BaseException:
                    w.abort()
                    raise
                res.rows += r.rows
                res.files.append(path)
                res.warnings += [f"{fname}: {x}" for x in r.warnings]
            if len(sheets) > 1:
                res.warnings.append(f"В книге несколько листов — создано файлов: {len(sheets)}.")
    except ExportError as exc:
        raise TableError(str(exc)) from exc
    skipped = getattr(reader, "skipped_long", 0)
    if skipped:
        res.warnings.append(f"{skipped} строк были длиннее заголовка — лишние значения отброшены.")
    return res
