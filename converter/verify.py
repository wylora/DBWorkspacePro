"""Контроль качества: проверка результата преобразования доступными средствами."""
from __future__ import annotations

import json
import zipfile
from html.parser import HTMLParser
from pathlib import Path


class VerifyError(Exception):
    pass


def _utf8(path: Path, limit: int = 20 * 1048576) -> str:
    data = path.read_bytes()[:limit]
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        if len(data) == limit:
            return data[: exc.start].decode("utf-8-sig", errors="ignore")
        raise VerifyError("Файл не в кодировке UTF-8.") from exc


def verify_output(path: Path, fmt: str) -> list[str]:
    """Бросает VerifyError, если файл повреждён; возвращает мелкие замечания."""
    path = Path(path)
    if not path.is_file():
        raise VerifyError("Выходной файл не создан.")
    size = path.stat().st_size
    if size == 0:
        raise VerifyError("Выходной файл пустой.")
    notes: list[str] = []
    with open(path, "rb") as fh:
        head = fh.read(8)
    try:
        if fmt == "docx":
            from docx import Document
            Document(str(path))
        elif fmt == "xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(str(path), read_only=True)
            if not wb.sheetnames:
                raise VerifyError("В книге Excel нет листов.")
            wb.close()
        elif fmt == "pdf":
            if not head.startswith(b"%PDF-"):
                raise VerifyError("Нет сигнатуры PDF.")
            with open(path, "rb") as fh:
                fh.seek(max(0, size - 2048))
                if b"%%EOF" not in fh.read():
                    raise VerifyError("PDF оборван (нет маркера конца файла).")
            try:
                from pypdf import PdfReader
                if len(PdfReader(str(path)).pages) == 0:
                    raise VerifyError("В PDF нет страниц.")
            except ImportError:
                pass
        elif fmt in ("odt", "ods"):
            with zipfile.ZipFile(path) as z:
                names = z.namelist()
                if "content.xml" not in names or "mimetype" not in names:
                    raise VerifyError("Файл OpenDocument неполный.")
                if not z.read("mimetype").startswith(b"application/vnd.oasis.opendocument"):
                    raise VerifyError("Неверный тип OpenDocument.")
        elif fmt == "rtf":
            if not head.startswith(b"{\\rtf"):
                raise VerifyError("Нет сигнатуры RTF.")
        elif fmt in ("doc", "xls"):
            if not head.startswith(b"\xd0\xcf\x11\xe0"):
                raise VerifyError(f"Нет сигнатуры {fmt.upper()} (OLE).")
        elif fmt == "html":
            text = _utf8(path)
            if "<" not in text:
                raise VerifyError("В HTML нет разметки.")
            HTMLParser().feed(text)
        elif fmt == "json":
            if size <= 100 * 1048576:
                text = _utf8(path, limit=size)                       # проверяем файл целиком, а не только начало
                try:
                    json.loads(text)
                except json.JSONDecodeError as exc:
                    raise VerifyError(f"Получен некорректный JSON: {exc.msg}.") from exc
            else:                                                    # очень большой файл: начало и конец
                with open(path, "rb") as fh:
                    first = fh.read(1024).lstrip(b"\xef\xbb\xbf \r\n\t")[:1]
                    fh.seek(max(0, size - 1024))
                    last = fh.read().rstrip(b" \r\n\t")[-1:]
                _utf8(path)
                if first not in (b"[", b"{") or last not in (b"]", b"}"):
                    raise VerifyError("Получен некорректный или оборванный JSON.")
        elif fmt in ("md", "txt", "sql", "csv", "tsv"):
            _utf8(path)
        else:
            notes.append("Формат результата не имеет встроенной проверки.")
    except VerifyError:
        raise
    except zipfile.BadZipFile as exc:
        raise VerifyError("Файл повреждён (не удаётся прочитать архив).") from exc
    except Exception as exc:  # noqa: BLE001
        raise VerifyError(f"Файл повреждён или не открывается ({exc.__class__.__name__}).") from exc
    return notes
